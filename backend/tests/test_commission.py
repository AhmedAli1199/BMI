"""BMI's commission structure: each clause of Matt's note as a worked example."""
import io
import uuid
from datetime import date, datetime, timezone

import openpyxl
import pytest

from app.models import CommissionStatement, SalesEdition, SalesEditionCost, SalesOrder, SalesOrderCredit, SalesRep, SalesTitle
from app.sales import commission as cm
from app.sales.commission_seed import load_structure
from app.sales.reference import ensure_reference_data
from tests.conftest import identity_headers, make_user


@pytest.fixture()
def w(db_session):
    ensure_reference_data(db_session)
    load_structure(db_session)
    titles = {t.slug: t for t in db_session.query(SalesTitle).all()}
    reps = {r.code: r for r in db_session.query(SalesRep).all()}
    eds: dict = {}

    def ed(slug, name, pub, kind="issue"):
        key = (slug, name)
        if key not in eds:
            e = SalesEdition(id=uuid.uuid4(), title_id=titles[slug].id, year=pub.year, name=name, edition_date=pub, kind=kind)
            db_session.add(e)
            db_session.flush()
            eds[key] = e
        return eds[key]

    def book(edition, client, value, rep, booked, *, split=None, agency=None, company_id=None, status="booked"):
        o = SalesOrder(id=uuid.uuid4(), edition_id=edition.id, client_name=client, value_gbp=value, rep_id=reps[rep].id, booked_on=booked,
                       agency_commission_gbp=agency, company_id=company_id, status=status)
        db_session.add(o)
        db_session.flush()
        for code, amt in (split or {rep: value}).items():
            db_session.add(SalesOrderCredit(id=uuid.uuid4(), order_id=o.id, rep_id=reps[code].id, amount_gbp=amt))
        db_session.flush()
        return o

    admin = make_user(db_session, role="admin")
    return {"db": db_session, "ed": ed, "book": book, "reps": reps, "titles": titles, "admin": admin, "h": identity_headers(admin)}


def st(w, code, period):
    return cm.statement(cm.Ctx.build(w["db"]), w["reps"][code], period)


def test_sally_magazine_base_and_new_business(w):
    old = w["ed"]("selling-travel", "Jan 2025", date(2025, 1, 15))
    w["book"](old, "Returning Tours", 4000, "SP", date(2024, 12, 1))
    mar = w["ed"]("selling-travel", "Mar 2026", date(2026, 3, 10))
    w["book"](mar, "Returning Tours", 6000, "SP", date(2026, 1, 20))
    w["book"](mar, "Brand New Cruises", 3000, "SP", date(2026, 1, 21))
    w["book"](mar, "Shared Client", 1000, "SP", date(2026, 1, 22), split={"SP": 500, "ST": 500})
    w["book"](mar, "Cancelled Co", 9000, "SP", date(2026, 1, 22), status="cancelled")
    s = st(w, "SP", "2026-03")
    by = {ln["client"]: ln for ln in s["lines"]}
    assert set(by) == {"Returning Tours", "Brand New Cruises", "Shared Client"}
    assert by["Returning Tours"]["new_business"] == "returning" and "Jan 2025" in by["Returning Tours"]["new_business_reason"]
    assert by["Brand New Cruises"]["new_business"] == "new" and by["Brand New Cruises"]["new_business_gbp"] == 60
    assert by["Shared Client"]["share_gbp"] == 500 and by["Shared Client"]["split"]
    # 5.5% of 9,500 + 2% on 3,000 new + 2% on the 500 share of the new shared client
    assert s["totals"]["base_gbp"] == 522.5 and s["totals"]["new_business_gbp"] == 70 and s["totals"]["total_gbp"] == 592.5
    # Steven's half of the shared booking is at his own 5%
    assert st(w, "ST", "2026-03")["totals"]["base_gbp"] == 25


def test_new_business_rule_24_months_same_day_and_later_bookings(w):
    e = w["ed"]("selling-travel", "Apr 2026", date(2026, 4, 10))
    w["book"](w["ed"]("selling-travel", "Feb 2024", date(2024, 2, 1)), "Long Gone Ltd", 2000, "SP", date(2024, 1, 5))
    w["book"](e, "Long Gone Ltd", 1000, "SP", date(2026, 2, 1))           # last spent 25 months before: new again
    w["book"](e, "Fresh Air", 1000, "SP", date(2026, 1, 1))               # first booking
    w["book"](e, "Fresh Air", 1000, "SP", date(2026, 1, 1))               # same day: part of the first deal
    w["book"](e, "Fresh Air", 1000, "SP", date(2026, 1, 10))              # later: returning
    lines = st(w, "SP", "2026-04")["lines"]
    got = sorted((ln["client"], ln["booked_on"], ln["new_business"]) for ln in lines)
    assert got == [("Fresh Air", "2026-01-01", "new"), ("Fresh Air", "2026-01-01", "new"), ("Fresh Air", "2026-01-10", "returning"),
                   ("Long Gone Ltd", "2026-02-01", "new")]
    # a longer first deal is a setting
    cm.settings(w["db"]).first_deal_days = 30
    w["db"].flush()
    assert all(ln["new_business"] == "new" for ln in st(w, "SP", "2026-04")["lines"])


def test_spend_on_another_title_and_similar_names(w, client):
    w["book"](w["ed"]("tbtm", "Winter 2025", date(2025, 12, 1)), "Delta Air Lines", 5000, "KH", date(2025, 10, 1))
    e = w["ed"]("selling-travel", "May 2026", date(2026, 5, 10))
    exact = w["book"](e, "Delta Air Lines Inc", 1000, "SP", date(2026, 3, 1))   # same customer, any product: returning
    similar = w["book"](e, "Delta", 1000, "SP", date(2026, 3, 2))               # similar name: a person decides
    lines = {ln["client"]: ln for ln in st(w, "SP", "2026-05")["lines"]}
    assert lines["Delta Air Lines Inc"]["new_business"] == "returning"
    assert lines["Delta"]["new_business"] == "check" and lines["Delta"]["new_business_gbp"] == 0
    r = client.post("/api/commission/statement/approve", json={"rep_id": str(w["reps"]["SP"].id), "period": "2026-05"}, headers=w["h"])
    assert r.status_code == 409 and "need a decision" in r.json()["detail"]
    r = client.put(f"/api/commission/orders/{similar.id}/new-business", json={"decision": "new", "reason": "Delta Holidays is a different company"}, headers=w["h"])
    assert r.json()["status"] == "new" and r.json()["decided_by"] == "manager"
    s = st(w, "SP", "2026-05")
    assert s["checks"] == 0 and s["totals"]["new_business_gbp"] == 20
    assert exact.id  # untouched


def test_david_temporary_new_business_rate(w):
    for pub, nb in ((date(2027, 5, 20), 50), (date(2027, 6, 10), 20)):
        e = w["ed"]("selling-australia", f"Issue {pub:%b}", pub)
        w["book"](e, f"Newco {pub:%b}", 1000, "DW", date(2027, 1, 5))
        s = st(w, "DW", pub.strftime("%Y-%m"))
        assert s["totals"]["base_gbp"] == 50 and s["totals"]["new_business_gbp"] == nb


def test_sue_flat_rate_and_agency_cut(w):
    e = w["ed"]("obh", "105", date(2026, 3, 19))
    w["book"](e, "Airline One", 2000, "SW", date(2026, 1, 1), agency=200)
    s = st(w, "SW", "2026-03")
    assert s["lines"][0]["share_gbp"] == 1800 and s["totals"]["total_gbp"] == 315  # 17.5% of 1,800, no new-business top-up


def test_guide_bonus_and_newsletter_threshold(w, client):
    feb = w["ed"]("selling-travel-supplements", "Guide to the Caribbean 2026", date(2026, 2, 15), kind="guide")
    w["book"](feb, "Caribbean Board", 6000, "SP", date(2025, 11, 1))
    r = client.put(f"/api/commission/editions/{feb.id}/new-guide", json={"on": True, "rep_id": str(w["reps"]["SP"].id)}, headers=w["h"])
    assert r.status_code == 200 and r.json()["new_contract_guide"]
    s = st(w, "SP", "2026-02")
    assert [b["kind"] for b in s["bonuses"]] == ["new_guide"] and s["totals"]["bonuses_gbp"] == 200
    assert s["lines"][0]["base_rate"] == 0.025
    jan = w["ed"]("visit-usa-online", "Jan 2026", date(2026, 1, 31))
    mar = w["ed"]("visit-usa-online", "Mar 2026", date(2026, 3, 31))
    apr = w["ed"]("visit-usa-online", "Apr 2026", date(2026, 4, 30))
    w["book"](jan, "Brand USA", 5000, "SP", date(2025, 12, 1))
    w["book"](mar, "Brand USA", 4000, "SP", date(2026, 2, 1))
    w["book"](apr, "Brand USA", 2000, "SP", date(2026, 3, 1))
    assert not [b for b in st(w, "SP", "2026-01")["bonuses"] if b["kind"] == "threshold"]
    assert [b["amount_gbp"] for b in st(w, "SP", "2026-03")["bonuses"] if b["kind"] == "threshold"] == [200]
    assert not [b for b in st(w, "SP", "2026-04")["bonuses"] if b["kind"] == "threshold"]


def test_kirsty_new_client_bonus_once(w):
    a = w["ed"]("tbtm", "Spring 2026", date(2026, 3, 1))
    b = w["ed"]("tbtm", "Summer 2026", date(2026, 6, 1))
    w["book"](a, "Corporate Travel Co", 3000, "KH", date(2026, 1, 5))
    w["book"](a, "Corporate Travel Co", 1000, "KH", date(2026, 1, 5))
    w["book"](b, "Corporate Travel Co", 3000, "KH", date(2026, 1, 5))   # same first deal, publishes later
    s = st(w, "KH", "2026-03")
    assert [x["amount_gbp"] for x in s["bonuses"]] == [100]
    assert s["totals"]["base_gbp"] == 200 and s["totals"]["new_business_gbp"] == 80
    assert st(w, "KH", "2026-06")["bonuses"] == []


def test_kirsty_event_profit_after_sign_off(w, client, db_session):
    kirsty = make_user(db_session, role="sales")
    w["reps"]["KH"].user_id = kirsty.id
    db_session.flush()
    kh = identity_headers(kirsty)
    ev = w["ed"]("tbtm-events", "Dinner Club Spring", date(2026, 4, 22), kind="event")
    w["book"](ev, "Sponsor A", 12000, "KH", date(2026, 2, 1))
    w["book"](ev, "Sponsor B", 8000, "SW", date(2026, 2, 1))
    for label, amt in (("Venue", 9000), ("Food and drink", 3500)):
        assert client.post(f"/api/sales/editions/{ev.id}/costs", json={"kind": "cost", "label": label, "amount_gbp": amt}, headers=kh).status_code == 201
    s = st(w, "KH", "2026-04")
    assert s["events"] == [] and s["totals"]["base_gbp"] == 240   # 2% of her own 12,000; no profit share until signed off
    assert client.post(f"/api/commission/editions/{ev.id}/costs-final", json={"on": True}, headers=kh).json()["costs_final_at"]
    assert client.post(f"/api/commission/editions/{ev.id}/costs-signoff", json={"on": True}, headers=kh).status_code == 403
    assert client.post(f"/api/commission/editions/{ev.id}/costs-signoff", json={"on": True}, headers=w["h"]).status_code == 200
    # locked once signed off
    assert client.post(f"/api/sales/editions/{ev.id}/costs", json={"kind": "cost", "label": "Late bill", "amount_gbp": 10}, headers=kh).status_code == 409
    month = date.today().strftime("%Y-%m")
    e = st(w, "KH", month)["events"][0]
    assert (e["income_gbp"], e["costs_gbp"], e["profit_gbp"], e["amount_gbp"]) == (20000, 12500, 7500, 750)
    # a loss pays nothing
    ev2 = w["ed"]("tbtm-events", "Lunch Forum", date(2026, 5, 1), kind="event")
    w["book"](ev2, "Sponsor C", 1000, "KH", date(2026, 2, 1))
    db_session.add(SalesEditionCost(id=uuid.uuid4(), edition_id=ev2.id, kind="cost", label="Venue", amount_gbp=3000))
    ev2.costs_final_at = ev2.costs_signed_off_at = datetime.now(timezone.utc)
    db_session.flush()
    loss = next(x for x in st(w, "KH", month)["events"] if x["edition"].endswith("Lunch Forum"))
    assert loss["loss"] and loss["amount_gbp"] == 0


def test_approve_freezes_and_later_changes_carry_forward(w, client):
    e = w["ed"]("obh", "104", date(2026, 1, 20))
    o = w["book"](e, "Airline Two", 1000, "SW", date(2025, 12, 1))
    r = client.post("/api/commission/statement/approve", json={"rep_id": str(w["reps"]["SW"].id), "period": "2026-01"}, headers=w["h"])
    assert r.status_code == 200 and r.json()["approved"] and r.json()["totals"]["total_gbp"] == 175
    o.value_gbp = 2000
    w["db"].query(SalesOrderCredit).filter_by(order_id=o.id).one().amount_gbp = 2000
    w["db"].flush()
    jan = st(w, "SW", "2026-01")
    assert jan["totals"]["total_gbp"] == 175 and jan["approved"]["changed_since_gbp"] == 175
    feb = st(w, "SW", "2026-02")
    assert feb["adjustments"][0]["amount_gbp"] == 175 and feb["totals"]["total_gbp"] == 175
    client.post("/api/commission/statement/approve", json={"rep_id": str(w["reps"]["SW"].id), "period": "2026-02"}, headers=w["h"])
    assert st(w, "SW", "2026-03")["adjustments"] == []   # paid once
    assert st(w, "SW", "2026-01")["approved"]["changed_since_gbp"] == 0
    assert w["db"].query(CommissionStatement).count() == 2


def test_not_in_plan_fallback_scoping_and_excel(w, client, db_session):
    e = w["ed"]("selling-australia", "Aus 2026", date(2026, 2, 1))
    w["book"](e, "Tourism Australia", 1000, "SP", date(2025, 6, 1))   # Sally has no Selling Australia rule
    s = st(w, "SP", "2026-02")
    assert s["lines"][0]["group"] == cm.FALLBACK and s["lines"][0]["base_rate"] == 0.02 and s["lines"][0]["flags"]
    sally = make_user(db_session, role="sales")
    w["reps"]["SP"].user_id = sally.id
    db_session.flush()
    h = identity_headers(sally)
    assert client.get(f"/api/commission/statement?rep_id={w['reps']['SW'].id}&period=2026-02", headers=h).status_code == 404
    ov = client.get("/api/commission/overview?year=2026", headers=h).json()
    assert [r["rep"]["code"] for r in ov["reps"]] == ["SP"] and ov["scoped_to_me"]
    x = client.get(f"/api/commission/statement.xlsx?rep_id={w['reps']['SP'].id}&period=2026-02", headers=h)
    ws = openpyxl.load_workbook(io.BytesIO(x.content)).active
    assert "Sally Parker" in ws["A1"].value and ws.cell(5, 4).value == "Tourism Australia"
    assert client.post("/api/commission/plans/load", headers=h).status_code == 403


def test_one_off_supplements_only_compare_by_name(w):
    from app.sales.analytics import equivalent_editions
    lc = w["ed"]("selling-travel-supplements", "Los Cabos", date(2025, 3, 1))
    st_louis = w["ed"]("selling-travel-supplements", "St Louis", date(2026, 3, 1))
    carib_old = w["ed"]("selling-travel-supplements", "Caribbean annual 2025", date(2025, 11, 1))
    carib = w["ed"]("selling-travel-supplements", "Caribbean annual 2026", date(2026, 11, 1))
    jan25 = w["ed"]("selling-travel", "JanFeb", date(2025, 1, 15))
    jan26 = w["ed"]("selling-travel", "JanFebMar", date(2026, 1, 20))
    m = equivalent_editions(w["db"], [st_louis, carib, jan26])
    assert st_louis.id not in m and lc.id  # a different destination isn't "last year's edition"
    assert m[carib.id].id == carib_old.id   # the same guide a year on is
    assert m[jan26.id].id == jan25.id       # regular issues still fall back to the nearest date
