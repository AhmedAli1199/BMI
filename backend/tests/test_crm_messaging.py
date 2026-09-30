"""Lookup, export, duplicate, group-from-selection, reminders,
notifications and mail merge."""
import io
import uuid
from datetime import datetime, timedelta, timezone

import openpyxl
import pytest

from app.models import Company, Contact, Email, Group, GroupMembership, HistoryEntry
from app.models.contact_channel import Address
from app.models.messaging import MailAccount, MailMerge, MailMergeRecipient, Notification
from app.services import mail_merge as mm
from app.services import notify, outlook
from app.services.reminders import fire_due_reminders
from tests.conftest import identity_headers, make_user


def _contact(db, first, last, *, email=None, city=None, country=None, company=None, title=None, **kw):
    c = Contact(id=uuid.uuid4(), source_db="manual", source_act_id=str(uuid.uuid4()), first_name=first, last_name=last,
                full_name=f"{first} {last}", job_title=title, company_id=company.id if company else None, **kw)
    db.add(c)
    db.flush()
    if email:
        db.add(Email(id=uuid.uuid4(), source_db="manual", source_act_id=str(uuid.uuid4()), contact_id=c.id,
                     address=email, is_primary=True))
    if city or country:
        db.add(Address(id=uuid.uuid4(), source_db="manual", source_act_id=str(uuid.uuid4()), contact_id=c.id,
                       line1="1 High St", city=city, country=country, postal_code="AB1 2CD", is_primary=True))
    db.flush()
    return c


@pytest.fixture()
def people(db_session):
    va = Company(id=uuid.uuid4(), source_db="manual", source_act_id=str(uuid.uuid4()), name="Zeta Airways")
    db_session.add(va)
    db_session.flush()
    return {
        "amy": _contact(db_session, "Amy", "Zed", email="amy@zeta.com", city="London", country="UK", company=va, title="Buyer"),
        "bob": _contact(db_session, "Bob", "Young", email="bob@x.com", city="Paris", country="France", title="CEO"),
        "cat": _contact(db_session, "Cat", "Xavier", city="Leeds", country="UK", title="Head Buyer"),
        "company": va,
    }


def test_lookup_filters_and_sort(client, people):
    r = client.get("/api/contacts", params={"country": "UK", "sort": "city"}).json()
    names = [i["first_name"] for i in r["items"] if i["first_name"] in ("Amy", "Bob", "Cat")]
    assert names == ["Cat", "Amy"]  # Leeds, London
    assert client.get("/api/contacts", params={"company": "zeta"}).json()["items"][0]["city"] == "London"
    r = client.get("/api/contacts", params={"title": "buyer", "sort": "name", "desc": True}).json()
    assert [i["last_name"] for i in r["items"] if i["last_name"] in ("Zed", "Xavier")] == ["Zed", "Xavier"]
    ids = client.get("/api/contacts/lookup-ids", params={"city": "paris"}).json()
    assert str(people["bob"].id) in ids


def test_export_selection_to_xlsx(client, people):
    r = client.get("/api/contacts/export", params={"ids": [str(people["amy"].id), str(people["bob"].id)], "sort": "first_name"})
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.values)
    assert rows[0][0] == "First name" and [x[0] for x in rows[1:]] == ["Amy", "Bob"]
    assert rows[1][4] == "Zeta Airways" and rows[1][5] == "amy@zeta.com"


def test_duplicate_contact_copies_company_and_groups(client, db_session, people):
    g = Group(id=uuid.uuid4(), name="Buyers", source_db="manual", source_act_id=str(uuid.uuid4()))
    db_session.add(g)
    db_session.flush()
    db_session.add(GroupMembership(id=uuid.uuid4(), group_id=g.id, contact_id=people["amy"].id))
    db_session.flush()
    r = client.post(f"/api/contacts/{people['amy'].id}/duplicate",
                    json={"first_name": "Ann", "last_name": "Other", "email": "ann@zeta.com", "copy_groups": True})
    assert r.status_code == 201, r.text
    new = r.json()
    assert db_session.get(Contact, uuid.UUID(new["id"])).company_id == people["company"].id
    assert db_session.query(GroupMembership).filter_by(contact_id=uuid.UUID(new["id"]), group_id=g.id).count() == 1


def test_group_from_selection_and_bulk_add(client, db_session, people):
    r = client.post("/api/groups", json={"name": "Mailing Oct", "contact_ids": [str(people["amy"].id), str(people["bob"].id)]})
    assert r.status_code == 201, r.text
    gid = r.json()["id"]
    r = client.post(f"/api/groups/{gid}/members/add", json={"contact_ids": [str(people["bob"].id), str(people["cat"].id)]})
    assert r.json() == {"added": 1, "already_members": 1}
    ws = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/groups/{gid}/export").content)).active
    assert ws.max_row == 4


def test_reminder_fires_once_and_emails(client, db_session, people, monkeypatch):
    me = make_user(db_session, role="sales")
    h = identity_headers(me)
    due = datetime.now(timezone.utc) + timedelta(minutes=1)
    r = client.post("/api/reminders", headers=h, json={"due_at": due.isoformat(), "note": "Chase media pack",
                                                        "contact_id": str(people["amy"].id)})
    assert r.status_code == 201 and r.json()["about"] == "Amy Zed"
    assert fire_due_reminders(db_session) == 0  # not due yet
    assert fire_due_reminders(db_session, now=due + timedelta(seconds=5)) == 1
    assert fire_due_reminders(db_session, now=due + timedelta(minutes=10)) == 0  # exactly once
    n = client.get("/api/notifications", headers=h).json()
    assert n["unread"] == 1 and n["items"][0]["title"] == "Reminder: Amy Zed"
    assert n["items"][0]["link"] == f"/contacts/{people['amy'].id}"

    # No mailbox configured -> skipped, in-app only.
    assert notify.send_pending_emails(db_session) == (0, 0, 1)
    # With one, it sends.
    note = db_session.query(Notification).filter_by(user_id=me.id).one()
    note.email_status = "pending"
    db_session.flush()
    sent = []
    monkeypatch.setattr(notify, "automation_mailbox_ready", lambda db: True)
    monkeypatch.setattr(notify, "automation_mailbox", lambda db: "automation@bmi.test")
    monkeypatch.setattr(notify, "send_from_automation_mailbox", lambda mb, to, subj, body: sent.append((mb, to, subj)))
    assert notify.send_pending_emails(db_session) == (1, 0, 0)
    assert sent == [("automation@bmi.test", me.email, "Reminder: Amy Zed")]

    client.post("/api/notifications/read-all", headers=h)
    assert client.get("/api/notifications/unread-count", headers=h).json() == {"unread": 0}


def test_reminder_snooze_rearms_and_is_private(client, db_session):
    me, other = make_user(db_session), make_user(db_session)
    r = client.post("/api/reminders", headers=identity_headers(me),
                    json={"due_at": (datetime.now(timezone.utc) - timedelta(hours=1)).isoformat(), "email_me": False}).json()
    fire_due_reminders(db_session)
    assert db_session.get(Notification, db_session.query(Notification).filter_by(user_id=me.id).one().id).email_status is None
    later = (datetime.now(timezone.utc) + timedelta(days=1)).isoformat()
    r2 = client.patch(f"/api/reminders/{r['id']}", headers=identity_headers(me), json={"due_at": later}).json()
    assert r2["notified_at"] is None and r2["status"] == "open"
    assert client.patch(f"/api/reminders/{r['id']}", headers=identity_headers(other), json={"status": "done"}).status_code == 404
    assert client.get("/api/reminders", headers=identity_headers(other)).json() == []


def test_merge_field_rendering():
    ctx = {"first_name": "Amy", "company": ""}
    assert mm.render("Dear {{first_name}}, at {{ company | your company }}", ctx) == "Dear Amy, at your company"
    assert mm.unknown_fields("Hi {{frist_name}} {{first_name}}") == ["frist_name"]
    assert mm.to_html("a<b\nc") == '<div style="font-family:Calibri,Arial,sans-serif;font-size:11pt">a&lt;b<br>c</div>'


def test_recipients_preview_flags(client, db_session, people):
    me = make_user(db_session)
    people["bob"].is_unsubscribed = True
    db_session.flush()
    r = client.post("/api/mail-merge/recipients", headers=identity_headers(me), json={"kind": "lookup", "title": "buyer"}).json()
    got = {i["name"]: i for i in r["items"]}
    assert got["Amy Zed"]["email"] == "amy@zeta.com" and got["Cat Xavier"]["email"] is None
    r = client.post("/api/mail-merge/recipients", headers=identity_headers(me),
                    json={"kind": "company", "company_id": str(people["company"].id)}).json()
    assert r["label"] == "Company: Zeta Airways" and r["total"] == 1


@pytest.mark.parametrize("output,magic", [("word", b"PK"), ("labels", b"PK"), ("data", b"PK")])
def test_document_outputs(client, db_session, people, output, magic):
    me = make_user(db_session)
    ids = [str(people[k].id) for k in ("amy", "bob", "cat")]
    r = client.post("/api/mail-merge", headers=identity_headers(me), json={
        "output": output, "contact_ids": ids, "body": "Dear {{first_name}},\n\nThanks.", "subject": "Media pack",
        "record_history": "subject_only", "history_regarding": "Media pack letter"})
    assert r.status_code == 200, r.text
    assert r.content[:2] == magic
    letters = db_session.query(HistoryEntry).filter_by(history_type="Letter Sent").count()
    assert letters == (3 if output == "word" else 0)


def test_email_merge_needs_outlook_then_queues_and_sends(client, db_session, people, monkeypatch):
    me = make_user(db_session)
    h = identity_headers(me)
    payload = {"output": "email", "contact_ids": [str(people[k].id) for k in ("amy", "bob", "cat")],
               "subject": "Hello {{first_name}}", "body": "Hi {{first_name|there}}", "record_history": "email_full"}
    assert client.post("/api/mail-merge", headers=h, json=payload).status_code == 409
    db_session.add(MailAccount(id=uuid.uuid4(), user_id=me.id, email="me@bmi.test", refresh_token_enc=outlook.encrypt("rt")))
    db_session.flush()
    assert client.post("/api/mail-merge", headers=h, json={**payload, "body": "{{nope}}"}).status_code == 400

    m = client.post("/api/mail-merge", headers=h, json=payload).json()
    assert m["status"] == "queued" and m["total"] == 3 and m["skipped"] == 1  # Cat has no email

    sent = []

    def fake_send(db, acct, *, to, subject, html_body, **kw):
        if to == ["bob@x.com"]:
            raise RuntimeError("mailbox full")
        sent.append((to, subject))
    monkeypatch.setattr(outlook, "send_mail", fake_send)
    assert mm.send_queued(db_session, per_minute=10) == 1
    assert sent == [(["amy@zeta.com"], "Hello Amy")]
    detail = client.get(f"/api/mail-merge/{m['id']}", headers=h).json()
    assert (detail["status"], detail["sent"], detail["failed"], detail["skipped"]) == ("done", 1, 1, 1)
    hist = db_session.query(HistoryEntry).filter_by(entity_id=people["amy"].id, history_type="E-mail Sent").one()
    assert hist.subject == "Hello Amy" and hist.details == "Hi Amy"
    # Retry the failure.
    assert client.post(f"/api/mail-merge/{m['id']}/resume", headers=h).json()["status"] == "queued"
    monkeypatch.setattr(outlook, "send_mail", lambda db, acct, **kw: sent.append((kw["to"], kw["subject"])))
    mm.send_queued(db_session)
    assert client.get(f"/api/mail-merge/{m['id']}", headers=h).json()["sent"] == 2
    # Letters for the contact with no email.
    r = client.post(f"/api/mail-merge/{m['id']}/letters", headers=h)
    assert r.status_code == 200 and r.content[:2] == b"PK"


def test_auth_error_pauses_merge_and_notifies(db_session, people, monkeypatch):
    me = make_user(db_session)
    db_session.add(MailAccount(id=uuid.uuid4(), user_id=me.id, email="me@bmi.test", refresh_token_enc="x"))
    m = MailMerge(id=uuid.uuid4(), created_by_user_id=me.id, output="email", subject="S", body="B", status="queued", total=1)
    db_session.add(m)
    db_session.flush()
    db_session.add(MailMergeRecipient(id=uuid.uuid4(), merge_id=m.id, contact_id=people["amy"].id, email="amy@zeta.com"))
    db_session.flush()

    def boom(*a, **k):
        raise outlook.OutlookAuthError("expired")
    monkeypatch.setattr(outlook, "send_mail", boom)
    mm.send_queued(db_session)
    assert m.status == "failed" and "Reconnect" in m.error
    assert db_session.query(Notification).filter_by(user_id=me.id, kind="mail_merge").count() == 1


def test_outlook_state_and_token_encryption():
    uid = uuid.uuid4()
    state = outlook.make_state(uid, "/mail-merge")
    assert outlook.read_state(state) == (uid, "/mail-merge")
    with pytest.raises(outlook.OutlookAuthError):
        outlook.read_state(state[:-1] + ("0" if state[-1] != "0" else "1"))
    assert outlook.decrypt(outlook.encrypt("secret-refresh")) == "secret-refresh"


def test_templates_are_shared_but_owner_edits(client, db_session):
    a, b = make_user(db_session), make_user(db_session)
    t = client.post("/api/mail/templates", headers=identity_headers(a),
                    json={"name": "Renewal", "subject": "Renew {{company}}", "body": "Hi"}).json()
    assert [x["name"] for x in client.get("/api/mail/templates", headers=identity_headers(b)).json()] == ["Renewal"]
    assert client.put(f"/api/mail/templates/{t['id']}", headers=identity_headers(b),
                      json={"name": "X", "body": ""}).status_code == 403


def test_export_large_selection_by_post(client, people):
    r = client.post("/api/contacts/export", json={"ids": [str(people["cat"].id), str(people["amy"].id)], "title": "My pick"})
    assert r.status_code == 200 and 'filename="My pick.xlsx"' in r.headers["content-disposition"]
    rows = list(openpyxl.load_workbook(io.BytesIO(r.content)).active.values)
    assert [x[0] for x in rows[1:]] == ["Cat", "Amy"]  # kept in the order given
