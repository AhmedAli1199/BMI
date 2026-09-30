"""Sales Order Register: the SOR spreadsheet importer, the /sales API and
the three SOR automations. See app/models/sales.py."""
import uuid
from datetime import date, datetime, timedelta

import openpyxl
import pytest
from tests.conftest import identity_headers, make_user

import app.automations.sales_orders as sor_automations
from app.automations import get_kind
from app.models import Company, FieldChange, ReviewQueueItem, SalesEdition, SalesOrder, SalesOrderCredit, SalesRep, SalesTitle
from app.sales.matching import match_clients
from app.sales.reference import ensure_reference_data, rep_code_for
from app.sales.sor_import import import_sor, parse_date, parse_money, parse_pages, parse_sheet


# ---- Parsing ------------------------------------------------------------------

def test_parse_helpers():
    assert parse_pages("FP", "print") == 1
    assert parse_pages("DPS", "print") == 2
    assert parse_pages("2 x 1/2", "print") == 1
    assert parse_pages("5 FP", "print") == 5
    assert parse_pages("0.25", "print") == 0.25
    assert parse_pages("Banner", "print") is None
    assert parse_pages("FP", "digital") is None
    assert parse_money(1250.0) == (1250.0, None)
    assert parse_money("£1,250") == (1250.0, None)
    assert parse_money("40, 0000")[0] is None and parse_money("40, 0000")[1]
    assert parse_money("tbc") == (None, None)
    assert parse_date("15.05.25") == (date(2025, 5, 15), None)
    assert parse_date("19th/20th Jan")[0] is None
    assert rep_code_for("S.Thompson") == "ST" and rep_code_for("Steve") == "ST"
    assert rep_code_for(" sw ") == "SW" and rep_code_for("C.Blacwell") == "CB"


HEADER_TOP = [None, None, None, None, None, None, None, None, None, "Rate", None, None, "Invoice", "Invoice", "Page", "Invoice"]
HEADER = ["Date", "Client", None, None, None, "Size", "Series", "Position", "Salesper.", "US$", "£", None, "number", "value",
          "number", "difference", "Reason for difference", None, "S.Parker", "S.Thompson"]


def _sheet_rows(bookings: list[list]) -> list[list]:
    return [
        ["PUBLICATION", None, "OBH 105"],
        [],
        ["MONTH", None, "March/April/May", None, None, None, None, None, None, "Exchange rate", 1.35],
        [None, None, None, None, None, "Cumulative value of publication", None, None, None, None, 9999.0],
        HEADER_TOP,
        HEADER,
        *bookings,
    ]


def _row(date_, client, size, rep, gbp, inv=None, reason=None, sp=0.0, st=0.0):
    r = [date_, client, None, None, None, size, None, None, rep, None, gbp, None, inv, gbp if inv else None,
         None, 0.0, reason, None, sp, st]
    return r


def test_parse_sheet_reads_bookings_continuations_and_skips_cost_blocks():
    rows = _sheet_rows([
        _row("15.05.25", "Foodcase", "0.5", "SP", 1250.0, "INV-1", sp=1250.0),
        _row(None, None, None, None, 800.0, "INV-2"),  # second instalment, client left blank
        [],
        _row(None, "Room Hire", None, None, None),  # cost block typed into the client column
        [None, None, None, None, None, "Caledonian Club hire", None, None, None, None, 2257.45],
    ])
    sheet = parse_sheet("105", rows)
    assert sheet.exchange_rate == 1.35 and sheet.period == "March/April/May" and sheet.sheet_total == 9999.0
    assert [(r.client, r.continuation) for r in sheet.rows] == [("Foodcase", False), ("Foodcase", True)]
    assert sheet.rows[0].rep_credits == {"SP": 1250.0}


def test_parse_sheet_skips_templates_and_entry_lists():
    assert parse_sheet("TEMPLATE DO NOT COPY OVER", _sheet_rows([])) is None
    assert parse_sheet("Blank 2", _sheet_rows([])) is None
    assert parse_sheet("IND Entries ALPHA", [["Date", "Client", "No."], ["04.11.25", "Air Canada", 5]]) is None


def _write_book(path, sheets: dict[str, list[list]]):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(path)


def test_import_end_to_end_statuses_splits_and_planner_pair(db_session, tmp_path):
    y = tmp_path / "2026"
    y.mkdir()
    _write_book(y / "OBH 2026.xlsx", {"105": _sheet_rows([
        _row("15.05.25", "Foodcase", "FP", "SP", 2000.0, "INV-1", sp=2000.0),
        _row("16.05.25", "BA/Hungary", "3-pages", "SP/ST", 6000.0, "INV-2", sp=3000.0, st=3000.0),
        _row("17.05.25", "Driessen", "-", "SP", 0.0, reason="Cancelled CANX"),
        _row("18.05.25", "deSter", "-", "SP", 0.0, inv="Moved to OBH 107"),
        _row("19.05.25", "Monty's", "FP", "SP", 0.0, reason="Contra"),
    ]), "TEMPLATE DO NOT COPY OVER": _sheet_rows([])})
    _write_book(y / "OBH 2026-Copy(1).xlsx", {"105": _sheet_rows([_row("15.05.25", "Dupe", "FP", "SP", 1.0)])})
    planner = [_row("01.12.25", "CINCY Region", "FP", "SP", 5445.0, "INV-9", sp=5445.0)]
    working = [_row("01.12.25", "CINCY Region", "FP", "SP", 5445.0, "INV-9", reason="agency comm", sp=5445.0)]
    planner[0][13] = 4900.5
    _write_book(y / "Visit USA Travel Planner 2026 edition.xlsx", {"Travel Planner": _sheet_rows(planner)})
    _write_book(y / "Visit USA Travel Planner 2026.xlsx", {"Travel Planner": _sheet_rows(working)})

    report = import_sor(db_session, tmp_path)
    assert report.files == 2  # the Copy(1) and the planner working copy are not imported on their own
    orders = {o.client_name: o for o in db_session.query(SalesOrder).all()}
    assert set(orders) == {"Foodcase", "BA/Hungary", "Driessen", "deSter", "Monty's", "CINCY Region"}
    assert orders["Driessen"].status == "cancelled"
    assert orders["deSter"].status == "moved" and orders["deSter"].notes == "Moved to OBH 107"
    assert orders["Monty's"].status == "contra"
    assert float(orders["Foodcase"].pages) == 1

    credits = db_session.query(SalesOrderCredit).filter_by(order_id=orders["BA/Hungary"].id).all()
    reps = {r.id: r.code for r in db_session.query(SalesRep).all()}
    assert sorted((reps[c.rep_id], float(c.amount_gbp)) for c in credits) == [("SP", 3000.0), ("ST", 3000.0)]

    cincy = orders["CINCY Region"]
    assert cincy.invoice_note == "agency comm"  # merged from the working copy
    assert float(cincy.agency_commission_gbp) == pytest.approx(544.5)
    assert "Working copy showed an invoice of £5,445.00" in cincy.notes

    with pytest.raises(RuntimeError):
        import_sor(db_session, tmp_path)  # no silent double import
    assert import_sor(db_session, tmp_path, replace=True).orders == 6


# ---- API ------------------------------------------------------------------------

@pytest.fixture()
def sor(db_session):
    ensure_reference_data(db_session)
    title = db_session.query(SalesTitle).filter_by(slug="obh").one()
    reps = {r.code: r for r in db_session.query(SalesRep).all()}
    today = date.today()
    last = SalesEdition(id=uuid.uuid4(), title_id=title.id, year=today.year - 1, name="101", edition_date=date(today.year - 1, 3, 1))
    this = SalesEdition(id=uuid.uuid4(), title_id=title.id, year=today.year, name="105", edition_date=today - timedelta(days=30))
    db_session.add_all([last, this])
    db_session.flush()
    return {"title": title, "reps": reps, "last": last, "this": this}


def _order(db_session, ed, client, value, rep=None, booked_on=None, **kw):
    o = SalesOrder(id=uuid.uuid4(), edition_id=ed.id, client_name=client, value_gbp=value, rep_id=rep.id if rep else None,
                   booked_on=booked_on, **kw)
    db_session.add(o)
    db_session.flush()
    if rep and value:
        db_session.add(SalesOrderCredit(id=uuid.uuid4(), order_id=o.id, rep_id=rep.id, amount_gbp=value))
        db_session.flush()
    return o


def test_create_and_edit_order_keeps_credit_in_step_and_audits(client, db_session, sor):
    staff = make_user(db_session, role="admin")
    sp, st = sor["reps"]["SP"], sor["reps"]["ST"]
    resp = client.post("/api/sales/orders", headers=identity_headers(staff), json={
        "edition_id": str(sor["this"].id), "client_name": "Air Canada", "rep_id": str(sp.id), "value_gbp": 1500, "size": "FP"})
    assert resp.status_code == 201, resp.text
    order = resp.json()
    assert order["pages"] == 1 and order["credits"][0]["amount_gbp"] == 1500

    resp = client.patch(f"/api/sales/orders/{order['id']}", headers=identity_headers(staff),
                        json={"value_gbp": 1800, "rep_id": str(st.id)})
    assert [(c["code"], c["amount_gbp"]) for c in resp.json()["credits"]] == [("ST", 1800)]

    resp = client.patch(f"/api/sales/orders/{order['id']}", headers=identity_headers(staff),
                        json={"credits": [{"rep_id": str(sp.id), "amount_gbp": 900}, {"rep_id": str(st.id), "amount_gbp": 900}]})
    assert sorted(c["code"] for c in resp.json()["credits"]) == ["SP", "ST"]
    # a hand-set split is left alone by a later value edit
    resp = client.patch(f"/api/sales/orders/{order['id']}", headers=identity_headers(staff), json={"value_gbp": 2000})
    assert len(resp.json()["credits"]) == 2

    resp = client.patch(f"/api/sales/orders/{order['id']}", headers=identity_headers(staff), json={"invoice_number": "INV-5"})
    assert resp.json()["invoice_value_gbp"] == 2000 and resp.json()["invoiced_on"] == date.today().isoformat()

    changes = client.get(f"/api/sales/orders/{order['id']}/changes").json()
    assert {"value_gbp", "rep_id", "invoice_number"} <= {c["field"] for c in changes}
    assert all(c["changed_by"]["id"] == str(staff.id) for c in changes)


def test_only_staff_can_hard_delete(client, db_session, sor):
    o = _order(db_session, sor["this"], "Temp", 10)
    rep_user = make_user(db_session, role="sales")
    assert client.delete(f"/api/sales/orders/{o.id}", headers=identity_headers(rep_user)).status_code == 403
    assert client.delete(f"/api/sales/orders/{o.id}", headers=identity_headers(make_user(db_session, role="admin"))).status_code == 204


def test_edition_detail_compares_with_last_years_equivalent(client, db_session, sor):
    sp = sor["reps"]["SP"]
    _order(db_session, sor["last"], "Foodcase", 1000, sp, booked_on=sor["last"].edition_date - timedelta(days=60))
    _order(db_session, sor["this"], "Foodcase", 1200, sp, booked_on=date.today() - timedelta(days=90))
    _order(db_session, sor["this"], "Gone", 500, sp, status="cancelled")
    detail = client.get(f"/api/sales/editions/{sor['this'].id}").json()
    assert detail["label"] == "OBH 105"
    assert detail["booked_gbp"] == 1200 and detail["orders"] == 1 and detail["cancelled_or_moved"] == 1
    assert detail["previous"]["label"] == "OBH 101" and detail["previous_booked_gbp"] == 1000
    assert detail["uninvoiced"] == 1
    assert client.get(f"/api/sales/editions/{sor['this'].id}/export").status_code == 200


def test_overview_and_renewals(client, db_session, sor):
    sp = sor["reps"]["SP"]
    _order(db_session, sor["last"], "Foodcase", 1000, sp, booked_on=date(date.today().year - 1, 1, 5))
    _order(db_session, sor["last"], "Lapsed Ltd", 3000, sp, booked_on=date(date.today().year - 1, 1, 6))
    _order(db_session, sor["this"], "Foodcase", 1500, sp, booked_on=date.today() - timedelta(days=5))
    _order(db_session, sor["this"], "Brand New", 700, sp, booked_on=date.today() - timedelta(days=5), invoice_number="INV-1",
           invoice_value_gbp=700)
    ov = client.get("/api/sales/overview").json()
    assert ov["booked_gbp"] == 2200 and ov["last_year_total_gbp"] == 4000
    assert ov["uninvoiced_count"] == 1 and ov["renewal_candidates"] == 1
    assert ov["new_advertisers"] == 1  # "Brand New" never booked before

    ren = client.get(f"/api/sales/renewals?title_id={sor['title'].id}").json()
    assert [r["client_name"] for r in ren["items"]] == ["Lapsed Ltd"]
    assert ren["previous_advertisers"] == 2 and ren["rebooked"] == 1 and ren["retention_rate"] == 0.5


def test_commissions_are_scoped_for_sales_reps(client, db_session, sor):
    me = make_user(db_session, role="sales")
    sp, st = sor["reps"]["SP"], sor["reps"]["ST"]
    sp.user_id = me.id
    db_session.flush()
    _order(db_session, sor["this"], "Mine", 1000, sp)
    _order(db_session, sor["this"], "Theirs", 5000, st)
    mine = client.get("/api/sales/commissions", headers=identity_headers(me)).json()
    assert mine["scoped_to_me"] and [r["rep"]["code"] for r in mine["reps"]] == ["SP"]
    assert mine["reps"][0]["commission_gbp"] == 20  # 2% of £1,000
    everyone = client.get("/api/sales/commissions", headers=identity_headers(make_user(db_session, role="admin"))).json()
    assert {r["rep"]["code"] for r in everyone["reps"]} == {"SP", "ST"}


def test_order_filters(client, db_session, sor):
    _order(db_session, sor["this"], "Uninvoiced", 100, sor["reps"]["SP"])
    _order(db_session, sor["this"], "Invoiced", 100, sor["reps"]["SP"], invoice_number="INV-1", invoice_value_gbp=80)
    assert [o["client_name"] for o in client.get("/api/sales/orders?uninvoiced=true").json()["items"]] == ["Uninvoiced"]
    assert [o["client_name"] for o in client.get("/api/sales/orders?mismatched=true").json()["items"]] == ["Invoiced"]
    assert client.get("/api/sales/orders?search=invoic").json()["total"] == 2


def test_order_multi_filters_sort_facets_and_export(client, db_session, sor):
    sp, st = sor["reps"]["SP"], sor["reps"]["ST"]
    _order(db_session, sor["this"], "Beta Air", 500, sp, booked_on=date.today() - timedelta(days=5))
    _order(db_session, sor["this"], "Alpha Air", 2500, st, booked_on=date.today() - timedelta(days=50),
           invoice_number="INV-9", invoice_value_gbp=2500)
    _order(db_session, sor["this"], "Gamma Air", 900, st, status="cancelled")
    _order(db_session, sor["last"], "Delta Air", 1200, sp)
    this_year = str(sor["this"].year)

    def names(qs):
        return [o["client_name"] for o in client.get(f"/api/sales/orders?{qs}").json()["items"]]

    assert names(f"year={this_year}&sort=client&desc=false") == ["Alpha Air", "Beta Air", "Gamma Air"]
    assert names(f"year={this_year}&sort=value&desc=true")[0] == "Alpha Air"
    assert set(names("status=booked&status=cancelled&year=" + this_year)) == {"Alpha Air", "Beta Air", "Gamma Air"}
    assert set(names(f"rep_id={sp.id}")) == {"Beta Air", "Delta Air"}
    assert names("value_min=1000&value_max=2000") == ["Delta Air"]
    assert names(f"invoiced=false&year={this_year}") == ["Beta Air"]  # live, valued, no invoice
    assert names(f"booked_from={(date.today() - timedelta(days=10)).isoformat()}") == ["Beta Air"]
    assert names(f"edition_id={sor['last'].id}") == ["Delta Air"]

    f = client.get(f"/api/sales/orders/facets?year={this_year}&status=booked").json()
    # A dimension's own filter is ignored for its counts; the others apply.
    assert f["status"] == {"booked": 2, "cancelled": 1}
    assert f["year"][this_year] == 2 and f["year"][str(sor["last"].year)] == 1
    assert f["rep"][str(st.id)] == 1 and f["rep"][str(sp.id)] == 1
    assert f["invoiced"] == {"true": 1, "false": 1}

    r = client.get(f"/api/sales/orders/export?year={this_year}&sort=client&desc=false")
    assert r.status_code == 200
    import io

    import openpyxl
    rows = list(openpyxl.load_workbook(io.BytesIO(r.content)).active.values)
    assert [x[0] for x in rows[1:]] == ["Alpha Air", "Beta Air", "Gamma Air"]


# ---- Automations ----------------------------------------------------------------

class _NoCloseSession:
    """Lets a scanner (which opens and closes its own SessionLocal) run
    inside the test's rolled-back transaction."""
    def __init__(self, s):
        self._s = s

    def __getattr__(self, name):
        return getattr(self._s, name)

    def close(self):
        pass


def test_client_matching_links_exact_and_queues_likely(db_session, sor):
    exact = Company(id=uuid.uuid4(), name="Foodcase Ltd", source_db="onboard", source_act_id=str(uuid.uuid4()))
    likely = Company(id=uuid.uuid4(), name="Gategroup International", source_db="onboard", source_act_id=str(uuid.uuid4()))
    db_session.add_all([exact, likely])
    db_session.flush()
    a = _order(db_session, sor["this"], "Foodcase", 100)
    b = _order(db_session, sor["this"], "Gategroup Intl", 100)
    linked, queued = match_clients(db_session)
    db_session.refresh(a)
    assert linked == 1 and a.company_id == exact.id
    item = db_session.query(ReviewQueueItem).filter_by(kind="sor_client_match").one()
    assert queued == 1 and item.payload["client_name"] == "Gategroup Intl"

    get_kind("sor_client_match").handler(db_session, item, "link", {})
    db_session.flush()
    assert b.company_id == likely.id


def test_invoice_chase_queues_and_records_invoice(db_session, sor, monkeypatch):
    monkeypatch.setattr(sor_automations, "SessionLocal", lambda: _NoCloseSession(db_session))
    o = _order(db_session, sor["this"], "Late Payer", 900, sor["reps"]["SP"])
    _order(db_session, sor["this"], "Free", 0, sor["reps"]["SP"])
    sor_automations.scan_uninvoiced()
    items = db_session.query(ReviewQueueItem).filter_by(kind="sor_invoice_missing").all()
    assert [i.payload["order_id"] for i in items] == [str(o.id)]

    sor_automations.scan_uninvoiced()  # already pending -> not queued twice
    assert db_session.query(ReviewQueueItem).filter_by(kind="sor_invoice_missing").count() == 1

    get_kind("sor_invoice_missing").handler(db_session, items[0], "record_invoice", {"invoice_number": "INV-77", "invoice_value": "850"})
    assert o.invoice_number == "INV-77" and float(o.invoice_value_gbp) == 850


def test_renewal_scan_drafts_once_per_client(db_session, sor, monkeypatch):
    monkeypatch.setattr(sor_automations, "SessionLocal", lambda: _NoCloseSession(db_session))
    monkeypatch.setattr(sor_automations, "is_configured", lambda: False)
    anniversary_soon = date.today().replace(year=date.today().year - 1) + timedelta(days=20)
    _order(db_session, sor["last"], "Lapsed Ltd", 2000, sor["reps"]["SP"], booked_on=anniversary_soon, size="FP")
    _order(db_session, sor["last"], "Far Off", 2000, sor["reps"]["SP"], booked_on=anniversary_soon + timedelta(days=200))
    sor_automations.scan_renewals()
    sor_automations.scan_renewals()
    items = db_session.query(ReviewQueueItem).filter_by(kind="renewal_due").all()
    assert [i.payload["client_name"] for i in items] == ["Lapsed Ltd"]
    assert "a full page in OBH 101" in items[0].payload["original_text"]


def test_import_keeps_cancel_reason_order_refs_and_extra_columns(db_session, tmp_path):
    y = tmp_path / "2026"
    y.mkdir()
    header_top = HEADER_TOP[:5] + ["Seats"] + HEADER_TOP[6:]
    rows = [
        ["PUBLICATION", None, "People Awards"], [],
        ["MONTH", None, "15th September", None, None, None, None, None, None, "Exchange rate", 1.7],
        [],
        header_top,
        HEADER,
        _row("21.04.26", "Deborah Short", "Judge ticket cancelled 4/9", "KH", 0.0, inv=94368633.0),
        _row("21.04.26", "Guest", "One ticket", "KH", 0.0, inv=94368634.0),
        _row("28.01.26", "Virgin Atlantic", "Ruby sponsorship", "KH", 10000.0, inv="INV-3250",
             reason="TO BE ON NEXT QUARTER INVOICE", sp=7012.5),
    ]
    rows[-1][13] = 7012.5
    for r in rows[6:]:
        r[5] = r[5]  # size stays in column 5 (under "Size")
    _write_book(y / "TBTM Events 2026.xlsx", {"People Awards": rows})
    import_sor(db_session, tmp_path)
    o = {x.client_name: x for x in db_session.query(SalesOrder).all()}
    assert o["Deborah Short"].status == "cancelled"
    assert o["Deborah Short"].status_reason == "Judge ticket cancelled 4/9"
    assert o["Guest"].order_ref == "94368634" and o["Guest"].invoice_number is None
    assert o["Virgin Atlantic"].import_warning is None  # credit = invoiced amount is a normal sheet pattern


def test_explained_differences_are_part_invoiced_not_mismatched(client, db_session, sor):
    _order(db_session, sor["this"], "Explained", 10000, sor["reps"]["SP"], invoice_number="INV-1",
           invoice_value_gbp=7012.5, invoice_note="TO BE ON NEXT QUARTER INVOICE")
    _order(db_session, sor["this"], "Unexplained", 500, sor["reps"]["SP"], invoice_number="INV-2", invoice_value_gbp=400)
    assert [x["client_name"] for x in client.get("/api/sales/orders?mismatched=true").json()["items"]] == ["Unexplained"]
    assert [x["client_name"] for x in client.get("/api/sales/orders?part_invoiced=true").json()["items"]] == ["Explained"]
    ed = client.get(f"/api/sales/editions/{sor['this'].id}").json()
    assert ed["paid_orders"] == 2 and ed["invoiced_orders"] == 2


def test_size_fractions_and_hidden_commission_columns(db_session, tmp_path):
    from app.sales.sor_import import fmt_size

    assert fmt_size(2 / 3) == "2/3" and fmt_size(0.5) == "1/2" and fmt_size(2.0) == "2" and fmt_size("DPS") == "DPS"

    y = tmp_path / "2025"
    y.mkdir()
    rows = _sheet_rows([
        # TripStax: K.Hicks' visible column 1500; the hidden S.Thompson column holds a stale 1500 too.
        _row("29.01.25", "TripStax", 2 / 3, "KH", 2500.0, "INV-2738", sp=0.0, st=1500.0),
        # Commission column shifted a row: shows the neighbour's value.
        _row("30.01.25", "Daytona", "FP", "SP", 2222.22, "INV-1", sp=2750.0),
    ])
    rows[5][18], rows[5][19] = "K.Hicks", "S.Thompson"
    rows[6][18], rows[6][19] = 1500.0, 1500.0
    rows[6][13] = 1500.0  # invoiced £1,500 on INV-2738 (the rest on a separate invoice)
    path = y / "TBTM Print and Digital 2025.xlsx"
    _write_book(path, {"Dec Print": rows})
    wb = openpyxl.load_workbook(path)
    wb["Dec Print"].column_dimensions["T"].hidden = True  # S.Thompson
    wb.save(path)

    import_sor(db_session, tmp_path)
    reps = {r.id: r.code for r in db_session.query(SalesRep).all()}
    o = {x.client_name: x for x in db_session.query(SalesOrder).all()}
    credits = lambda order: {reps[c.rep_id]: float(c.amount_gbp) for c in db_session.query(SalesOrderCredit).filter_by(order_id=order.id)}  # noqa: E731
    assert o["TripStax"].size == "2/3"
    assert credits(o["TripStax"]) == {"KH": 1500.0}  # hidden column ignored
    assert credits(o["Daytona"]) == {"SP": 2222.22}  # misaligned sheet figure replaced by the booking value
    assert "shifted" in o["Daytona"].import_warning
