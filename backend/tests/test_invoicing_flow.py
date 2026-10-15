"""The invoicing flow agreed on the 15 Oct call: when to invoice (and Samantha's own dates), Ready to invoice cards,
the lock once a booking is checked or invoiced, unlocking, holiday cover, the copy contact, terms per publication,
and adverts found in the magazine with no booking."""
import io
import uuid
from datetime import date, timedelta

import docx
from tests.conftest import identity_headers, make_user
from tests.test_orders import w  # noqa: F401 - fixture

from app.automations import invoice_due, runtime_settings
from app.models import Notification, ReviewQueueItem, SalesDeal, SalesOrder
from app.sales.invoicing import plan_date, quarter_start


def test_when_to_invoice_dates():
    today = date(2026, 10, 15)
    feb = date(2027, 2, 10)
    assert quarter_start(feb, 7) == date(2027, 1, 1)
    assert plan_date("asap", feb, None, today, 7, 3) == date(2027, 1, 1)                     # Matt: Oct booking, Feb issue -> January
    assert plan_date("asap", date(2026, 11, 20), None, today, 7, 3) == today                # same quarter: straight away
    assert plan_date("asap", None, None, today, 7, 3) == today
    assert plan_date("on_publication", feb, None, today, 7, 3) == feb
    assert plan_date("after_run", feb, None, today, 7, 3) == date(2027, 2, 13)
    assert plan_date("on_date", feb, date(2026, 12, 1), today, 7, 3) == date(2026, 12, 1)
    assert plan_date("next_quarter", date(2027, 6, 1), None, today, 7, 3) == date(2027, 7, 1)   # no budget left this year
    assert plan_date("next_year", date(2027, 6, 1), None, today, 7, 3) == date(2027, 7, 1)
    assert plan_date("next_year", date(2026, 11, 1), None, today, 7, 3) == date(2027, 7, 1)
    assert plan_date("on_publication", None, None, today, 7, 3) is None


def _booking(w, pub: date, **kw):
    ed = w["ed"]("selling-travel", f"Issue {uuid.uuid4().hex[:4]}", pub)
    return {"client_name": "Travel Wisconsin", "contact_name": "Keely", "contact_email": "keely@example.com",
            "confirmation_address": "Cellet Travel Services Ltd", "rep_id": str(w["reps"]["SP"].id), "booked_on": date.today().isoformat(),
            "lines": [{"description": "Full page", "unit_price": 2000, "list_price": 2000, "placements": [{"edition_id": str(ed.id)}]}],
            "status": "confirmed", **kw}, ed


def test_ready_to_invoice_cards_and_accounts_dates(w, client):
    db = w["db"]
    sam = make_user(db, role="sales", name="Samantha Elmer")
    sam.can_approve_bookings = True
    sally = make_user(db, role="sales", name="Sally Parker")
    w["reps"]["SP"].user_id = sally.id
    db.commit()
    src = w["titles"]["selling-travel"].crm_source_db
    sam_h, sally_h = identity_headers(sam, access=[(src, None)]), identity_headers(sally, access=[(src, None)])

    # publishes in a later quarter: not due yet, then due once accounts pick today... via "remind me on"
    later = quarter_start(date.today(), 7) + timedelta(days=200)
    body, ed = _booking(w, later, invoice_instructions="Client sends a PO first")
    d = client.post("/api/sales/deals", json=body, headers=sally_h).json()
    row = d["schedule"][0]
    assert row["invoice_due"] == quarter_start(later, 7).isoformat() and not row["ready_to_invoice"]
    assert db.query(ReviewQueueItem).filter_by(kind="invoice_due").count() == 0
    # Samantha brings it forward to today -> card now; the salesperson is told
    r = client.post(f"/api/sales/deals/{d['id']}/invoice-date", json={"order_ids": [row["order_id"]], "due_on": date.today().isoformat()},
                    headers=sam_h)
    assert r.status_code == 200, r.text
    assert r.json()["schedule"][0]["ready_to_invoice"] and r.json()["schedule"][0]["invoice_due_set"]
    assert client.post(f"/api/sales/deals/{d['id']}/invoice-date", json={"order_ids": [row["order_id"]], "due_on": None},
                       headers=sally_h).status_code == 403
    card = db.query(ReviewQueueItem).filter_by(kind="invoice_due", status="pending").one()
    labels = {x["label"]: x["value"] for x in card.payload["details"]}
    assert labels["Invoice instructions"] == "Client sends a PO first" and "Full page" in labels["Due to invoice"]
    assert db.query(Notification).filter_by(user_id=sam.id, kind="invoice_due").count() == 1
    # "Remind me on" a later date closes it for now
    db.commit()
    soon = (date.today() + timedelta(days=10)).isoformat()
    assert client.post(f"/api/review-queue/{card.id}/actions/remind", json={"fields": {"remind_on": soon}}, headers=sam_h).status_code == 200
    db.expire_all()
    assert db.get(SalesOrder, uuid.UUID(row["order_id"])).invoice_due_on.isoformat() == soon

    # due now (same quarter, "as soon as allowed") -> a card straight away; "Invoiced" records the number
    body2, _ = _booking(w, date.today() + timedelta(days=1))
    d2 = client.post("/api/sales/deals", json=body2, headers=sally_h).json()
    card2 = db.query(ReviewQueueItem).filter_by(kind="invoice_due", status="pending").one()
    assert card2.payload["deal_id"] == d2["id"]
    db.commit()
    assert client.post(f"/api/review-queue/{card2.id}/actions/invoiced", json={"fields": {"invoice_number": "INV-9001"}}, headers=sally_h).status_code != 200
    assert client.post(f"/api/review-queue/{card2.id}/actions/invoiced", json={"fields": {"invoice_number": "INV-9001"}}, headers=sam_h).status_code == 200
    db.expire_all()
    assert db.query(SalesOrder).filter_by(deal_id=uuid.UUID(d2["id"])).one().invoice_number == "INV-9001"
    # invoiced = locked for the salesperson, even though nobody pressed "Looks right"
    got = client.get(f"/api/sales/deals/{d2['id']}", headers=sally_h).json()
    assert got["locked"]
    assert client.put(f"/api/sales/deals/{d2['id']}", json={**body2, "notes": "x"}, headers=sally_h).status_code == 409

    # a pencilled booking gets neither a check nor an invoice reminder
    body3, _ = _booking(w, date.today(), status="pencilled")
    client.post("/api/sales/deals", json=body3, headers=sally_h)
    assert db.query(ReviewQueueItem).filter_by(kind="invoice_due", status="pending").count() == 0

    # a set date needs the date, and one in an earlier quarter gets a gentle warning
    body4, _ = _booking(w, later, invoice_plan="on_date")
    assert client.post("/api/sales/deals", json=body4, headers=sally_h).status_code == 422
    d4 = client.post("/api/sales/deals", json={**body4, "invoice_on": date.today().isoformat()}, headers=sally_h).json()
    assert any("quarter before" in x for x in d4["warnings"])


def test_unlock_and_holiday_cover(w, client):
    db = w["db"]
    sam = make_user(db, role="sales", name="Samantha Elmer")
    sam.can_approve_bookings = True
    sally = make_user(db, role="sales", name="Sally Parker")
    matt = make_user(db, role="admin", name="Matt Bonner")
    w["reps"]["SP"].user_id = sally.id
    db.commit()
    src = w["titles"]["selling-travel"].crm_source_db
    sam_h, sally_h, matt_h = (identity_headers(u, access=[(src, None)]) for u in (sam, sally, matt))
    body, _ = _booking(w, date.today() + timedelta(days=200))
    d = client.post("/api/sales/deals", json=body, headers=sally_h).json()
    check = db.query(ReviewQueueItem).filter_by(kind="order_check", status="pending").one()
    db.commit()
    # Matt can't approve while Samantha's in; can once "Accounts are away" is on
    assert client.post(f"/api/review-queue/{check.id}/actions/approve", json={}, headers=matt_h).status_code == 404
    runtime_settings.set_override(db, "accounts_away", True)
    db.commit()
    assert client.post(f"/api/review-queue/{check.id}/actions/approve", json={}, headers=matt_h).status_code == 200
    db.expire_all()
    got = client.get(f"/api/sales/deals/{d['id']}", headers=sally_h).json()
    assert got["locked"]
    assert client.post(f"/api/sales/deals/{d['id']}/unlock", json={}, headers=sally_h).status_code == 403
    # unlocked for Sally: she changes it, it locks again and goes back for checking
    r = client.post(f"/api/sales/deals/{d['id']}/unlock", json={"note": "Go ahead and fix the address"}, headers=sam_h)
    assert r.status_code == 200 and r.json()["unlocked"] and not r.json()["locked"]
    assert db.query(Notification).filter_by(user_id=sally.id, kind="booking_unlocked").count() == 1
    r = client.put(f"/api/sales/deals/{d['id']}", json={**body, "confirmation_address": "Cellet Travel, Bloxham"}, headers=sally_h)
    assert r.status_code == 200, r.text
    assert not r.json()["unlocked"] and r.json()["check"]["state"] == "waiting"
    assert db.query(ReviewQueueItem).filter_by(kind="order_check", status="pending").count() == 1


def test_copy_contact_terms_per_publication_and_unbooked_ads(w, client):
    db = w["db"]
    h = w["h"]
    st_id = str(w["titles"]["selling-travel"].id)
    cur = client.get("/api/sales/deals/settings", headers=h).json()
    client.put("/api/sales/deals/settings", json={**cur, "terms": "General BMI terms: bmipublishing.co.uk/terms",
                                                  "terms_by_title": {st_id: "Selling Travel terms: sellingtravel.co.uk/terms"}}, headers=h)
    body, ed = _booking(w, date.today() + timedelta(days=30), copy_contact_name="Ana at Creative Co", copy_contact_email="ana@creative.example",
                        invoice_currency="USD")
    d = client.post("/api/sales/deals", json=body, headers=h).json()
    assert d["invoice_currency"] == "USD" and d["copy_contact_email"] == "ana@creative.example"
    doc = client.get(f"/api/sales/deals/{d['id']}/document", headers=h).json()
    assert doc["terms"].startswith("Selling Travel terms") and "Ana at Creative Co" in doc["copy_contact"]
    word = docx.Document(io.BytesIO(client.get(f"/api/sales/deals/{d['id']}/document.docx", headers=h).content))
    text = "\n".join(p.text for p in word.paragraphs) + "\n".join(c.text for t in word.tables for r in t.rows for c in r.cells)
    assert "Copy and artwork from: Ana at Creative Co" in text and "Selling Travel terms" in text

    # an advert found in the issue with no booking: a filler, or ask the salesperson to make the booking
    r = client.post("/api/sales/deals/unbooked-ad", json={"edition_id": str(ed.id), "client_name": "Visit Malta", "page": "34",
                                                          "outcome": "filler"}, headers=h)
    assert r.status_code == 201
    filler = db.query(SalesOrder).filter_by(edition_id=ed.id, filler=True).one()
    assert float(filler.value_gbp) == 0 and "page 34" in filler.notes
    sally = make_user(db, role="sales", name="Sally Parker")
    w["reps"]["SP"].user_id = sally.id
    db.commit()
    r = client.post("/api/sales/deals/unbooked-ad", json={"edition_id": str(ed.id), "client_name": "Visit Malta", "page": "35",
                                                          "rep_id": str(w["reps"]["SP"].id), "outcome": "needs_booking"}, headers=h)
    assert r.status_code == 201
    note = db.query(Notification).filter_by(user_id=sally.id, kind="unbooked_ad").one()
    assert f"edition={ed.id}" in note.link and "page 35" in note.body
