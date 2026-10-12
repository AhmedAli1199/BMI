"""Rate card revamp: brands, ways prices are quoted, offers, publishers' rights, history + put back, next year, media-pack load."""
import io
import uuid
from datetime import date

import openpyxl
import pytest

from app.models import RateOffer, SalesRate, SalesRep, SalesTitle
from app.sales.reference import ensure_reference_data
from app.sales.renewals import current_price
from tests.conftest import identity_headers, make_user

Y = 2026


@pytest.fixture()
def world(db_session):
    ensure_reference_data(db_session)
    admin = make_user(db_session, role="admin")
    kirsty = make_user(db_session, role="sales", name="Kirsty")
    rep = db_session.query(SalesRep).filter_by(code="KH").one()
    rep.user_id = kirsty.id
    db_session.flush()
    return {"admin": identity_headers(admin), "kirsty": identity_headers(kirsty, access=[("sellingtravel", None)]),
            "nobody": identity_headers(make_user(db_session, role="sales"))}


def test_overview_and_media_pack_load(client, db_session, world):
    ov = client.get(f"/api/rate-card?year={Y}", headers=world["admin"]).json()
    assert [b["key"] for b in ov["brands"]] == ["obh", "tbtm", "stm"]
    obh = ov["brands"][0]
    assert obh["products"] == 0 and obh["seed_available"] > 20 and obh["can_edit"]
    r = client.post(f"/api/rate-card/brands/obh/load-media-pack?year={Y}", headers=world["admin"]).json()
    assert r["added"] == obh["seed_available"]
    page = client.get(f"/api/rate-card/brands/obh?year={Y}", headers=world["admin"]).json()
    print_items = next(s for s in page["sections"] if s["key"] == "print")["items"]
    full = next(i for i in print_items if i["product"] == "Full page")
    assert full["price_gbp"] == 2990 and full["needs_check"] and "FP" in full["aliases"] and full["price_label"] == "£2,990"
    web = next(s for s in page["sections"] if s["key"] == "website")["items"]
    assert web[0]["title_name"] == "OBH Web" and web[0]["price_label"].endswith("per month")
    assert next(i for i in print_items if i["product"] == "Double page")["price_label"] == "Price on request"
    assert page["brand"]["seed_available"] == 0
    assert client.post(f"/api/rate-card/brands/obh/load-media-pack?year={Y}", headers=world["admin"]).json()["added"] == 0
    # offers came too
    assert any(o["kind"] == "series" for o in page["offers"])


def test_publisher_edits_only_own_brand(client, db_session, world):
    from app.automations import runtime_settings

    body = {"brand": "tbtm", "year": Y, "section": "print", "product": "Full page", "price_gbp": 2850}
    # by default only admins and data managers change prices (Matt, 8 Oct)
    assert client.post("/api/rate-card/items", json=body, headers=world["kirsty"]).status_code == 403
    assert not next(b for b in client.get(f"/api/rate-card?year={Y}", headers=world["kirsty"]).json()["brands"] if b["key"] == "tbtm")["can_edit"]
    runtime_settings.set_override(db_session, "rate_card_publishers_can_edit", True)
    db_session.flush()
    assert client.post("/api/rate-card/items", json=body, headers=world["kirsty"]).status_code == 201
    assert client.post("/api/rate-card/items", json={**body, "brand": "obh"}, headers=world["kirsty"]).status_code == 403
    assert client.post("/api/rate-card/items", json={**body, "product": "Half page"}, headers=world["nobody"]).status_code == 403
    ov = client.get(f"/api/rate-card?year={Y}", headers=world["kirsty"]).json()
    assert {b["key"]: b["can_edit"] for b in ov["brands"]} == {"obh": False, "tbtm": True, "stm": False}


def test_add_validation_and_defaults(client, world):
    h = world["admin"]
    base = {"brand": "obh", "year": Y, "section": "website", "product": "Leaderboard", "price_gbp": 1000, "unit": "month"}
    r = client.post("/api/rate-card/items", json=base, headers=h).json()
    assert r["title_name"] == "OBH Web" and r["price_label"] == "£1,000 per month"
    assert client.post("/api/rate-card/items", json=base, headers=h).status_code == 409
    assert client.post("/api/rate-card/items", json={**base, "product": "X", "price_gbp": None}, headers=h).status_code == 422
    poa = client.post("/api/rate-card/items", json={**base, "product": "Takeover", "price_type": "poa", "price_gbp": 5}, headers=h).json()
    assert poa["price_gbp"] is None and poa["price_label"] == "Price on request"
    frm = client.post("/api/rate-card/items", json={**base, "product": "Hub", "price_type": "from", "price_gbp": 9500, "unit": "each"}, headers=h).json()
    assert frm["price_label"] == "From £9,500"
    assert client.post("/api/rate-card/items", json={**base, "product": "Y", "section": "nope"}, headers=h).status_code == 422


def test_edit_history_put_back_and_archive(client, world):
    h = world["admin"]
    item = client.post("/api/rate-card/items", json={"brand": "stm", "year": Y, "section": "print", "product": "Full page", "price_gbp": 3750}, headers=h).json()
    client.patch(f"/api/rate-card/items/{item['id']}", json={"price_gbp": 3900}, headers=h)
    hist = client.get(f"/api/rate-card/history?brand=stm&year={Y}", headers=h).json()
    change = next(c for c in hist if c["field"] == "price")
    assert change["old_value"].startswith("3750") and change["new_value"].startswith("3900") and change["can_put_back"]
    back = client.post(f"/api/rate-card/history/{change['id']}/put-back", headers=h).json()
    assert back["price_gbp"] == 3750
    # changed again since -> can't put back the older change
    client.patch(f"/api/rate-card/items/{item['id']}", json={"price_gbp": 4000}, headers=h)
    assert client.post(f"/api/rate-card/history/{change['id']}/put-back", headers=h).status_code == 409
    a = client.post(f"/api/rate-card/items/{item['id']}/archive", headers=h).json()
    assert a["archived"]
    page = client.get(f"/api/rate-card/brands/stm?year={Y}", headers=h).json()
    assert page["archived"][0]["id"] == item["id"] and not next(s for s in page["sections"] if s["key"] == "print")["items"]
    assert not client.post(f"/api/rate-card/items/{item['id']}/archive?restore=true", headers=h).json()["archived"]


def test_offers_validation(client, world):
    h = world["admin"]
    ok = {"brand": "tbtm", "year": Y, "kind": "volume", "label": "Book 2 save 10%", "rules": {"tiers": [{"qty": 2, "discount_pct": 10}]}, "section": "print"}
    assert client.post("/api/rate-card/offers", json=ok, headers=h).status_code == 201
    assert client.post("/api/rate-card/offers", json={**ok, "rules": {"tiers": [{"qty": 1, "discount_pct": 10}]}}, headers=h).status_code == 422
    assert client.post("/api/rate-card/offers", json={**ok, "kind": "early_bird", "rules": {}}, headers=h).status_code == 422
    assert client.post("/api/rate-card/offers", json={**ok, "kind": "early_bird", "rules": {}, "valid_until": "2027-02-01"}, headers=h).status_code == 201


def test_next_year_copy_with_rise(client, db_session, world):
    h = world["admin"]
    client.post(f"/api/rate-card/brands/tbtm/load-media-pack?year={Y}", headers=h)
    pv = client.post("/api/rate-card/next-year/preview", json={"brand": "tbtm", "from_year": Y, "raise_pct": 5, "round_to": 10}, headers=h).json()
    full = next(r for r in pv["rows"] if r["product"] == "Full page")
    assert full["old"] == 2850 and full["new"] == 2990  # 2,992.50 -> nearest £10
    poa = next(r for r in pv["rows"] if r["price_type"] == "poa")
    assert poa["new"] is None
    res = client.post("/api/rate-card/next-year", json={"brand": "tbtm", "from_year": Y, "raise_pct": 5, "round_to": 10,
                                                        "overrides": {full["id"]: 2995}}, headers=h).json()
    assert res["copied"] == len(pv["rows"]) and res["year"] == Y + 1
    nxt = client.get(f"/api/rate-card/brands/tbtm?year={Y + 1}", headers=h).json()
    fp = next(i for i in next(s for s in nxt["sections"] if s["key"] == "print")["items"] if i["product"] == "Full page")
    assert fp["price_gbp"] == 2995 and not fp["needs_check"]
    eb = next(o for o in nxt["offers"] if o["kind"] == "early_bird")
    assert eb["valid_until"] == "2028-02-01" and eb["needs_check"]
    assert client.post("/api/rate-card/next-year/preview", json={"brand": "tbtm", "from_year": Y}, headers=h).json()["rows"] == []


def test_renewal_price_matches_aliases_and_skips_archived(client, db_session, world):
    client.post(f"/api/rate-card/brands/obh/load-media-pack?year={Y}", headers=world["admin"])
    obh = db_session.query(SalesTitle).filter_by(slug="obh").one()
    assert float(current_price(db_session, obh.id, Y, "FP").price_gbp) == 2990
    assert current_price(db_session, obh.id, Y, "DPS") is None  # price on request
    fp = db_session.query(SalesRate).filter_by(title_id=obh.id, year=Y, product="Full page").one()
    fp.archived = True
    db_session.flush()
    assert current_price(db_session, obh.id, Y, "FP") is None


def test_confirm_all_and_export(client, world):
    h = world["admin"]
    client.post(f"/api/rate-card/brands/stm/load-media-pack?year={Y}", headers=h)
    assert client.post("/api/rate-card/confirm-all", json={"brand": "stm", "year": Y}, headers=h).json()["confirmed"] > 20
    assert client.get(f"/api/rate-card/brands/stm?year={Y}", headers=h).json()["brand"]["needs_check"] == 0
    x = client.get(f"/api/rate-card/export.xlsx?year={Y}&brand=stm", headers=h)
    ws = openpyxl.load_workbook(io.BytesIO(x.content)).active
    assert "Selling Travel" in ws["A1"].value and any(c.value == "Full page" for c in ws["A"])


def test_sizes_pages_and_own_sections(client, world):
    """Sizes in mm or px with bleed, page counts, and a brand's own section (cover wraps, belly bands)."""
    h = world["admin"]
    sec = client.post("/api/rate-card/brands/obh/sections", json={"label": "Cover wraps & belly bands", "hint": "Special formats"}, headers=h)
    assert sec.status_code == 201 and sec.json()["key"].startswith("x-") and sec.json()["custom"]
    key = sec.json()["key"]
    assert client.post("/api/rate-card/brands/obh/sections", json={"label": "cover wraps & belly bands"}, headers=h).status_code == 409
    item = client.post("/api/rate-card/items", json={"brand": "obh", "year": Y, "section": key, "product": "Belly band", "price_gbp": 4500,
                                                     "width": 420, "height": 80, "bleed_mm": 3}, headers=h)
    assert item.status_code == 201, item.text
    it = item.json()
    assert it["size_unit"] == "mm" and it["size_label"] == "420 x 80 mm, 3 mm bleed" and it["title_name"] == "OnBoard Hospitality (OBH)"
    adv = client.post("/api/rate-card/items", json={"brand": "obh", "year": Y, "section": "sponsored", "product": "Advertorial",
                                                    "price_gbp": 3990, "pages": 2}, headers=h).json()
    assert adv["size_label"] == "2 pages"
    web = client.post("/api/rate-card/items", json={"brand": "obh", "year": Y, "section": "website", "product": "MPU", "price_gbp": 500,
                                                    "unit": "month", "width": 300, "height": 250, "size_unit": "px"}, headers=h).json()
    assert web["size_label"] == "300 x 250 px"
    bad = {"brand": "obh", "year": Y, "section": "print", "product": "Odd", "price_gbp": 1}
    assert client.post("/api/rate-card/items", json={**bad, "width": 100}, headers=h).status_code == 422  # height missing
    assert client.post("/api/rate-card/items", json={**bad, "width": 1, "height": 1, "size_unit": "px", "bleed_mm": 3}, headers=h).status_code == 422
    assert client.post("/api/rate-card/items", json={**bad, "section": "x-nope"}, headers=h).status_code == 422
    # clearing the size
    cleared = client.patch(f"/api/rate-card/items/{it['id']}", json={"width": None, "height": None}, headers=h).json()
    assert cleared["size_label"] is None and cleared["size_unit"] is None
    page = client.get(f"/api/rate-card/brands/obh?year={Y}", headers=h).json()
    own = next(s for s in page["sections"] if s["key"] == key)
    assert own["label"] == "Cover wraps & belly bands" and [i["product"] for i in own["items"]] == ["Belly band"]
    assert key in page["brand"]["sections"]
    # a section with prices can't be removed; rename works
    assert client.delete(f"/api/rate-card/sections/{own['id']}", headers=h).status_code == 409
    assert client.patch(f"/api/rate-card/sections/{own['id']}", json={"label": "Special formats"}, headers=h).json()["label"] == "Special formats"
    assert client.delete(f"/api/rate-card/sections/{own['id']}", headers=world["kirsty"]).status_code == 403
    # next year keeps sizes and pages
    client.post("/api/rate-card/next-year", json={"brand": "obh", "from_year": Y, "raise_pct": 0}, headers=h)
    nxt = client.get(f"/api/rate-card/brands/obh?year={Y + 1}", headers=h).json()
    assert next(i for s in nxt["sections"] for i in s["items"] if i["product"] == "Advertorial")["pages"] == 2
    assert next(i for s in nxt["sections"] for i in s["items"] if i["product"] == "MPU")["size_label"] == "300 x 250 px"
    rates = client.get(f"/api/sales/rates?year={Y}", headers=h).json()
    assert next(r for r in rates if r["product"] == "Advertorial")["pages"] == 2
