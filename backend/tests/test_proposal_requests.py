"""Proposal builder blocks, rich text, email-or-Word output and the examples library; and proposal requests spotted
in a salesperson's inbox with a reply drafted."""
import base64
import io
import uuid
from datetime import date, datetime, timedelta, timezone

import docx
import pytest

from app.automations import proposal_requests as pr
from app.automations.mail_parsing import ParsedMessage
from app.models import Company, Contact, Note, Proposal, ReviewQueueItem, SalesEdition, SalesRate, SalesRep, SalesTitle
from app.models.messaging import MailAccount
from app.proposals.examples import clean_url
from app.sales.reference import ensure_reference_data
from tests.conftest import identity_headers, make_user
from tests.test_sales_orders import _NoCloseSession

YEAR = date.today().year
PNG = base64.b64encode(bytes.fromhex("89504e470d0a1a0a0000000d4948445200000001000000010806000000"
                                     "1f15c4890000000d49444154789c63f8ffffff7f0009fb03fd2a86e38a0000000049454e44ae426082")).decode()


@pytest.fixture()
def tbtm(db_session):
    ensure_reference_data(db_session)
    title = db_session.query(SalesTitle).filter_by(slug="tbtm").one()
    db_session.add_all([
        SalesRate(id=uuid.uuid4(), title_id=title.id, year=YEAR, section="print", product="Full page", price_gbp=2850),
        SalesRate(id=uuid.uuid4(), title_id=title.id, year=YEAR, section="sponsored", product="Double page sponsored content", price_gbp=4750),
        SalesRate(id=uuid.uuid4(), title_id=title.id, year=YEAR, section="newsletter", product="Newsletter banner", price_gbp=1250),
    ])
    db_session.add(SalesEdition(id=uuid.uuid4(), title_id=title.id, year=YEAR + 1, name="Spring", kind="issue",
                                edition_date=date.today() + timedelta(days=60), ad_deadline=date.today() + timedelta(days=30), theme="TMCs and tech"))
    kirsty = make_user(db_session, role="sales", email="kirsty.hicks@bmipublishing.co.uk", name="Kirsty Hicks")
    rep = db_session.query(SalesRep).filter_by(code="KH").first()
    if rep:
        rep.user_id = kirsty.id
    company = Company(id=uuid.uuid4(), name="Radisson Hotels", source_db="prospects", source_act_id=uuid.uuid4().hex[:12])
    db_session.add(company)
    db_session.flush()
    return {"title": title, "user": kirsty, "h": identity_headers(kirsty, access=[(title.crm_source_db or "prospects", None)]), "company": company}


def test_blocks_rich_text_and_both_outputs(client, db_session, tbtm, monkeypatch):
    rate = db_session.query(SalesRate).filter_by(title_id=tbtm["title"].id, product="Full page").one()
    p = client.post("/api/proposals", json={"company_id": str(tbtm["company"].id), "title_id": str(tbtm["title"].id), "use_ai": False,
                                            "lines": [{"product": "Full page", "rate_id": str(rate.id), "source": "rate_card"}]},
                    headers=tbtm["h"]).json()
    ex = client.get("/api/proposal-examples?brand=tbtm", headers=tbtm["h"]).json()
    assert len(ex["items"]) > 30 and "urldefense" not in str(ex["items"])
    jetblue = next(i for i in ex["items"] if "JetBlue" in i["label"])
    sections = [
        {"kind": "intro", "heading": "Hello", "html": f"<p>Hi <strong>Chueca</strong>,</p><table><tr><td>A</td><td>B</td></tr></table>"
                                                      f"<p><img src=\"data:image/png;base64,{PNG}\"></p><script>alert(1)</script>"},
        {"kind": "block_package", "heading": "Your package", "block": {"kind": "package", "config": {}}},
        {"kind": "block_prices", "heading": "Prices", "block": {"kind": "prices", "config": {}}},
        {"kind": "block_examples", "heading": "Examples", "block": {"kind": "examples", "config": {"ids": [jetblue["id"]]}}},
    ]
    r = client.patch(f"/api/proposals/{p['id']}", json={"sections": sections, "output": "email"}, headers=tbtm["h"])
    assert r.status_code == 200, r.text
    saved = r.json()
    assert saved["output"] == "email" and "<script>" not in saved["sections"][0]["html"]
    prev = client.get(f"/api/proposals/{p['id']}/preview", headers=tbtm["h"]).json()
    assert "£2,850" in prev["html"] and "JetBlue" in prev["html"] and "Double page sponsored content" in prev["html"]

    # the Word file has the same content, tables and the picture included
    w = client.get(f"/api/proposals/{p['id']}/download", headers=tbtm["h"])
    d = docx.Document(io.BytesIO(w.content))
    text = "\n".join(par.text for par in d.paragraphs)
    cells = " ".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert "Chueca" in text and "£2,850" in cells and "JetBlue" in text and len(d.tables) >= 3 and d.inline_shapes

    # sent as the email itself: pictures go inline, the Word file can come too
    from app.services import outlook
    sent = {}
    monkeypatch.setattr(outlook, "send_mail", lambda db, acct, **kw: sent.update(kw))
    db_session.add(MailAccount(id=uuid.uuid4(), user_id=tbtm["user"].id, email="kirsty.hicks@bmipublishing.co.uk", refresh_token_enc="x"))
    db_session.flush()
    r = client.post(f"/api/proposals/{p['id']}/send", json={"to": ["buyer@radisson.test"], "subject": "Your proposal", "body": "Hi Chueca,",
                                                            "as_body": True, "attach_word": True}, headers=tbtm["h"])
    assert r.status_code == 200, r.text
    assert "cid:" in sent["html_body"] and "£2,850" in sent["html_body"]
    kinds = [a[1] for a in sent["attachments"]]
    assert "image/png" in kinds and any("wordprocessingml" in k for k in kinds)


def test_clean_url_unwraps_security_scanner_links():
    u = "https://urldefense.proofpoint.com/v2/url?u=https-3A__issuu.com_bmipublishingltd_docs_tbtm-2D99_38&d=Dw"
    assert clean_url(u) == "https://issuu.com/bmipublishingltd/docs/tbtm-99/38"


def _msg(**kw) -> ParsedMessage:
    base = dict(message_id=f"m-{uuid.uuid4().hex[:6]}", internet_message_id=None, from_address="marta@agency.test", from_name="Marta Ruiz",
                subject="Rates for our hotel client", received_at=datetime.now(timezone.utc), conversation_id=f"c-{uuid.uuid4().hex[:6]}",
                body_text="Hi Kirsty, could you send your rates and media pack? We'd like sponsored content and a display advert "
                          "for our global hospitality client in the spring. Budget around £10k.")
    base.update(kw)
    return ParsedMessage(**base)


def test_the_free_checks():
    box = "kirsty.hicks@bmipublishing.co.uk"
    domains = {"bmipublishing.co.uk"}
    assert pr.candidate(_msg(), box, domains)
    assert not pr.candidate(_msg(from_address=box), box, domains)                                         # her own email
    assert not pr.candidate(_msg(body_text="Lunch on Friday?", subject="Lunch"), box, domains)            # nothing about prices
    assert not pr.candidate(_msg(from_address="bev@bmipublishing.co.uk"), box, domains)                    # internal chat
    passed_on = _msg(from_address="bev@bmipublishing.co.uk", body_text="Passing this on.\n---------- Forwarded message ---------\n"
                                                                        "From: Leila <leila@pr.test>\nCould we feature our client? What are the costs?")
    assert pr.candidate(passed_on, box, domains)                                                           # a colleague passing one on


def test_request_is_found_drafted_and_answered(client, db_session, tbtm, monkeypatch):
    from app.automations import llm, runtime_settings
    from app.services import outlook

    box = "kirsty.hicks@bmipublishing.co.uk"
    real = runtime_settings.get_csv
    monkeypatch.setattr(runtime_settings, "get_csv", lambda db, k: [box] if k == "proposal_request_mailboxes" else real(db, k))
    m = _msg()
    monkeypatch.setattr("app.graph_client.list_messages_since", lambda mailbox, since: [{"id": m.message_id, "_m": m}])
    monkeypatch.setattr(pr, "parse_message", lambda raw: raw["_m"])
    monkeypatch.setattr(pr, "SessionLocal", lambda: _NoCloseSession(db_session))
    monkeypatch.setattr(llm, "is_configured", lambda: True)
    monkeypatch.setattr(llm, "extract_json", lambda *a, **k: {
        "is_request": True, "confidence": 0.9, "request_type": "agency_brief", "magazine": "The Business Travel Magazine",
        "wants": "sponsored content and a display advert", "budget": "around £10k", "timing": "spring", "agency": True,
        "client_company": "Global Hotels", "contact_name": "Marta Ruiz", "contact_email": "marta@agency.test", "sector": "hotel",
        "summary": "An agency wants rates for a hotel client."})
    monkeypatch.setattr(llm, "draft_text", lambda *a, **k: "Thanks so much for thinking of us for 3 campaigns.")   # has a figure: dropped
    pr.scan()
    pr.scan()   # the same email again isn't queued twice
    item = db_session.query(ReviewQueueItem).filter_by(kind="proposal_request").one()
    p = item.payload
    assert p["owner_user_id"] == str(tbtm["user"].id) and p["facets"]["type"] == "Agency brief" and p["facets"]["client"] == "New prospect"
    draft = p["original_text"]
    assert "Hi Marta," in draft and "£4,750" in draft and "£2,850" in draft and "3 campaigns" not in draft and "brief on behalf of Global Hotels" in draft
    assert "Spring" in draft and "Kirsty" in draft
    db_session.commit()

    # only Kirsty (and managers) see it
    other = identity_headers(make_user(db_session, role="sales"), access=[(tbtm["title"].crm_source_db or "prospects", None)])
    assert client.get("/api/review-queue?kind=proposal_request", headers=other).json()["total"] == 0
    assert client.get("/api/review-queue?kind=proposal_request", headers=tbtm["h"]).json()["total"] == 1

    # send: a reply in the same thread from her Outlook; the new client goes into the CRM with a note
    db_session.add(MailAccount(id=uuid.uuid4(), user_id=tbtm["user"].id, email=box, refresh_token_enc="x"))
    db_session.flush()
    got = {}
    monkeypatch.setattr(outlook, "reply_to", lambda db, acct, mid, **kw: got.update(mid=mid, **kw))
    r = client.post(f"/api/review-queue/{item.id}/actions/send", json={"note": draft}, headers=tbtm["h"])
    assert r.status_code == 200, r.text
    assert got["mid"] == m.message_id and got["to"] == ["marta@agency.test"] and "£4,750" in got["html_body"]
    contact = db_session.query(Contact).filter_by(full_name="Marta Ruiz").one()
    assert db_session.query(Company).filter_by(name="Global Hotels").one().id == contact.company_id
    assert db_session.query(Note).filter_by(entity_id=contact.id, note_type="Proposal request").count() == 1


def test_request_becomes_a_full_proposal(client, db_session, tbtm, monkeypatch):
    m = _msg()
    info = {"is_request": True, "request_type": "general", "magazine": None, "wants": "sponsored content", "agency": False,
            "client_company": "SkyJet", "contact_name": "Ana", "contact_email": "ana@skyjet.test", "sector": "airline"}
    from app.automations import llm
    monkeypatch.setattr(llm, "is_configured", lambda: False)
    item = pr.queue(db_session, "kirsty.hicks@bmipublishing.co.uk", tbtm["user"], m, info)   # no magazine named: Kirsty's brand
    db_session.commit()
    assert item.payload["title_id"] == str(tbtm["title"].id)
    r = client.post(f"/api/review-queue/{item.id}/actions/proposal", json={}, headers=tbtm["h"])
    assert r.status_code == 200, r.text
    db_session.refresh(item)
    prop = db_session.get(Proposal, uuid.UUID(item.payload["proposal_id"]))
    assert prop.output == "email" and prop.request_item_id == item.id
    assert any(s.get("block", {}).get("kind") == "examples" for s in prop.sections)
    assert prop.lines and prop.lines[0]["product"] == "Double page sponsored content"
    assert item.payload["details"][0]["href"] == f"/sales/proposals/{prop.id}"
