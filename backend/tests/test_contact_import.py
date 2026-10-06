"""Act! feedback B4: import contacts from a spreadsheet - column matching, review, import, undo."""
import io
import uuid

import openpyxl
import pytest

from app.contacts import import_engine as E
from app.contacts import import_mapper as M
from app.contacts import import_parse as P
from app.models import Company, Contact, Email, Group, GroupMembership, Note, Publication
from app.models.contact_channel import Address, Phone
from tests.conftest import identity_headers, make_user

# ---- the matcher -----------------------------------------------------------------------

def fields(headers, rows=None, custom=None):
    rows = rows or [["x"] * len(headers)] * 3
    return {r.header: r.field for r in M.match_columns(headers, rows, custom)}


@pytest.mark.parametrize("header", ["First Name", "first name", "FIRST NAME", "firstname", "First_Name", "first-name", "FirstName", "Forename",
                                    "Given name", "Given Names", "fname", "F Name", "Christian name", "Contact First Name", "First", "Frist Name", "firstnaem"])
def test_first_name_variants(header):
    assert fields([header, "Surname"])[header] == "first_name"


@pytest.mark.parametrize("header", ["Last Name", "LASTNAME", "Surname", "surnmae", "Family name", "Last_Name", "lname", "Second name"])
def test_last_name_variants(header):
    assert fields(["First Name", header])[header] == "last_name"


@pytest.mark.parametrize("header,target", [
    ("E-mail", "email"), ("Email Address", "email"), ("e_mail", "email"), ("EMAIL", "email"), ("Work Email", "email"), ("emial", "email"),
    ("Email 2", "email_2"), ("E-mail 2 Address", "email_2"), ("Secondary Email", "email_2"), ("Email3", "email_3"),
    ("Phone", "phone"), ("Tel", "phone"), ("Telephone Number", "phone"), ("Business Phone", "phone"), ("Work Tel", "phone"), ("Direct Dial", "phone"),
    ("Mobile", "mobile"), ("Mobile Phone", "mobile"), ("Cell", "mobile"), ("Mob No", "mobile"), ("Mobile phone number", "mobile"),
    ("Home Phone", "home_phone"), ("Fax", "fax"), ("Phone 2", "other_phone"),
    ("Company", "company"), ("Company Name", "company"), ("Organisation", "company"), ("Organization", "company"), ("Employer", "company"),
    ("Job Title", "job_title"), ("Position", "job_title"), ("Role", "job_title"), ("Designation", "job_title"),
    ("Postcode", "postcode"), ("Post Code", "postcode"), ("Zip", "postcode"), ("Zip Code", "postcode"), ("Postal Code", "postcode"), ("Business Postal Code", "postcode"),
    ("City", "city"), ("Town", "city"), ("Business City", "city"), ("County", "state"), ("State/Province", "state"), ("Region", "state"),
    ("Country", "country"), ("Country/Region", "country"),
    ("Address", "address_line1"), ("Address 1", "address_line1"), ("Address Line 1", "address_line1"), ("Street", "address_line1"), ("Business Street", "address_line1"),
    ("Address 2", "address_line2"), ("Address Line 3", "address_line3"),
    ("Date of Birth", "birthdate"), ("DOB", "birthdate"), ("Notes", "notes"), ("Comments", "notes"), ("Remarks", "notes"),
    ("Department", "department"), ("Salutation", "salutation"), ("Do Not Email", "unsubscribed"), ("Opt out", "unsubscribed"),
])
def test_header_variants(header, target):
    got = fields(["First Name", "Surname", header])[header]
    assert got == target, f"{header!r} -> {got}"


def test_full_name_headers():
    for h in ("Name", "Full Name", "Contact", "Contact Name"):
        assert fields([h, "Email"])[h] == "full_name"


def test_name_next_to_surname_is_first_name():
    assert fields(["Name", "Surname"], [["Ann", "Lee"]] * 3)["Name"] == "first_name"


def test_values_identify_a_poorly_named_column():
    rows = [["ann@x.com", "SW1A 1AA", "Mrs", "UK", "07700 900123"], ["bob@y.com", "M1 1AA", "Mr", "England", "07700 900124"], ["c@z.com", "EC1A 1BB", "Dr", "Wales", "07700 900125"]]
    got = fields(["Column 1", "Col B", "Salutation or title", "Place", "Contact no"], rows)
    assert got["Column 1"] == "email" and got["Col B"] == "postcode" and got["Place"] == "country"
    assert got["Contact no"] in ("mobile", "phone")


def test_title_means_prefix_or_job_title_depending_on_values():
    assert fields(["Title"], [["Mr"], ["Mrs"], ["Dr"]])["Title"] == "name_prefix"
    assert fields(["Title"], [["Sales Director"], ["Buyer"], ["Head of Travel"]])["Title"] == "job_title"


def test_wrong_heading_is_not_trusted():
    # a column called "Phone" that holds email addresses
    r = M.match_columns(["Phone"], [["a@x.com"], ["b@y.com"], ["c@z.com"]])[0]
    assert r.field == "email"


def test_ids_are_skipped_and_unknowns_left_alone():
    got = fields(["ID", "Favourite colour", "Customer ID"], [["1", "red", "C1"], ["2", "blue", "C2"]])
    assert got["ID"] == "skip" and got["Customer ID"] == "skip" and got["Favourite colour"] is None
    res = {r.header: r for r in M.match_columns(["Favourite colour"], [["red"], ["blue"]])}
    assert res["Favourite colour"].level == "none"


def test_ambiguous_columns_are_left_for_the_user(monkeypatch):
    # two different fields about equally likely for the same heading: don't guess
    monkeypatch.setattr(M, "header_scores", lambda h, idx: [M.Candidate("phone", 0.90, "a"), M.Candidate("mobile", 0.88, "b")])
    r = M.match_columns(["Whatever"], [["x"], ["y"]])[0]
    assert r.level == "ambiguous" and r.field is None and len(r.candidates) >= 2 and "or" in r.reason


def test_second_email_and_phone_columns_fill_the_next_slots():
    got = fields(["Email", "E-mail Address", "Phone", "Telephone"])
    assert got["Email"] == "email" and got["E-mail Address"] == "email_2" and got["Phone"] == "phone" and got["Telephone"] == "other_phone"


def test_named_custom_fields_are_matched_by_label():
    assert fields(["ABTA No", "Surname"], custom={"custom:user3": ["ABTA number", "user3"]})["ABTA No"] == "custom:user3"


def test_outlook_and_act_exports_map_completely():
    outlook = ["First Name", "Last Name", "Company", "Job Title", "Business Street", "Business City", "Business State", "Business Postal Code",
               "Business Country/Region", "Business Phone", "Mobile Phone", "Home Phone", "E-mail Address", "E-mail 2 Address", "Notes"]
    got = fields(outlook)
    assert list(got.values()) == ["first_name", "last_name", "company", "job_title", "address_line1", "city", "state", "postcode", "country",
                                  "phone", "mobile", "home_phone", "email", "email_2", "notes"]


# ---- parsing --------------------------------------------------------------------------------

def xlsx_bytes(*sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_skips_title_rows_and_picks_the_busy_sheet():
    data = xlsx_bytes(("Cover", [["BMI list"]]), ("Data", [["Contacts export 2026"], [], ["Name", "Email"], ["Ann Lee", "ann@x.com"], ["Bob Marsh", "bob@y.com"]]))
    kind, p = P.parse("x.xlsx", data)
    assert p.sheet_name == "Data" and len(p.sheets) == 2
    row, has = P.detect_header(p.rows)
    headers, body = P.split_rows(p.rows, row, has)
    assert headers == ["Name", "Email"] and len(body) == 2


def test_parse_csv_semicolons_and_cp1252():
    data = "Prénom;Nom;E-mail\nAnne;Lée;a@x.com\n".encode("cp1252")
    _, p = P.parse("a.csv", data)
    assert p.rows[0] == ["Prénom", "Nom", "E-mail"] and p.rows[1][1] == "Lée"


def test_no_header_row_is_detected():
    rows = [["Ann Lee", "ann@x.com", "020 7946 0001"], ["Bob Marsh", "bob@y.com", "020 7946 0002"]]
    row, has = P.detect_header(rows)
    assert has is False
    headers, body = P.split_rows(rows, row, has)
    assert headers == ["A", "B", "C"] and len(body) == 2


def test_bad_files_get_plain_messages():
    with pytest.raises(P.ImportFileError):
        P.parse("x.docx", b"hello")
    with pytest.raises(P.ImportFileError):
        P.parse("x.xlsx", b"PK not really")


# ---- cleaning --------------------------------------------------------------------------------

def test_name_splitting():
    assert E.split_name("Dr Ann Marie van der Berg Jr") == {"prefix": "Dr", "first": "Ann", "middle": "Marie", "last": "van der Berg", "suffix": "Jr"}
    assert E.split_name("Lee, Ann")["first"] == "Ann" and E.split_name("Lee, Ann")["last"] == "Lee"
    assert E.split_name("Madonna") == {"first": "Madonna", "one_word": True}
    assert E.tidy_name("JOHN SMITH") == "John Smith" and E.tidy_name("MCDONALD") == "McDonald" and E.tidy_name("McDonald") == "McDonald"


def test_email_and_phone_cleaning():
    assert E.parse_emails("Ann <ANN@X.com>; bob@y.com, nonsense") == (["ann@x.com", "bob@y.com"], ["nonsense"])
    assert E.parse_phone("4.47E+11")[1] and E.parse_phone("123")[1] and E.parse_phone("n/a") == (None, None)
    assert E.parse_phone("+44 (0)20 7946 0001") == ("+44 (0)20 7946 0001", None)


# ---- the whole import, through the API -------------------------------------------------------

@pytest.fixture()
def world(db_session):
    pub = db_session.query(Publication).filter_by(slug="onboardhospitality").first()
    if not pub:
        pub = Publication(id=uuid.uuid4(), slug="onboardhospitality", name="Onboard Hospitality")
        db_session.add(pub)
        db_session.flush()
    admin = make_user(db_session, role="admin")
    return {"h": identity_headers(admin), "admin": admin}


def upload(client, world, rows, name="people.xlsx"):
    data = xlsx_bytes(("Sheet1", rows)) if name.endswith(".xlsx") else "\n".join(",".join(r) for r in rows).encode()
    r = client.post("/api/contact-imports", files={"file": (name, data)}, data={"source_db": "onboardhospitality"}, headers=world["h"])
    assert r.status_code == 201, r.text
    return r.json()


ROWS = [["First Name", "SURNAME", "E-mail", "Telephone", "Mobile", "Company", "Position", "Post Code", "Town", "Notes"],
        ["ann", "LEE", "Ann@Sunny.example", "020 7946 0001", "07700 900123", "Sunny Travel Ltd", "Travel Counsellor", "sw1a1aa", "London", "Met at show"],
        ["Bob", "Marsh", "bob@y.example", "", "", "Sunny Travel", "Buyer", "", "", ""],
        ["", "", "", "", "", "Nobody Ltd", "", "", "", ""],
        ["Cy", "Dunn", "not-an-email", "12", "", "", "", "", "", ""]]


def test_upload_maps_columns_automatically(client, world):
    imp = upload(client, world, ROWS)
    got = {c["header"]: c["field"] for c in imp["columns"]}
    assert got == {"First Name": "first_name", "SURNAME": "last_name", "E-mail": "email", "Telephone": "phone", "Mobile": "mobile", "Company": "company",
                   "Position": "job_title", "Post Code": "postcode", "Town": "city", "Notes": "notes"}
    assert imp["row_count"] == 4 and imp["problems"] == [] and imp["options"]["source_db"] == "onboardhospitality"


def test_review_then_import_then_undo(client, db_session, world):
    imp = upload(client, world, ROWS)
    rv = client.post(f"/api/contact-imports/{imp['id']}/review", json={}, headers=world["h"]).json()
    assert rv["totals"]["new"] == 3 and rv["totals"]["error"] == 1 and rv["totals"]["warnings"] >= 1
    by_n = {r["n"]: r for r in rv["rows"]}
    assert by_n[3]["status"] == "error" and any("isn't a valid email" in i["text"] for i in by_n[4]["issues"])

    r = client.post(f"/api/contact-imports/{imp['id']}/commit", headers=world["h"])
    assert r.status_code == 200 and r.json()["result"]["created"] == 3
    ann = db_session.query(Contact).filter_by(first_name="Ann", last_name="Lee").one()
    assert ann.full_name == "Ann Lee" and ann.source_db == "onboardhospitality" and ann.custom_fields["_import"]["row"] == 1
    assert [e.address for e in db_session.query(Email).filter_by(contact_id=ann.id)] == ["ann@sunny.example"]
    assert {(p.type_label, p.number) for p in db_session.query(Phone).filter_by(contact_id=ann.id)} == {("Business", "020 7946 0001"), ("Mobile", "07700 900123")}
    assert db_session.query(Address).filter_by(contact_id=ann.id).one().postal_code == "SW1A 1AA"
    assert db_session.query(Note).filter_by(entity_id=ann.id).one().body == "Met at show"
    # "Sunny Travel Ltd" and "Sunny Travel" are one company
    assert db_session.query(Company).filter(Company.name.ilike("sunny travel%")).count() == 1
    bob = db_session.query(Contact).filter_by(last_name="Marsh").one()
    assert bob.company_id == ann.company_id

    u = client.post(f"/api/contact-imports/{imp['id']}/undo", headers=world["h"]).json()
    assert u["deleted"] == 3 and u["kept"] == 0
    db_session.expire_all()
    assert db_session.query(Contact).filter(Contact.last_name.in_(["Lee", "Marsh", "Dunn"])).count() == 0
    assert db_session.query(Company).filter(Company.name.ilike("sunny travel%")).count() == 0


def test_existing_people_are_found_and_options_are_respected(client, db_session, world):
    c = Contact(id=uuid.uuid4(), source_db="onboardhospitality", source_act_id="old1", first_name="Ann", last_name="Lee", full_name="Ann Lee", job_title=None, custom_fields={})
    db_session.add(c)
    db_session.flush()
    db_session.add(Email(id=uuid.uuid4(), source_db="onboardhospitality", source_act_id="oe1", contact_id=c.id, address="ann@sunny.example", is_primary=True))
    db_session.flush()
    imp = upload(client, world, ROWS)
    rv = client.post(f"/api/contact-imports/{imp['id']}/review", json={"status": "update"}, headers=world["h"]).json()
    assert rv["total_rows"] == 1 and rv["rows"][0]["match"]["id"] == str(c.id) and rv["rows"][0]["match"]["kind"] == "email"
    # skip option
    client.patch(f"/api/contact-imports/{imp['id']}", json={"options": {"on_duplicate": "skip"}}, headers=world["h"])
    assert client.post(f"/api/contact-imports/{imp['id']}/review", json={}, headers=world["h"]).json()["totals"]["skip_existing"] == 1
    # fill blanks (default): the job title is added, nothing duplicated
    client.patch(f"/api/contact-imports/{imp['id']}", json={"options": {"on_duplicate": "fill_blanks"}}, headers=world["h"])
    r = client.post(f"/api/contact-imports/{imp['id']}/commit", headers=world["h"]).json()
    assert r["result"]["updated"] == 1 and r["result"]["created"] == 2
    db_session.expire_all()
    assert db_session.get(Contact, c.id).job_title == "Travel Counsellor"
    assert db_session.query(Email).filter_by(contact_id=c.id).count() == 1
    assert db_session.query(Phone).filter_by(contact_id=c.id).count() == 2
    # undo puts the contact back as it was
    client.post(f"/api/contact-imports/{imp['id']}/undo", headers=world["h"])
    db_session.expire_all()
    assert db_session.get(Contact, c.id).job_title is None and db_session.query(Phone).filter_by(contact_id=c.id).count() == 0


def test_manual_mapping_custom_fields_group_and_exclusions(client, db_session, world):
    imp = upload(client, world, [["Who", "Mail", "ABTA", "Segment"], ["Ann Lee", "ann@x.example", "A1234", "Counsellors"], ["Bob Marsh", "bob@y.example", "B9", "Counsellors"]])
    got = {c["header"]: c for c in imp["columns"]}
    assert got["Mail"]["field"] == "email"
    patch = {"mapping": {"0": {"field": "full_name"}, "2": {"field": "new_custom", "name": "ABTA number"}, "3": {"field": "group"}},
             "options": {"new_group_name": "Imported test"}, "excluded": [2]}
    r = client.patch(f"/api/contact-imports/{imp['id']}", json=patch, headers=world["h"])
    assert r.status_code == 200 and r.json()["problems"] == []
    done = client.post(f"/api/contact-imports/{imp['id']}/commit", headers=world["h"]).json()
    assert done["result"]["created"] == 1 and done["result"]["skipped"] == 1
    ann = db_session.query(Contact).filter_by(last_name="Lee").one()
    assert ann.first_name == "Ann" and ann.custom_fields["abta_number"] == "A1234"
    names = {g.name for g in db_session.query(Group).join(GroupMembership, GroupMembership.group_id == Group.id).filter(GroupMembership.contact_id == ann.id)}
    assert names == {"Imported test", "Counsellors"}
    # the new custom field now exists, named properly
    assert {f["key"]: f["label"] for f in client.get("/api/contacts/fields").json()}["custom:abta_number"] == "ABTA number"


def test_problems_block_the_import(client, world):
    imp = upload(client, world, [["Name", "Email", "Notes"], ["Ann Lee", "a@x.example", "hi"]])
    r = client.patch(f"/api/contact-imports/{imp['id']}", json={"mapping": {"0": {"field": "email"}}}, headers=world["h"]).json()
    assert any("both set to Email" in p for p in r["problems"])
    assert client.post(f"/api/contact-imports/{imp['id']}/commit", headers=world["h"]).status_code == 422


def test_other_database_is_out_of_reach_for_a_rep(client, db_session, world):
    rep = make_user(db_session, role="sales")
    h = identity_headers(rep, access=[("sellingtravel", None)])
    data = xlsx_bytes(("S", [["Name", "Email"], ["Ann Lee", "a@x.example"]]))
    r = client.post("/api/contact-imports", files={"file": ("p.xlsx", data)}, data={"source_db": "onboardhospitality"}, headers=h)
    assert r.status_code == 201 and r.json()["options"]["source_db"] != "onboardhospitality"
    iid = r.json()["id"]
    assert client.patch(f"/api/contact-imports/{iid}", json={"options": {"source_db": "onboardhospitality"}}, headers=h).status_code == 403
    assert client.get(f"/api/contact-imports/{iid}", headers=identity_headers(make_user(db_session, role="sales"))).status_code == 404


def test_template_and_report_download(client, world):
    assert client.get("/api/contact-imports/template.xlsx", headers=world["h"]).status_code == 200
    imp = upload(client, world, ROWS)
    rep = client.get(f"/api/contact-imports/{imp['id']}/report.xlsx", headers=world["h"])
    ws = openpyxl.load_workbook(io.BytesIO(rep.content)).active
    assert ws.max_row == 5 and "Can't import" in {c.value for c in ws["B"]}
