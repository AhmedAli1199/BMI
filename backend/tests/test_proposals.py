"""Proposal builder (SALES-020) + logging (SALES-009)."""
import io
import uuid
from datetime import date

import docx
import pytest

from app.models import Company, Note, SalesEdition, SalesOrder, SalesRate, SalesTitle
from app.models.messaging import Reminder
from app.proposals import drafting
from app.sales.reference import ensure_reference_data
from tests.conftest import identity_headers, make_user

YEAR = date.today().year


@pytest.fixture()
def world(db_session):
    ensure_reference_data(db_session)
    title = db_session.query(SalesTitle).filter_by(slug="obh").one()
    company = Company(id=uuid.uuid4(), name="Delta Air Lines", source_db="onboardhospitality", source_act_id=uuid.uuid4().hex[:12])
    db_session.add(company)
    ed = SalesEdition(id=uuid.uuid4(), title_id=title.id, year=YEAR - 1, name="99")
    db_session.add(ed)
    db_session.flush()
    db_session.add(SalesOrder(id=uuid.uuid4(), edition_id=ed.id, client_name="Delta Air Lines", company_id=company.id,
                              value_gbp=4200, size="FP", booked_on=date(YEAR - 1, 3, 1)))
    rate = SalesRate(id=uuid.uuid4(), title_id=title.id, year=YEAR, product="FP", price_gbp=4500)
    db_session.add(rate)
    db_session.flush()
    user = make_user(db_session, role="sales")
    return {"title": title, "company": company, "rate": rate, "user": user, "h": identity_headers(user)}


def create(client, w, **kw):
    body = {"company_id": str(w["company"].id), "title_id": str(w["title"].id), "use_ai": False,
            "lines": [{"product": "FP", "qty": 2, "source": "rate_card", "unit_price": 1}], **kw}
    return client.post("/api/proposals", json=body, headers=w["h"])


def test_create_prices_from_rate_card_not_browser(client, world):
    r = create(client, world)
    assert r.status_code == 201, r.text
    p = r.json()
    assert p["lines"][0]["unit_price"] == 4500  # browser sent 1
    assert p["total_gbp"] == 9000
    assert p["template"] == "obh"
    assert [s["heading"] for s in p["sections"]] == ["Introduction", "Your history with us", "What we propose", "Investment", "Next steps"]
    assert "£4,200" in p["sections"][1]["body"]
    assert p["drafted_by"] == "template"


def test_flags_for_missing_inputs(client, db_session, world):
    other = Company(id=uuid.uuid4(), name="Newco", source_db="onboardhospitality", source_act_id=uuid.uuid4().hex[:12])
    db_session.add(other)
    db_session.flush()
    r = client.post("/api/proposals", json={"company_id": str(other.id), "title_id": str(world["title"].id), "use_ai": False,
                                            "year": YEAR + 5}, headers=world["h"])
    flags = " ".join(r.json()["flags"])
    assert "No bookings on record" in flags and "rate card" in flags and "No products" in flags


def test_unknown_rate_card_product_rejected(client, world):
    r = client.post("/api/proposals", json={"company_id": str(world["company"].id), "title_id": str(world["title"].id),
                                            "lines": [{"product": "Gold sponsorship", "source": "rate_card"}]}, headers=world["h"])
    assert r.status_code == 422


def test_guard_rejects_invented_figures():
    facts = "Delta 2026 £4,200 4200 FP 2"
    allowed = drafting.numbers_in(facts)
    assert drafting.guard("You booked an FP for £4,200 in 2026.", allowed)
    assert not drafting.guard("You will reach 85,000 readers.", allowed)
    assert not drafting.guard("Take 15% off at £3,000.", allowed)


def test_ai_text_with_invented_figures_falls_back(monkeypatch, client, world):
    from app.automations import llm
    monkeypatch.setattr(llm, "is_configured", lambda: True)
    monkeypatch.setattr(llm, "extract_json", lambda *a, **k: {
        "intro": "Hello and thank you.", "history": "We reach 85,000 readers.", "proposal": "A full page.", "next_steps": "- Reply"})
    body = {"company_id": str(world["company"].id), "title_id": str(world["title"].id), "lines": []}
    p = client.post("/api/proposals", json=body, headers=world["h"]).json()
    by_kind = {s["kind"]: s["body"] for s in p["sections"]}
    assert by_kind["intro"] == "Hello and thank you."          # clean AI text kept
    assert "85,000" not in by_kind["history"]                  # invented figure replaced by standard wording
    assert p["drafted_by"] == "ai"


def test_edit_download_and_finish(client, db_session, world):
    p = create(client, world).json()
    secs = p["sections"]
    secs[0]["body"] = "Edited **intro**"
    r = client.patch(f"/api/proposals/{p['id']}", json={"sections": secs, "campaign_name": "Delta Air Lines 2026/27"}, headers=world["h"])
    assert r.status_code == 200
    d = client.get(f"/api/proposals/{p['id']}/download", headers=world["h"])
    assert d.status_code == 200 and "Delta Air Lines 2026-27 proposal.docx" in d.headers["content-disposition"]
    document = docx.Document(io.BytesIO(d.content))
    text = "\n".join(x.text for x in document.paragraphs)
    assert "Edited intro" in text and "Total: £9,000.00" in text
    headers = [x.text for sec in document.sections for part in (sec.header, sec.first_page_header) for x in part.paragraphs]
    assert any("Delta Air Lines 2026/27" in h for h in headers) and not any("Campaign/Advertiser" in h for h in headers)

    f = client.post(f"/api/proposals/{p['id']}/finish", json={"via": "downloaded", "follow_up_days": 10}, headers=world["h"])
    assert f.status_code == 200 and f.json()["status"] == "sent" and f.json()["follow_up_due"]
    note = db_session.query(Note).filter_by(note_type="Proposal sent", entity_id=world["company"].id).one()
    assert "£9,000" in note.body and "before VAT" in note.body
    assert db_session.query(Reminder).filter_by(company_id=world["company"].id, status="open").count() == 1
    assert client.post(f"/api/proposals/{p['id']}/finish", json={}, headers=world["h"]).status_code == 409
    assert client.delete(f"/api/proposals/{p['id']}", headers=world["h"]).status_code == 409


def test_rep_sees_only_own_proposals(client, db_session, world):
    p = create(client, world).json()
    other = make_user(db_session, role="sales")
    h2 = identity_headers(other)
    assert client.get(f"/api/proposals/{p['id']}", headers=h2).status_code == 404
    assert client.get("/api/proposals", headers=h2).json() == []
    admin = make_user(db_session, role="admin")
    assert len(client.get("/api/proposals", headers=identity_headers(admin)).json()) == 1


def _connect(db_session, user):
    from app.models.messaging import MailAccount
    db_session.add(MailAccount(id=uuid.uuid4(), user_id=user.id, email="rep@bmi.test", refresh_token_enc="x"))
    db_session.flush()


def test_send_from_outlook_attaches_word_and_logs(monkeypatch, client, db_session, world):
    from app.services import outlook
    sent = {}
    monkeypatch.setattr(outlook, "send_mail", lambda db, acct, **kw: sent.update(kw))
    p = create(client, world).json()
    body = {"to": ["buyer@delta.test"], "subject": "Proposal", "body": "Hi"}
    assert client.post(f"/api/proposals/{p['id']}/send", json=body, headers=world["h"]).status_code == 409  # Outlook not connected
    _connect(db_session, world["user"])
    r = client.post(f"/api/proposals/{p['id']}/send", json=body, headers=world["h"])
    assert r.status_code == 200 and r.json()["status"] == "sent" and r.json()["sent_via"] == "outlook"
    name, ctype, data = sent["attachments"][0]
    assert name.endswith("proposal.docx") and data[:2] == b"PK"
    note = db_session.query(Note).filter_by(note_type="Proposal sent", entity_id=world["company"].id).one()
    assert "buyer@delta.test" in note.body
    assert client.post(f"/api/proposals/{p['id']}/send", json=body, headers=world["h"]).status_code == 409


def test_failed_send_is_not_logged(monkeypatch, client, db_session, world):
    from app.services import outlook
    def boom(*a, **k):
        raise RuntimeError("refused")
    monkeypatch.setattr(outlook, "send_mail", boom)
    _connect(db_session, world["user"])
    p = create(client, world).json()
    r = client.post(f"/api/proposals/{p['id']}/send", json={"to": ["a@b.test"], "subject": "s", "body": "b"}, headers=world["h"])
    assert r.status_code == 502
    assert client.get(f"/api/proposals/{p['id']}", headers=world["h"]).json()["status"] == "draft"
    assert db_session.query(Note).filter_by(note_type="Proposal sent").count() == 0


def test_email_draft_uses_contact_address(client, db_session, world):
    from app.models import Contact, Email
    c = Contact(id=uuid.uuid4(), source_db="onboardhospitality", source_act_id="c1", first_name="Ann", last_name="Lee", full_name="Ann Lee", company_id=world["company"].id)
    db_session.add(c)
    db_session.flush()
    db_session.add(Email(id=uuid.uuid4(), source_db="onboardhospitality", source_act_id="e1", contact_id=c.id, address="ann@delta.test", is_primary=True))
    db_session.flush()
    p = create(client, world, contact_id=str(c.id)).json()
    d = client.get(f"/api/proposals/{p['id']}/email-draft", headers=world["h"]).json()
    assert d["to"] == ["ann@delta.test"] and d["body"].startswith("Hi Ann") and not d["outlook_connected"]


def test_revise_with_an_instruction(client, world, monkeypatch):
    from app.automations import llm
    p = create(client, world).json()
    monkeypatch.setattr(llm, "is_configured", lambda: True)
    seen = {}

    def fake(system, user, **kw):
        seen["user"] = user
        return {"summary": "Shorter intro and a new Awards section.", "sections": [
            {"kind": "intro", "heading": "Hello Delta", "body": "A short hello."},
            {"kind": "custom", "heading": "Onboard Awards", "body": "Enter the Awards this year."},
            {"kind": "proposal", "heading": "Our idea", "body": "Two full pages at £4,500 each."}]}
    monkeypatch.setattr(llm, "extract_json", fake)
    r = client.post(f"/api/proposals/{p['id']}/revise", json={"instruction": "make the intro shorter and add an Awards section"}, headers=world["h"])
    assert r.status_code == 200, r.text
    out = r.json()
    assert [s["heading"] for s in out["proposal"]["sections"]] == ["Hello Delta", "Onboard Awards", "Our idea", "Investment"]
    assert out["summary"].startswith("Shorter") and len(out["previous_sections"]) == 5
    assert "make the intro shorter" in seen["user"]
    # an invented figure is refused and nothing changes
    monkeypatch.setattr(llm, "extract_json", lambda *a, **k: {"sections": [{"kind": "intro", "heading": "Hi", "body": "We reach 95,000 readers."}]})
    r = client.post(f"/api/proposals/{p['id']}/revise", json={"instruction": "sell harder"}, headers=world["h"])
    assert r.status_code == 422 and "95000" in r.json()["detail"]
    assert client.get(f"/api/proposals/{p['id']}", headers=world["h"]).json()["sections"][0]["heading"] == "Hello Delta"
    # unless the salesperson gave that figure themselves
    r = client.post(f"/api/proposals/{p['id']}/revise", json={"instruction": "mention we reach 95,000 readers"}, headers=world["h"])
    assert r.status_code == 200


def test_transcribe_needs_audio(client, world, monkeypatch):
    from app.automations import llm
    assert client.post("/api/proposals/transcribe", content=b"", headers=world["h"]).status_code == 422
    monkeypatch.setattr(llm, "transcribe_audio", lambda audio, mime, **k: "make it friendlier")
    r = client.post("/api/proposals/transcribe", content=b"abc", headers={**world["h"], "Content-Type": "audio/webm"})
    assert r.json() == {"text": "make it friendlier"}


def test_several_issues_titles_options_and_discounts(client, db_session, world):
    """Matt's ask: issue 108 plus digital in one proposal, a year's programme as an option, % off a line and the total."""
    t = world["title"]
    web = db_session.query(SalesTitle).filter_by(slug="obh-web").one()
    eds = [SalesEdition(id=uuid.uuid4(), title_id=t.id, year=YEAR + 1, name=n, edition_date=date(YEAR + 1, m, 10), kind="issue")
           for n, m in (("106", 6), ("107", 8), ("108", 12))]
    db_session.add_all(eds)
    db_session.flush()
    lines = [
        {"product": "FP", "source": "rate_card", "issues": [str(eds[2].id)], "option": "Option A: Issue 108"},
        {"product": "FP", "source": "rate_card", "issues": [str(e.id) for e in eds], "discount_pct": 0.1, "option": "Option B: The year"},
        {"product": "Newsletter banner", "source": "manual", "unit_price": 500, "qty": 2, "title_id": str(web.id)},
    ]
    r = create(client, world, lines=lines, kind="mixed", discount_pct=0.05)
    assert r.status_code == 201, r.text
    p = r.json()
    tots = {x["option"]: x for x in p["totals"]}
    assert tots["Option A: Issue 108"]["subtotal_gbp"] == 4500 + 1000
    assert tots["Option B: The year"]["subtotal_gbp"] == 3 * 4500 * 0.9 + 1000
    assert tots["Option B: The year"]["total_gbp"] == round((3 * 4500 * 0.9 + 1000) * 0.95, 2)
    assert p["total_gbp"] == tots["Option A: Issue 108"]["total_gbp"] and p["kind_label"] == "A mixed package"
    year_line = next(ln for ln in p["lines"] if ln.get("option") == "Option B: The year")
    assert year_line["qty"] == 3 and year_line["issue_labels"][0].startswith("Issue 106")
    assert set(p["titles"]) == {t.name, web.name}
    assert "two options" in next(s for s in p["sections"] if s["kind"] == "proposal")["body"]
    # Word file shows each option and its total
    d = client.get(f"/api/proposals/{p['id']}/download", headers=world["h"])
    text = "\n".join(par.text for par in docx.Document(io.BytesIO(d.content)).paragraphs)
    assert "Option B: The year total" in text and "less 10%" in text and "Less 5% discount" in text and "Runs in Issue 106" in text
    # the client picks option B: the order starts from it, with its issues and discounts
    pre = client.get(f"/api/sales/deals/prefill?proposal_id={p['id']}&option=Option B: The year", headers=world["h"]).json()
    assert pre["option"] == "Option B: The year" and pre["discount_pct"] == 0.05
    assert [len(ln["placements"]) for ln in pre["lines"]] == [3, 2] and pre["lines"][0]["discount_pct"] == 0.1
    # change the total discount without touching the lines
    r = client.patch(f"/api/proposals/{p['id']}", json={"discount_pct": 0}, headers=world["h"])
    assert r.json()["total_gbp"] == 5500
