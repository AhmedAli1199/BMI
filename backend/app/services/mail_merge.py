"""Mail merge - Act!'s Write > Mail Merge, rebuilt.

Outputs:
  email  - personalised emails sent from the user's own Outlook (queued,
           throttled by the sender job below, resumable, per-recipient log)
  word   - one letter per contact, one page each, as a .docx
  labels - address labels (3 x 7 = 21 per A4 sheet, Avery L7160 size) .docx
  data   - the merge data itself (.xlsx / .csv) for Word's own mail merge,
           Mailchimp, or a printer

Merge fields are {{field}} or {{field|fallback}} - e.g.
"Dear {{first_name|colleague}}," - see FIELDS for the list.
"""
from __future__ import annotations

import csv
import html
import io
import logging
import re
import uuid
from datetime import date, datetime, timezone

from sqlalchemy import select, update
from sqlalchemy.orm import Session

from app.models import HistoryEntry, User
from app.models.messaging import MailAccount, MailAttachment, MailMerge, MailMergeRecipient
from app.services.contact_export import HEADERS, contact_rows

logger = logging.getLogger("app.services.mail_merge")

MANUAL_SOURCE_DB = "manual"

# key -> (label, example)
FIELDS: dict[str, tuple[str, str]] = {
    "first_name": ("First name", "Sarah"),
    "last_name": ("Last name", "Jones"),
    "full_name": ("Full name", "Sarah Jones"),
    "salutation": ("Salutation (Act! 'Dear' name)", "Sarah"),
    "job_title": ("Job title", "Head of Onboard Retail"),
    "company": ("Company", "Virgin Atlantic"),
    "email": ("Email", "sarah.jones@example.com"),
    "phone": ("Phone", "+44 20 7000 0000"),
    "address_1": ("Address line 1", "The VHQ"),
    "address_2": ("Address line 2", "Fleming Way"),
    "city": ("City", "Crawley"),
    "state": ("County / State", "West Sussex"),
    "postcode": ("Postcode", "RH10 9DF"),
    "country": ("Country", "United Kingdom"),
    "address_block": ("Full address block", "The VHQ\nFleming Way\nCrawley\nRH10 9DF"),
    "my_name": ("My name", "Your name"),
    "my_email": ("My email", "you@bmipublishing.co.uk"),
    "today": ("Today's date", date.today().strftime("%d %B %Y")),
}
_IDX = {h: i for i, h in enumerate(HEADERS)}
_FIELD_RE = re.compile(r"\{\{\s*([a-z_0-9]+)\s*(?:\|([^}]*))?\}\}")


def contexts(db: Session, ids: list[uuid.UUID], me: User | None = None) -> list[tuple[uuid.UUID, dict]]:
    rows = contact_rows(db, ids)
    out = []
    today = date.today().strftime("%d %B %Y")
    # contact_rows drops ids that no longer exist - realign against the survivors
    from app.models import Contact

    existing = [i for i in ids if i in set(db.scalars(select(Contact.id).where(Contact.id.in_(ids))))] if ids else []
    for cid, r in zip(existing, rows):
        g = lambda h: (r[_IDX[h]] or "") if r[_IDX[h]] is not None else ""  # noqa: E731
        first, last = str(g("First name")), str(g("Last name"))
        addr_lines = [str(g(h)) for h in ("Address 1", "Address 2", "Address 3", "City", "County/State", "Postcode", "Country") if g(h)]
        out.append((cid, {
            "first_name": first, "last_name": last, "full_name": " ".join(x for x in (first, last) if x),
            "salutation": str(g("Salutation")) or first, "job_title": str(g("Job title")), "company": str(g("Company")),
            "email": str(g("Email")), "phone": str(g("Phone") or g("Mobile")), "address_1": str(g("Address 1")),
            "address_2": str(g("Address 2")), "city": str(g("City")), "state": str(g("County/State")),
            "postcode": str(g("Postcode")), "country": str(g("Country")), "address_block": "\n".join(addr_lines),
            "my_name": me.name if me else "", "my_email": me.email if me else "", "today": today,
            "_unsubscribed": bool(g("Unsubscribed")), "_bounced": bool(g("Bounced")),
        }))
    return out


def render(template: str, ctx: dict) -> str:
    def sub(m):
        val = ctx.get(m.group(1))
        return str(val) if val else (m.group(2) or "").strip()
    return _FIELD_RE.sub(sub, template or "")


def unknown_fields(*templates: str | None) -> list[str]:
    found = {m.group(1) for t in templates if t for m in _FIELD_RE.finditer(t)}
    return sorted(found - FIELDS.keys())


def is_html(body: str) -> bool:
    return bool(re.search(r"<(p|div|br|table|span|a|b|strong|em|ul|ol|h\d)(\s[^>]*)?/?>", body or "", re.I))


def to_html(body: str) -> str:
    """Rendered body -> email HTML. Plain text keeps its line breaks; a
    body that's already HTML is used as-is."""
    if is_html(body):
        return body
    esc = html.escape(body).replace("\n", "<br>")
    return f"<div style=\"font-family:Calibri,Arial,sans-serif;font-size:11pt\">{esc}</div>"


def to_text(body: str) -> str:
    if not is_html(body):
        return body
    t = re.sub(r"<br\s*/?>|</p>|</div>", "\n", body, flags=re.I)
    return html.unescape(re.sub(r"<[^>]+>", "", t)).strip()


# ---- Documents --------------------------------------------------------------

def letters_docx(items: list[tuple[str]]) -> io.BytesIO:
    """One page per letter."""
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(11)
    for i, text in enumerate(items):
        for line in text.split("\n"):
            doc.add_paragraph(line)
        if i < len(items) - 1:
            doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def labels_docx(blocks: list[str]) -> io.BytesIO:
    """3 columns x 7 rows per A4 page (63.5 x 38.1 mm - Avery L7160)."""
    from docx import Document
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.enum.text import WD_BREAK
    from docx.shared import Mm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(15.1), Mm(0)
    sec.left_margin, sec.right_margin = Mm(7.2), Mm(0)
    doc.styles["Normal"].font.name, doc.styles["Normal"].font.size = "Calibri", Pt(10)
    per_page = 21
    for start in range(0, max(len(blocks), 1), per_page):
        page = blocks[start:start + per_page]
        table = doc.add_table(rows=7, cols=3)
        table.autofit = False
        for r, row in enumerate(table.rows):
            row.height, row.height_rule = Mm(38.1), WD_ROW_HEIGHT_RULE.EXACTLY
            for c, cell in enumerate(row.cells):
                cell.width = Mm(66)
                idx = r * 3 + c
                text = page[idx] if idx < len(page) else ""
                cell.paragraphs[0].text = ""
                for j, line in enumerate(text.split("\n")):
                    p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.left_indent = Mm(3)
                    p.add_run(line)
        if start + per_page < len(blocks):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


DATA_KEYS = [k for k in FIELDS if k not in ("my_name", "my_email", "today", "address_block")]


def data_file(ctxs: list[dict], fmt: str) -> io.BytesIO:
    header = [FIELDS[k][0] for k in DATA_KEYS]
    rows = [[c.get(k, "") for k in DATA_KEYS] for c in ctxs]
    if fmt == "csv":
        s = io.StringIO()
        w = csv.writer(s)
        w.writerow(header)
        w.writerows(rows)
        return io.BytesIO(("﻿" + s.getvalue()).encode("utf-8"))
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Merge data"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---- History ----------------------------------------------------------------

def record_history(db: Session, contact_id: uuid.UUID, history_type: str, subject: str | None, details: str | None,
                   user_id: uuid.UUID | None) -> None:
    db.add(HistoryEntry(
        id=uuid.uuid4(), source_db=MANUAL_SOURCE_DB, source_act_id=str(uuid.uuid4()), entity_type="contact",
        entity_id=contact_id, history_type=history_type, subject=(subject or "")[:512] or None, details=details,
        occurred_at=datetime.now(timezone.utc), created_by_user_id=user_id,
    ))


# ---- Sender job -------------------------------------------------------------

def _split(addrs: str | None) -> list[str]:
    return [a.strip() for a in re.split(r"[;,]", addrs or "") if a.strip()]


def send_queued(db: Session, per_minute: int | None = None) -> int:
    """One tick: sends up to per_minute messages for each active merge
    (each merge is one user's mailbox, so each gets its own budget).
    Returns messages sent."""
    from app.services import outlook
    from app.services.notify import create_notification

    from app.automations.runtime_settings import get_int

    per_minute = per_minute or get_int(db, "mail_merge_per_minute")
    total = 0
    merges = db.scalars(select(MailMerge).where(MailMerge.output == "email",
                                                MailMerge.status.in_(("queued", "sending")))
                        .order_by(MailMerge.created_at)).all()
    for m in merges:
        acct = db.scalar(select(MailAccount).where(MailAccount.user_id == m.created_by_user_id)) if m.created_by_user_id else None
        if not acct:
            m.status, m.error = "failed", "No Outlook mailbox connected."
            db.commit()
            continue
        me = db.get(User, m.created_by_user_id)
        m.status = "sending"
        db.commit()
        atts = [(a.filename, a.content_type, a.data) for a in db.scalars(select(MailAttachment).where(MailAttachment.merge_id == m.id))]
        batch = db.scalars(select(MailMergeRecipient).where(MailMergeRecipient.merge_id == m.id,
                                                            MailMergeRecipient.status == "queued")
                           .order_by(MailMergeRecipient.id).limit(per_minute)).all()
        ctx_by_id = dict(contexts(db, [r.contact_id for r in batch if r.contact_id], me))
        stop = False
        for r in batch:
            ctx = ctx_by_id.get(r.contact_id)
            if not ctx or not r.email:
                r.status, r.error = "skipped", "contact deleted or no email"
                m.skipped += 1
                db.commit()
                continue
            subject, body = render(m.subject or "", ctx), render(m.body, ctx)
            try:
                outlook.send_mail(db, acct, to=[r.email], subject=subject, html_body=to_html(body),
                                  cc=_split(m.cc), bcc=_split(m.bcc), attachments=atts)
            except outlook.RateLimited:
                stop = True
                break
            except outlook.OutlookAuthError as exc:
                m.status, m.error = "failed", f"{exc} Reconnect Outlook in Settings, then resume."
                db.commit()
                stop = True
                break
            except Exception as exc:
                r.status, r.error = "failed", str(exc)[:500]
                m.failed += 1
                db.commit()
                continue
            r.status, r.sent_at = "sent", datetime.now(timezone.utc)
            m.sent += 1
            total += 1
            if r.contact_id and m.record_history != "none":
                record_history(db, r.contact_id, "E-mail Sent", subject,
                               to_text(body) if m.record_history == "email_full" else None, m.created_by_user_id)
            db.commit()
        if stop:
            if m.status == "failed" and m.created_by_user_id:
                create_notification(db, m.created_by_user_id, "mail_merge", "Mail merge paused",
                                    m.error, f"/mail-merge/{m.id}")
                db.commit()
            continue
        remaining = db.scalar(select(MailMergeRecipient.id).where(MailMergeRecipient.merge_id == m.id,
                                                                  MailMergeRecipient.status == "queued").limit(1))
        if remaining is None:
            m.status, m.finished_at = "done", datetime.now(timezone.utc)
            if m.created_by_user_id:
                create_notification(db, m.created_by_user_id, "mail_merge", f"Mail merge sent: {m.subject or ''}".strip(),
                                    f"{m.sent} sent, {m.failed} failed, {m.skipped} skipped.", f"/mail-merge/{m.id}",
                                    email=False)
            db.commit()
    return total


def cancel(db: Session, m: MailMerge) -> None:
    n = db.execute(update(MailMergeRecipient).where(MailMergeRecipient.merge_id == m.id,
                                                    MailMergeRecipient.status == "queued")
                   .values(status="skipped", error="cancelled")).rowcount
    m.skipped += n
    m.status, m.finished_at = "cancelled", datetime.now(timezone.utc)
