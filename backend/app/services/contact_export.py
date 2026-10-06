"""Contacts -> .xlsx, in lookup order. Used by the contacts-list and group
Export buttons. Batched queries (no per-contact round trips) so a
5,000-contact group exports in a couple of seconds."""
from __future__ import annotations

import io
import uuid
from collections import defaultdict

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Company, Contact, Email, Group, GroupMembership
from app.models.contact_channel import Address, Phone

HEADERS = ["First name", "Last name", "Salutation", "Job title", "Company", "Email", "Phone", "Mobile",
           "Address 1", "Address 2", "Address 3", "City", "County/State", "Postcode", "Country",
           "Database", "Groups", "Unsubscribed", "Bounced"]


def _first(rows, key=lambda r: 0):
    return sorted(rows, key=key)[0] if rows else None


def contact_rows(db: Session, ids: list[uuid.UUID]) -> list[list]:
    """Flat rows in the given order - shared by the Excel export and the
    mail-merge data file."""
    if not ids:
        return []
    out = []
    for i in range(0, len(ids), 2000):
        chunk = ids[i:i + 2000]
        contacts = {c.id: c for c in db.scalars(select(Contact).where(Contact.id.in_(chunk)))}
        comp_ids = {c.company_id for c in contacts.values() if c.company_id}
        companies = {c.id: c.name for c in db.execute(select(Company.id, Company.name).where(Company.id.in_(comp_ids)))} if comp_ids else {}
        emails, phones, addrs, groups = defaultdict(list), defaultdict(list), defaultdict(list), defaultdict(list)
        for e in db.scalars(select(Email).where(Email.contact_id.in_(chunk))):
            emails[e.contact_id].append(e)
        for p in db.scalars(select(Phone).where(Phone.contact_id.in_(chunk))):
            phones[p.contact_id].append(p)
        for a in db.scalars(select(Address).where(Address.contact_id.in_(chunk))):
            addrs[a.contact_id].append(a)
        for cid, name in db.execute(select(GroupMembership.contact_id, Group.name)
                                    .join(Group, Group.id == GroupMembership.group_id)
                                    .where(GroupMembership.contact_id.in_(chunk))):
            groups[cid].append(name)
        for cid in chunk:
            c = contacts.get(cid)
            if not c:
                continue
            email = _first(emails[cid], key=lambda e: not e.is_primary)
            addr = _first(addrs[cid], key=lambda a: not a.is_primary)
            mobile = next((p.number for p in phones[cid] if (p.type_label or "").lower().startswith("mob")), None)
            phone = next((p.number for p in phones[cid] if p.number and p.number != mobile), None)
            out.append([
                c.first_name, c.last_name, c.salutation, c.job_title,
                companies.get(c.company_id) or c.company_name_freetext,
                email.address if email else None, phone, mobile,
                addr.line1 if addr else None, addr.line2 if addr else None, addr.line3 if addr else None,
                addr.city if addr else None, addr.state if addr else None,
                addr.postal_code if addr else None, addr.country if addr else None,
                c.source_db, ", ".join(sorted(set(groups[cid]))),
                "Yes" if (c.is_unsubscribed or c.is_email_opted_out) else "", "Yes" if c.has_bounced else "",
            ])
    return out


def _cell(v):
    if v is None or isinstance(v, (str, int, float)):
        return v
    return ", ".join(map(str, v)) if isinstance(v, list) else str(v)


def contacts_xlsx(db: Session, ids: list[uuid.UUID], title: str = "Contacts") -> io.BytesIO:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook(write_only=False)
    ws = wb.active
    ws.title = title[:31]
    # The Act! custom fields (ABTA number, Type...) ride along as extra columns - only the ones
    # at least one exported contact has a value for, under the names set in app/contacts/fields.py.
    from app.contacts.fields import get_custom_labels, pretty_key
    customs: dict[uuid.UUID, dict] = {}
    for i in range(0, len(ids), 2000):
        for cid, cf in db.execute(select(Contact.id, Contact.custom_fields).where(Contact.id.in_(ids[i:i + 2000]))):
            customs[cid] = cf or {}
    keys = sorted({k for cf in customs.values() for k, v in cf.items() if not k.startswith("_") and v not in (None, "", [], {})})
    labels = get_custom_labels(db)
    ws.append(HEADERS + [labels.get(k) or pretty_key(k) for k in keys])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE4F0")
    present = [i for i in ids if i in customs]
    for cid, row in zip(present, contact_rows(db, present)):
        cf = customs[cid]
        ws.append(row + [_cell(cf.get(k)) for k in keys])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDEFGHIJKLMNOPQRS", (14, 18, 10, 26, 28, 32, 16, 16, 28, 20, 16, 16, 14, 10, 14, 13, 30, 12, 9)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
