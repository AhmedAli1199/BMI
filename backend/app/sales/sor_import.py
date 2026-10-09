"""One-time import of BMI's SOR workbooks (2023-2026) into the Sales
Order Register tables. Run via scripts/import_sor.py.

The workbooks share one core layout - a summary block at the top
(MONTH, Exchange rate, Cumulative value...), then a two-row header
(Date | Client | Size | Series | Position | Salesper. | Rate US$ | £ |
Invoice number | Invoice value | Page number | Invoice difference |
Reason for difference | one commission column per rep) - with a few
per-title variations (events have tables/seats, awards have entries,
TBTM adds a Total column, Visit USA adds Planner/Online columns). So
columns are found by their *labels*, never their position.

What the importer deliberately does NOT do is guess: a date it can't
read, a value typed as "40, 0000", a salesperson code it doesn't know -
each is imported as-is with an `import_warning` the UI surfaces, so a
person fixes it once instead of the importer silently getting it wrong.

Duplicate files:
- "-Copy(1)" files (every 2023 workbook has one) are skipped.
- Visit USA Travel Planner ships as a pair each year: a working copy
  ("...2026.xls", with the rep's notes: "need po", "agency comm") and a
  reconciled "...edition.xls" (notes removed, every difference zeroed).
  Same bookings. The edition file is imported; the working copy's notes
  and pre-reconciliation invoice values are merged onto the matching
  rows so nothing is lost.
"""
from __future__ import annotations

import logging
import re
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from sqlalchemy.orm import Session

from app.models import EditionFeature, SalesEdition, SalesEditionCost, SalesOrder, SalesOrderCredit, SalesRep, SalesTitle
from app.sales.reference import ensure_reference_data, rep_code_for, title_for_file

logger = logging.getLogger("app.sales.sor_import")

Cell = str | float | datetime | None

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}


# ---- Reading --------------------------------------------------------------

def hidden_columns(path: Path) -> dict[str, set[int]]:
    """Sheet name -> 0-based indexes of columns hidden in Excel. Hidden
    commission columns often hold stale copied figures (Dec Print 2025:
    L.Merrigan / S.Thompson / D.Clare hidden, each showing £1,500 on a
    K.Hicks booking) - the person maintaining the sheet never sees them,
    so neither should the import."""
    try:
        if path.suffix.lower() == ".xlsx":
            import openpyxl
            from openpyxl.utils import column_index_from_string

            wb = openpyxl.load_workbook(path, read_only=False)
            out = {}
            for ws in wb.worksheets:
                hidden = set()
                for key, dim in ws.column_dimensions.items():
                    if dim.hidden:
                        lo = column_index_from_string(key) - 1
                        hi = (dim.max or lo + 1) - 1
                        hidden.update(range(min(lo, hi), max(lo, hi) + 1))
                out[ws.title] = hidden
            return out
        import xlrd

        book = xlrd.open_workbook(str(path), formatting_info=True)
        return {sh.name: {c for c, info in sh.colinfo_map.items() if info.hidden} for sh in book.sheets()}
    except Exception:  # noqa: BLE001 - formatting info is a nicety; never block an import on it
        return {}


def fmt_size(v: Cell) -> str | None:
    """A size cell as the sheet shows it. Excel stores a typed "2/3" as
    0.666..., so small fractions are turned back into "2/3"."""
    if v is None:
        return None
    if isinstance(v, float):
        if 0 < v < 1:
            from fractions import Fraction

            f = Fraction(v).limit_denominator(12)
            if abs(float(f) - v) < 1e-6:
                return f"{f.numerator}/{f.denominator}"
        return f"{v:g}"
    return str(v)


def read_workbook(path: Path) -> dict[str, list[list[Cell]]]:
    """Sheet name -> rows of plain Python values (str / float / datetime /
    None), for both legacy .xls (xlrd) and .xlsx (openpyxl)."""
    if path.suffix.lower() == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        out = {}
        for ws in wb.worksheets:
            out[ws.title] = [[_clean(v) for v in row] for row in ws.iter_rows(values_only=True)]
        return out

    import xlrd

    book = xlrd.open_workbook(str(path))
    out = {}
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    except Exception:  # noqa: BLE001 - corrupt date cell, keep raw
                        row.append(cell.value)
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    row.append(None)
                else:
                    row.append(_clean(cell.value))
            rows.append(row)
        out[sh.name] = rows
    return out


def _clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return float(v)
    return v


# ---- Value parsing --------------------------------------------------------

def parse_money(v: Cell) -> tuple[float | None, str | None]:
    """(value, warning). Numbers pass through; text like "£1,250" parses;
    anything ambiguous ("40, 0000") returns (None, warning)."""
    if v is None:
        return None, None
    if isinstance(v, float):
        return round(v, 2), None
    if isinstance(v, datetime):
        return None, f"expected an amount, found a date ({v:%d/%m/%Y})"
    s = str(v).replace("£", "").replace("$", "").strip()
    if re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?|-?\d+(\.\d+)?", s):
        return round(float(s.replace(",", "")), 2), None
    if not re.search(r"\d", s):
        return None, None  # a word ("tbc", "FOC") - not a value, not worth a warning
    return None, f"couldn't read amount {s!r}"


def parse_date(v: Cell) -> tuple[date | None, str | None]:
    if v is None:
        return None, None
    if isinstance(v, datetime):
        return v.date(), None
    if isinstance(v, float):
        return None, None
    s = str(v).strip()
    m = re.fullmatch(r"(\d{1,2})[./](\d{1,2})[./](\d{2}|\d{4})", s)
    if m:
        d, mo, y = int(m[1]), int(m[2]), int(m[3])
        y = y + 2000 if y < 100 else y
        try:
            return date(y, mo, d), None
        except ValueError:
            pass
    return None, f"couldn't read date {s!r}"


def parse_pages(size: str | None, product_line: str) -> float | None:
    """Page-equivalents for print sizes - what the sheet header's "Number
    pages booked" adds up. None for anything that isn't a page (banners,
    listings, event tickets)."""
    if not size or product_line not in ("print",):
        return None
    s = size.lower().replace("½", "1/2").strip()
    m = re.fullmatch(r"(\d+)\s*x\s*(\d+)/(\d+)(\s*page)?", s)
    if m:
        return int(m[1]) * int(m[2]) / int(m[3])
    m = re.fullmatch(r"(\d+(?:\.\d+)?)?\s*(fp|full page|page|dps)\b.*", s)
    if m:
        n = float(m[1]) if m[1] else 1.0
        return n * (2 if m[2] == "dps" else 1)
    m = re.fullmatch(r"(\d+)/(\d+)(\s*page)?", s)
    if m:
        return int(m[1]) / int(m[2])
    m = re.fullmatch(r"(\d+(?:\.\d+)?)", s)
    if m and float(m[1]) <= 10:
        return float(m[1])
    return None


def edition_date_for(year: int, sheet: str, period: Cell) -> date | None:
    if isinstance(period, datetime):
        return period.date()
    for text in (str(period or ""), sheet):
        m = re.search(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*", text.lower())
        if m:
            day = re.search(r"\b(\d{1,2})(?:st|nd|rd|th)\b", text.lower())
            try:
                return date(year, MONTHS[m[1]], int(day[1]) if day else 1)
            except ValueError:
                return date(year, MONTHS[m[1]], 1)
    return None


def edition_kind(product_line: str, sheet: str, period: Cell) -> str:
    if product_line == "events":
        return "event"
    if product_line == "awards":
        return "awards"
    if product_line == "digital" or re.search(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)", sheet.lower()):
        return "month" if product_line == "digital" else "issue"
    return "issue"


# ---- Sheet parsing --------------------------------------------------------

def _norm(v: Cell) -> str:
    return re.sub(r"\s+", " ", str(v)).strip().lower() if v is not None else ""


@dataclass
class ParsedRow:
    row_index: int
    client: str
    cells: dict[str, Cell]
    rep_credits: dict[str, float]
    text: str
    continuation: bool = False
    raw: list = field(default_factory=list)
    extra: dict[str, str] = field(default_factory=dict)


@dataclass
class ParsedSheet:
    name: str
    period: Cell = None
    exchange_rate: float | None = None
    sheet_total: float | None = None
    rows: list[ParsedRow] = field(default_factory=list)
    # (sheet row number, cells) for every non-booking line under the header,
    # blank lines included as [] - read by cost_lines() on events sheets.
    other_rows: list[tuple[int, list]] = field(default_factory=list)
    gbp_col: int = 0


_COLS = {
    "date": ("date",),
    "size": ("size", "sponsor"),
    "series": ("series",),
    "position": ("position",),
    "rep": ("salesper.", "sales"),
    "usd": ("rate us$", "us$", "rate $", "rate"),
    "gbp": ("£",),
    "invoice_number": ("invoice number",),
    "invoice_value": ("invoice value",),
    "reason": ("reason for difference",),
}


def parse_sheet(name: str, rows: list[list[Cell]], hidden: set[int] | None = None) -> ParsedSheet | None:
    if re.search(r"template|do ?n.?t copy|^(blank|sheet)\s*\d*\b", name.strip(), re.IGNORECASE):
        return None  # the copy-from template and unused placeholder sheets
    hdr = next((i for i, r in enumerate(rows[:15]) if any(_norm(v).endswith("client") for v in r)), None)
    if hdr is None or hdr == 0:
        return None

    width = max(len(r) for r in rows[hdr - 1: hdr + 1])
    top = rows[hdr - 1] + [None] * width
    labels = [_norm(f"{top[i] or ''} {rows[hdr][i] if i < len(rows[hdr]) else ''}") for i in range(width)]
    col: dict[str, int] = {}
    client_col = next(i for i, lab in enumerate(labels) if lab.endswith("client"))
    for key, names in _COLS.items():
        for n in names:
            idx = next((i for i, lab in enumerate(labels) if lab == n and i not in col.values()), None)
            if idx is not None:
                col[key] = idx
                break
    if "gbp" not in col:
        return None  # an entries/attendee list (e.g. "IND Entries ALPHA"), not bookings

    # Commission columns: everything right of the reason/difference column
    # whose header is a known rep.
    start = max(col.get("reason", 0), col.get("invoice_value", 0)) + 1
    rep_cols = {i: rep_code_for(rows[hdr][i]) for i in range(start, len(rows[hdr]))
                if rows[hdr][i] is not None and rep_code_for(rows[hdr][i]) and i not in (hidden or set())}

    # Any other labelled column left of the commission block (Seats, Table
    # no., Paid?, Travel Planner / Online, Page number...) is kept as-is.
    # Running totals and the computed difference column are not data.
    def _pretty(i: int) -> str | None:
        # Headers span two rows ("Page" over "number", "Travel" over
        # "Planner"); join them, dropping a leading "Rate"/"Invoice" that
        # belongs to the neighbouring mapped column.
        parts = [v for v in (top[i], rows[hdr][i] if i < len(rows[hdr]) else None)
                 if isinstance(v, str) and v.strip() and not re.fullmatch(r"[\d.,\s]+", v)]
        label = re.sub(r"\s+", " ", " ".join(parts)).strip()
        if len(parts) == 2:
            label = re.sub(r"^(rate|invoice)\s+", "", label, flags=re.IGNORECASE)
        return label or None
    used = set(col.values()) | {client_col}
    extra_cols = {}
    for i in range(start):
        label = _pretty(i) if i not in used else None
        if label and _norm(label) not in ("invoice difference", "difference", "total", "invoice", "rate", "number", "value"):
            extra_cols[i] = label

    parsed = ParsedSheet(name=name.strip(), gbp_col=col["gbp"])
    for r in rows[:hdr]:
        for i, v in enumerate(r):
            lab = _norm(v)
            nxt = next((x for x in r[i + 1:] if x is not None), None)
            if lab in ("month", "date") and parsed.period is None:
                # An empty MONTH cell must not pick up the next label along.
                if not (isinstance(nxt, str) and _norm(nxt) in ("exchange rate", "pages sold")) and not isinstance(nxt, float):
                    parsed.period = nxt
            elif lab == "exchange rate" and isinstance(nxt, float):
                parsed.exchange_rate = nxt
            elif lab.startswith("cumulative value") and parsed.sheet_total is None:
                parsed.sheet_total, _ = parse_money(nxt)

    last_client: str | None = None
    last_row = -10
    for ri in range(hdr + 1, len(rows)):
        r = rows[ri] + [None] * width
        client = r[client_col]
        client = None if client is None or isinstance(client, float) else str(client).strip() or None
        value, _ = parse_money(r[col["gbp"]])
        inv = r[col["invoice_number"]] if "invoice_number" in col else None
        has_invoice = inv is not None and not isinstance(inv, float) and bool(re.search(r"\d", str(inv)))
        continuation = False
        if client is None:
            # A second instalment / extra invoice on the row under a
            # booking, with the client cell left blank.
            # Only in the block of lines directly under a booking (no blank
            # line between) - a £ figure further down is a cost/P&L block
            # (events sheets), not another instalment.
            empty = all(v is None or v == 0 for v in r)
            if empty:
                last_row = -10  # a blank line ends the booking above
                parsed.other_rows.append((ri + 1, []))
                continue
            if last_client and last_row >= 0 and ((value or 0) > 0 or has_invoice):
                client, continuation = last_client, True
            else:
                parsed.other_rows.append((ri + 1, r))
                continue  # a package line ("Enews") under a booking - part of it, no value of its own
        if client.lower().startswith(("total", "totals")):
            continue
        text = " ".join(str(v) for v in r if v is not None)
        looks_like_booking = (
            (value or 0) != 0 or has_invoice
            or ("rep" in col and r[col["rep"]] is not None)
            or ("date" in col and r[col["date"]] is not None)
            # Whole words: "Edinburgh 27th January Contracted Doubletree" is
            # an event-day heading, not a contra booking.
            or re.search(r"\bCANX\b|CANCEL|\bCONTRA\b|MOVED TO", text.upper())
        )
        if not looks_like_booking:
            parsed.other_rows.append((ri + 1, r))
            continue  # cost breakdowns, section labels ("Direct Costs", "Room Hire") typed into the client column
        last_client = client
        last_row = ri
        credits = {}
        for i, code in rep_cols.items():
            amount, _ = parse_money(r[i])
            if amount:
                credits[code] = credits.get(code, 0) + amount
        parsed.rows.append(ParsedRow(
            row_index=ri + 1,
            client=client,
            cells={k: r[i] for k, i in col.items()},
            rep_credits=credits,
            text=text,
            continuation=continuation,
            raw=r,
            extra={lab: (f"{r[i]:g}" if isinstance(r[i], float) else r[i].strftime("%d/%m/%Y") if isinstance(r[i], datetime) else str(r[i]).strip())
                   for i, lab in extra_cols.items() if r[i] is not None and r[i] != 0 and str(r[i]).strip()},
        ))
    return parsed


# ---- Import ---------------------------------------------------------------

# The sheet's own sums - kept for reference, never added up again.
_SUMMARY = re.compile(r"\btotals?\b|^(event profit|profit\b|net\b|\d+\s*% of profit|cumulative|estimated costs$|av cost|"
                      r"(plus|inc|ex)\.? vat\b)", re.IGNORECASE)
_DAY = re.compile(r"^\W*(\d{1,2}(st|nd|rd|th)\b\s+[A-Za-z]|[A-Za-z]{3,}\s+\d{1,2}(st|nd|rd|th)\b)", re.IGNORECASE)
_INCOME = re.compile(r"^(income|revenue|sponsorship|sponsors?)\b", re.IGNORECASE)


def _is_label(v: Cell) -> bool:
    return isinstance(v, str) and bool(re.search(r"[A-Za-z]{2}", v)) and not re.fullmatch(r"\s*[£$]?[\d.,\s]+\s*", v)


def cost_lines(other_rows: list[tuple[int, list]], last_col: int | None = None) -> list[dict]:
    """The costs / P&L area of an events sheet, as lines: venue hire, food,
    AV, travel, photographer, BMI overheads... plus the free-text lines
    around them (contracted on..., deposit paid...) and the sheet's own
    totals for reference. Only text-labelled lines are kept - a bare
    number on its own is a running total. Figures right of `last_col` (the
    £ column) are the sheet's running totals and invoice columns, never a
    cost.

    Headings: in a block of text-only lines (no blank line between) that
    leads into figures, the first line is the heading - an event day or
    venue ("Edinburgh 27th January", "Doubletree - Sue and Stuart
    attending") - for everything under it until the next such block."""
    entries: list[dict] = []
    gap = True
    for row_no, r in other_rows:
        if not r or all(v is None or v == 0 or (isinstance(v, str) and not v.strip()) for v in r):
            gap = True
            continue
        labels = [v.strip() for v in r if _is_label(v)]
        if not labels or len(labels[0]) < 3 or labels[0].lower() in ("total", "totals", "v", "notes"):
            continue
        label = labels[0]
        cells = r[: last_col + 1] if last_col is not None else r
        nums = [round(float(v), 2) for v in cells if isinstance(v, (int, float)) and not isinstance(v, bool) and abs(v) >= 0.01]
        if not nums:  # "£261.50 inc VAT" typed as text in another cell
            for t in labels[1:]:
                m = re.search(r"£\s*([\d,]+(?:\.\d+)?)", t)
                if m:
                    nums = [round(float(m.group(1).replace(",", "")), 2)]
                    label = f"{label} ({t})"
                    break
        text = " ".join(labels).lower()
        if re.match(r"\s*£", label):  # "£2520.33 +VAT" - the sheet's own running total for the block
            m = re.search(r"£\s*([\d,]+(?:\.\d+)?)", label)
            kind, amount, inc_vat = "summary", (round(float(m.group(1).replace(",", "")), 2) if m else None), None
        elif not nums:
            kind, amount, inc_vat = "note", None, None
        else:
            kind = "summary" if _SUMMARY.search(label) else "income" if _INCOME.search(label) else "cost"
            amount = nums[0]
            inc_vat = nums[1] if len(nums) > 1 and "vat" in text and nums[1] > nums[0] else None
        entries.append({"kind": kind, "label": label[:500], "amount_gbp": amount, "amount_inc_vat_gbp": inc_vat,
                        "source_row": row_no, "_gap": gap})
        gap = False

    section: str | None = None
    i = 0
    while i < len(entries):
        e = entries[i]
        if e["kind"] == "note":
            j = i
            while j + 1 < len(entries) and entries[j + 1]["kind"] == "note" and not entries[j + 1]["_gap"]:
                j += 1
            leads_into_figures = j + 1 < len(entries) and not entries[j + 1]["_gap"]
            if leads_into_figures or e["_gap"]:
                # The heading is the line naming a day ("13th April Glasgow",
                # "Edinburgh 27th January") if the block has one, else its first line.
                dated = [k for k in range(i, j + 1) if _DAY.search(entries[k]["label"])]
                h = dated[-1] if dated else i
                for k in range(i, h):
                    entries[k]["section"] = section
                section = entries[h]["label"][:300]
                entries[h]["section"] = None
                for k in range(h + 1, j + 1):
                    entries[k]["section"] = section
            else:
                for k in range(i, j + 1):
                    entries[k]["section"] = section
            i = j + 1
            continue
        e["section"] = section
        i += 1
    for e in entries:
        e.pop("_gap")
    return entries


@dataclass
class ImportReport:
    files: int = 0
    cost_lines: int = 0
    skipped_files: list[str] = field(default_factory=list)
    editions: int = 0
    orders: int = 0
    warnings: int = 0
    # (title, year, edition, imported_total, sheet_total)
    checks: list[tuple[str, int, str, float, float | None]] = field(default_factory=list)


def _cell_with(row: ParsedRow, pattern: str) -> str | None:
    return next((str(v).strip() for v in row.raw if isinstance(v, str) and re.search(pattern, v, re.IGNORECASE)), None)


def _status_for(row: ParsedRow) -> tuple[str, str | None, str | None]:
    """(status, moved-to note, the sheet text that decided it). The text is
    kept so the app can show *why* - e.g. a "Size" cell reading "Judge
    ticket cancelled 4/9" that Excel clips to "Judge ticket" on screen."""
    moved_cell = _cell_with(row, r"moved to")
    if moved_cell:
        return "moved", moved_cell[:200], moved_cell
    cancel_cell = _cell_with(row, r"\bCANX\b|CANCELL?ED\b")
    if cancel_cell and not re.search(r"CANCELLED AGAINST", cancel_cell, re.IGNORECASE):
        return "cancelled", None, cancel_cell
    contra_cell = _cell_with(row, r"\bCONTRA\b")
    if contra_cell:
        return "contra", None, contra_cell
    return "booked", None, None


def _group_files(root: Path) -> list[tuple[int, Path, Path | None]]:
    """(year, file, working_copy_or_None) for every workbook worth importing."""
    out = []
    for year_dir in sorted(p for p in root.iterdir() if p.is_dir() and re.fullmatch(r"\d{4}", p.name)):
        files = sorted(p for p in year_dir.iterdir()
                       if p.suffix.lower() in (".xls", ".xlsx") and not p.name.startswith("."))
        files = [p for p in files if "copy(" not in p.name.lower()]
        planners = [p for p in files if title_for_file(p.name) and title_for_file(p.name).slug == "visit-usa-planner"]
        edition = next((p for p in planners if "edition" in p.name.lower()), None)
        working = next((p for p in planners if p is not edition), None) if edition else None
        for p in files:
            if p is working:
                continue
            out.append((int(year_dir.name), p, working if p is edition else None))
    return out


# Editorial-plan columns on an edition that a re-import must keep (see app/api/routes/editorial.py).
PLAN_FIELDS = ("editorial_deadline", "ad_deadline", "copy_deadline", "milestones", "theme", "distribution", "format",
               "plan_needs_check", "date_set_in_plan")


def import_sor(db: Session, root: Path, *, replace: bool = False) -> ImportReport:
    ensure_reference_data(db)
    # What people added in the app to imported editions - notes, online
    # link, cost lines typed in - is carried over to the re-imported edition
    # with the same title/year/name, so a refresh never loses it.
    kept: dict[tuple, dict] = {}
    if replace:
        for e in db.query(SalesEdition).filter(SalesEdition.source_file.isnot(None)):
            extra_costs = [
                {k: getattr(c, k) for k in ("kind", "label", "amount_gbp", "amount_inc_vat_gbp", "section", "sort_order")}
                for c in db.query(SalesEditionCost).filter(SalesEditionCost.edition_id == e.id, SalesEditionCost.source_row.is_(None))
            ]
            # ...and the editorial plan: deadlines, key dates, theme, planned features.
            plan = {k: getattr(e, k) for k in PLAN_FIELDS}
            features = [{k: getattr(f, k) for k in ("title", "description", "status", "sponsorable", "sort_order")}
                        for f in db.query(EditionFeature).filter(EditionFeature.edition_id == e.id)]
            has_plan = any(v not in (None, [], False) for v in plan.values()) or features
            if e.notes or e.digital_url or e.target_gbp or extra_costs or has_plan:
                kept[(e.title_id, e.year, e.name)] = {"notes": e.notes, "digital_url": e.digital_url,
                                                      "target_gbp": e.target_gbp, "costs": extra_costs,
                                                      "plan": plan, "features": features,
                                                      "edition_date": e.edition_date if e.date_set_in_plan else None}
        imported = db.query(SalesEdition.id).filter(SalesEdition.source_file.isnot(None))
        # Bookings made in the app (orders, typed-in bookings) on imported issues aren't in the spreadsheets, so a
        # refresh must not lose them: they wait on a placeholder issue and move to the re-imported one afterwards.
        held = _hold_app_bookings(db)
        db.query(SalesOrder).filter(SalesOrder.edition_id.in_(imported.scalar_subquery())).delete(synchronize_session=False)  # credits cascade
        db.query(SalesEdition).filter(SalesEdition.source_file.isnot(None)).delete(synchronize_session=False)
        db.flush()
    elif db.query(SalesEdition).filter(SalesEdition.source_file.isnot(None)).first():
        raise RuntimeError("SOR workbooks were already imported - rerun with replace=True to reimport.")
    else:
        held = {}

    titles = {t.slug: t for t in db.query(SalesTitle).all()}
    reps = {r.code: r for r in db.query(SalesRep).all()}
    report = ImportReport()

    for year, path, working_copy in _group_files(root):
        tdef = title_for_file(path.name)
        if not tdef:
            report.skipped_files.append(f"{year}/{path.name} (no matching title)")
            continue
        title = titles[tdef.slug]
        try:
            book = read_workbook(path)
            hidden = hidden_columns(path)
        except Exception as exc:  # noqa: BLE001 - one unreadable file shouldn't stop the rest
            report.skipped_files.append(f"{year}/{path.name} ({exc})")
            continue
        report.files += 1
        working = read_workbook(working_copy) if working_copy else {}

        for sheet_name, rows in book.items():
            sheet = parse_sheet(sheet_name, rows, hidden.get(sheet_name))
            if not sheet:
                continue
            wsheet = parse_sheet(sheet_name, working.get(sheet_name, [])) if sheet_name in working else None
            wrows = {(r.row_index, r.client): r for r in (wsheet.rows if wsheet else [])}

            name = sheet.name
            existing = db.query(SalesEdition).filter_by(title_id=title.id, year=year, name=name).one_or_none()
            if existing:  # two workbooks for one title+year with the same sheet name
                name = f"{name} ({path.stem})"[:120]
            ed = SalesEdition(
                id=uuid.uuid4(), title_id=title.id, year=year, name=name,
                period_label=(sheet.period.strftime("%B %Y") if isinstance(sheet.period, datetime) else
                              (str(sheet.period)[:200] if sheet.period is not None else None)),
                edition_date=edition_date_for(year, sheet.name, sheet.period),
                kind=edition_kind(tdef.product_line, sheet.name, sheet.period),
                status="closed" if year < date.today().year else "open",
                exchange_rate=sheet.exchange_rate,
                source_file=f"{year}/{path.name}", source_sheet=sheet_name,
                sheet_total_gbp=sheet.sheet_total,
            )
            db.add(ed)
            db.flush()
            report.editions += 1
            total = 0.0
            # Events sheets keep each event's direct costs / P&L under the
            # bookings - carried over so revenue and costs stay together.
            if tdef.product_line in ("events", "awards"):
                for n, line in enumerate(cost_lines(sheet.other_rows, sheet.gbp_col)):
                    db.add(SalesEditionCost(id=uuid.uuid4(), edition_id=ed.id, sort_order=n, **line))
                    report.cost_lines += 1
            if (k := kept.pop((title.id, year, name), None)):
                ed.notes, ed.digital_url, ed.target_gbp = k["notes"], k["digital_url"], k["target_gbp"]
                for f, v in (k.get("plan") or {}).items():
                    setattr(ed, f, v)
                if k.get("edition_date"):
                    ed.edition_date = k["edition_date"]
                for feat in k.get("features") or []:
                    db.add(EditionFeature(id=uuid.uuid4(), edition_id=ed.id, **feat))
                for c in k["costs"]:
                    db.add(SalesEditionCost(id=uuid.uuid4(), edition_id=ed.id, **{**c, "sort_order": c["sort_order"] + 10_000}))

            for row in sheet.rows:
                warnings = []
                if row.continuation:
                    warnings.append("client cell was blank on the sheet - assumed to be the same client as the row above")
                c = row.cells
                value, w = parse_money(c.get("gbp")); warnings += [w] if w else []
                usd, _ = parse_money(c.get("usd"))
                inv_value, _ = parse_money(c.get("invoice_value"))
                raw_inv_value = c.get("invoice_value")
                # Text typed into the invoice-value cell ("part of overpaid
                # credit of 4K rest on People Awards") is an explanation, not
                # an amount - kept as the reason, amount left unknown.
                inv_value_text = raw_inv_value.strip() if isinstance(raw_inv_value, str) and inv_value is None and raw_inv_value.strip() else None
                booked_on, w = parse_date(c.get("date")); warnings += [w] if w else []
                status, moved_note, status_reason = _status_for(row)

                raw_rep = c.get("rep")
                raw_rep = None if raw_rep is None or isinstance(raw_rep, float) else str(raw_rep).strip()
                parts = [p for p in re.split(r"[/&,+]", raw_rep or "") if p.strip()]
                codes = [rep_code_for(p) for p in parts]
                if raw_rep and (not codes or None in codes) and raw_rep.upper() not in ("TBC", "CONTRA"):
                    warnings.append(f"unknown salesperson {raw_rep!r}")
                codes = [x for x in codes if x]
                # Credits: the sheet's own commission columns when it filled
                # them in (that's how a shared "SP/ST" booking gets split),
                # else the full value to the salesperson(s) named.
                credits = dict(row.rep_credits)
                # Credits adding up to more than the booking, spread over
                # reps the Salesper. column doesn't name, are leftovers in
                # the sheet - keep only the named rep(s).
                if codes and value and sum(credits.values()) > value + 1:
                    named = {k: v for k, v in credits.items() if k in codes}
                    if named:
                        credits = named
                if not codes and credits:
                    codes = list(credits)
                if not credits and codes and value:
                    credits = {k: round(value / len(codes), 2) for k in codes}
                credited = sum(credits.values())
                if (value and credits and status == "booked" and abs(credited - value) > 1
                        and not (inv_value is not None and abs(credited - inv_value) <= 1)):
                    # Crediting the invoiced amount (e.g. part-invoiced, rest
                    # "next quarter") is a normal sheet pattern. A figure
                    # matching neither is a sheet error - usually the
                    # commission column shifted a row after rows were
                    # inserted/sorted (it holds the neighbouring booking's
                    # value), or a doubled figure. Credit the named rep(s)
                    # with the booking value instead, and say so.
                    fallback = codes or list(credits)
                    credits = {k: round(value / len(fallback), 2) for k in fallback}
                    warnings.append(f"the sheet's commission column showed £{credited:,.2f} for this £{value:,.2f} booking "
                                    f"(probably shifted from a neighbouring row) - credited the booking value instead")
                code = codes[0] if codes else None

                inv_no = c.get("invoice_number")
                if isinstance(inv_no, float):
                    inv_no = None if inv_no == 0 else f"{inv_no:.0f}"
                else:
                    inv_no = None if inv_no is None else str(inv_no).strip()
                order_ref = None
                if inv_no and re.fullmatch(r"\d{7,}", inv_no) and not value:
                    # An online ticket/order number on a £0 seat - not a BMI invoice.
                    order_ref, inv_no = inv_no, None
                if inv_no and not re.search(r"\d", inv_no):  # "Moved to OBH 107", "tbc" - a note, not an invoice
                    moved_note = moved_note or inv_no
                    inv_no = None
                reason = c.get("reason")
                reason = None if reason is None or isinstance(reason, float) else str(reason)
                if inv_value_text and inv_value_text.lower() not in ("value", "invoice value"):
                    reason = "; ".join(filter(None, [reason, inv_value_text]))
                notes = moved_note

                wrow = wrows.get((row.row_index, row.client))
                if wrow:
                    wreason = wrow.cells.get("reason")
                    if wreason is not None and not isinstance(wreason, float) and not reason:
                        reason = str(wreason)
                    winv, _ = parse_money(wrow.cells.get("invoice_value"))
                    if winv is not None and inv_value is not None and abs(winv - inv_value) > 0.5:
                        notes = "; ".join(filter(None, [notes, f"Working copy showed an invoice of £{winv:,.2f}"]))

                agency = None
                if reason and "agency" in reason.lower() and value and inv_value is not None and inv_value < value:
                    agency = round(value - inv_value, 2)

                size = fmt_size(c.get("size"))
                size = size[:120] if size else None
                num = lambda v: None if v is None else (f"{v:g}" if isinstance(v, float) else str(v))[:60]  # noqa: E731
                order_id = uuid.uuid4()
                db.add(SalesOrder(
                    id=order_id, edition_id=ed.id, client_name=row.client[:256],
                    rep_id=reps[code].id if code in reps else None,
                    booked_on=booked_on, size=size, pages=parse_pages(size, tdef.product_line),
                    series=num(c.get("series")), position=num(c.get("position")),
                    rate_usd=usd, value_gbp=value or 0, agency_commission_gbp=agency,
                    invoice_number=inv_no[:60] if inv_no else None, invoice_value_gbp=inv_value if inv_no else None,
                    invoice_note=reason, status=status, notes=notes,
                    source_file=f"{year}/{path.name}", source_sheet=sheet_name, source_row=row.row_index,
                    import_warning="; ".join(warnings) or None,
                    order_ref=order_ref, status_reason=status_reason, extra=row.extra,
                ))
                for rc, amount in credits.items():
                    if rc in reps and amount:
                        db.add(SalesOrderCredit(id=uuid.uuid4(), order_id=order_id, rep_id=reps[rc].id, amount_gbp=amount))
                report.orders += 1
                report.warnings += bool(warnings)
                total += value or 0
            report.checks.append((title.name, year, name, round(total, 2), sheet.sheet_total))
        db.flush()

    _link_moves(db)
    _restore_app_bookings(db, held)
    return report


def _hold_app_bookings(db: Session) -> dict:
    """Moves bookings made in the app off imported issues onto placeholder issues; returns {placeholder id: original key}."""
    rows = db.query(SalesOrder.edition_id).join(SalesEdition, SalesOrder.edition_id == SalesEdition.id).filter(
        SalesEdition.source_file.isnot(None), SalesOrder.source_file.is_(None)).distinct().all()
    held: dict = {}
    for (eid,) in rows:
        e = db.get(SalesEdition, eid)
        ph = SalesEdition(id=uuid.uuid4(), title_id=e.title_id, year=e.year, name=f"{e.name} [kept {uuid.uuid4().hex[:6]}]"[:120],
                          edition_date=e.edition_date, kind=e.kind, status=e.status)
        db.add(ph)
        db.flush()
        db.query(SalesOrder).filter(SalesOrder.edition_id == eid, SalesOrder.source_file.is_(None)).update(
            {SalesOrder.edition_id: ph.id}, synchronize_session=False)
        held[ph.id] = (e.title_id, e.year, e.name)
    db.flush()
    return held


def _restore_app_bookings(db: Session, held: dict) -> None:
    """Puts held bookings on the re-imported issue with the same title, year and name (or gives the placeholder its name
    back when the spreadsheets no longer have that issue). A booking the spreadsheet now has too is flagged as a possible double."""
    from app.sales.matching import normalise

    for ph_id, (title_id, year, name) in held.items():
        ph = db.get(SalesEdition, ph_id)
        new = db.query(SalesEdition).filter_by(title_id=title_id, year=year, name=name).one_or_none()
        if not new:
            ph.name = name
            continue
        sheet = [(normalise(o.client_name), round(float(o.value_gbp or 0), 2), o.source_row)
                 for o in db.query(SalesOrder).filter(SalesOrder.edition_id == new.id, SalesOrder.source_file.isnot(None))]
        for o in db.query(SalesOrder).filter(SalesOrder.edition_id == ph.id).all():
            o.edition_id = new.id
            twin = next((r for c, v, r in sheet if c == normalise(o.client_name) and v == round(float(o.value_gbp or 0), 2)), None)
            if twin is not None:
                o.import_warning = (f"The order register now has this booking too (row {twin}). Check it isn't counted twice - "
                                    "cancel one of them if it is.")
        db.flush()
        db.delete(ph)
    db.flush()


def _link_moves(db: Session) -> None:
    """Best effort: "Moved to OBH 107" -> link to that edition when the
    number/name is unambiguous within the same title and year."""
    moved = db.query(SalesOrder, SalesEdition).join(SalesEdition, SalesOrder.edition_id == SalesEdition.id).filter(
        SalesOrder.status == "moved", SalesOrder.moved_to_edition_id.is_(None)).all()
    for order, ed in moved:
        m = re.search(r"moved to\s+(.+)", (order.notes or "").lower())
        if not m:
            continue
        target = re.sub(r"^(obh|tbtm|connect)\s+", "", m[1].strip())
        candidates = db.query(SalesEdition).filter(
            SalesEdition.title_id == ed.title_id, SalesEdition.year.in_([ed.year, ed.year + 1]),
            SalesEdition.id != ed.id).all()
        hits = [c for c in candidates if c.name.lower().strip() == target or c.name.lower().endswith(target)]
        if len(hits) == 1:
            order.moved_to_edition_id = hits[0].id


def edition_window(ed: SalesEdition) -> tuple[date, date]:
    """(start, end) a reasonable 'this edition is live' window - used by
    invoice chasing to decide when an uninvoiced booking is overdue."""
    start = ed.edition_date or date(ed.year, 1, 1)
    return start, start + timedelta(days=45)
