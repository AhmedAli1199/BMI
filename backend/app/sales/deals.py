"""Multi-item orders (app/models/deal.py): pricing, keeping each item's booking in step,
the confirmation document, and rebooking next year.

Pricing, in the order a publisher thinks about it:
1. Each line has a price per insertion (the rate card's, or agreed), optionally its own % off,
   and runs in one or more issues / months / events ("placements"). Added-value lines are free.
2. Or the whole order is a package at one price, shared out across the items - by their rate
   card value (default), evenly, or by amounts typed per line - so each issue still gets its
   fair share of the revenue and commission is earned as each item runs.
3. A % off the whole order, then any agency's commission (taken off what the client pays,
   and off the salesperson's commissionable revenue, as the order register always did).
Every figure is rounded to the penny and the pennies are put on the biggest item, so the
items always add up exactly to the order total.
"""
from __future__ import annotations

import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, timezone

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (DealSettings, SalesDeal, SalesEdition, SalesOrder, SalesOrderCredit, SalesRep, SalesTitle)

ITEM_STATUS = {"pencilled": "pencilled", "confirmed": "booked", "cancelled": "cancelled"}


def _r(v: float) -> float:
    return round(v + 0.0, 2)


def _f(v, default: float = 0.0) -> float:
    try:
        return float(v) if v not in (None, "") else default
    except (TypeError, ValueError):
        return default


class OrderError(ValueError):
    """A plain-English reason the order can't be saved as it is."""


@dataclass
class Priced:
    lines: list[dict]                   # the lines, each with "placements" carrying value_gbp / agency_gbp
    gross_gbp: float = 0.0              # before the order discount (items mode: sum of lines; package: the package price)
    discount_gbp: float = 0.0
    total_gbp: float = 0.0              # what's agreed, before VAT
    agency_gbp: float = 0.0
    payable_gbp: float = 0.0            # what the client (or agency) is invoiced, before VAT
    rate_card_gbp: float = 0.0          # everything at rate card prices, free items included
    added_value_gbp: float = 0.0        # rate card value of what's given free
    warnings: list[str] = field(default_factory=list)

    @property
    def off_rate_card_pct(self) -> float | None:
        return _r(100 * (1 - self.total_gbp / self.rate_card_gbp)) if self.rate_card_gbp > 0 else None


def _spread(target: float, weights: list[float]) -> list[float]:
    """Shares `target` in proportion to weights, to the penny, adding up exactly."""
    if not weights:
        return []
    tw = sum(weights)
    raw = [target * w / tw for w in weights] if tw > 0 else [target / len(weights)] * len(weights)
    out = [_r(x) for x in raw]
    diff = _r(target - sum(out))
    if diff:
        i = max(range(len(out)), key=lambda k: (out[k], -k))
        out[i] = _r(out[i] + diff)
    return out


def price(lines: list[dict], *, pricing: str = "items", package_price: float | None = None, package_split: str = "rate_card",
          discount_pct: float = 0.0, agency_pct: float = 0.0) -> Priced:
    """Works out every item's value. `lines` as stored on SalesDeal.lines (placements need no ids here)."""
    lines = [dict(ln, placements=[dict(p) for p in ln.get("placements") or []]) for ln in lines]
    pr = Priced(lines=lines)
    slots: list[tuple[dict, dict]] = []  # (line, placement) for every paid placement
    for ln in lines:
        qty = max(1, int(_f(ln.get("qty"), 1)))
        ln["qty"] = qty
        unit, lst = _f(ln.get("unit_price")), _f(ln.get("list_price"))
        ld = min(max(_f(ln.get("discount_pct")), 0.0), 1.0)
        if not ln["placements"]:
            raise OrderError(f"“{ln.get('description') or 'A line'}” needs at least one issue, month or event.")
        for p in ln["placements"]:
            pr.rate_card_gbp += (lst or unit) * qty
            if ln.get("added_value"):
                pr.added_value_gbp += (lst or unit) * qty
                p["gross_gbp"] = 0.0
                continue
            p["gross_gbp"] = unit * qty * (1 - ld)
            slots.append((ln, p))
    if pricing == "package":
        if package_price is None or package_price < 0:
            raise OrderError("Enter the package price.")
        if not slots:
            raise OrderError("A package needs at least one item that isn't free.")
        if package_split == "manual":
            weights_by_line: dict[int, float] = {}
            for ln in lines:
                if not ln.get("added_value"):
                    weights_by_line[id(ln)] = _f(ln.get("share_gbp"))
            typed = sum(weights_by_line.values())
            if abs(typed - package_price) > 0.005:
                pr.warnings.append(f"The amounts typed per line add up to £{typed:,.2f}, not the package price of £{package_price:,.2f}, "
                                   "so they've been scaled to fit.")
            weights = []
            for ln, _p in slots:
                n = sum(1 for _ln, _ in slots if _ln is ln)
                weights.append(weights_by_line[id(ln)] / n)
        elif package_split == "even":
            weights = [1.0] * len(slots)
        else:  # rate_card: what each item would cost at rate card (or its typed price), so dearer items carry more
            weights = [(_f(ln.get("list_price")) or _f(ln.get("unit_price"))) * ln["qty"] for ln, _p in slots]
            if not any(weights):
                weights = [1.0] * len(slots)
                pr.warnings.append("No rate card prices to share the package by, so it's shared evenly.")
        gross = [*_spread(package_price, weights)]
        for (ln, p), g in zip(slots, gross):
            p["gross_gbp"] = g
        pr.gross_gbp = _r(package_price)
    else:
        pr.gross_gbp = _r(sum(p["gross_gbp"] for _ln, p in slots))
    od = min(max(discount_pct, 0.0), 1.0)
    pr.total_gbp = _r(pr.gross_gbp * (1 - od))
    pr.discount_gbp = _r(pr.gross_gbp - pr.total_gbp)
    values = _spread(pr.total_gbp, [p["gross_gbp"] for _ln, p in slots]) if slots else []
    for (_ln, p), v in zip(slots, values):
        p["value_gbp"] = v
    pr.agency_gbp = _r(pr.total_gbp * min(max(agency_pct, 0.0), 1.0))
    agency = _spread(pr.agency_gbp, values) if slots and pr.agency_gbp else [0.0] * len(slots)
    for (_ln, p), a in zip(slots, agency):
        p["agency_gbp"] = a
    for ln in lines:
        for p in ln["placements"]:
            p.setdefault("value_gbp", 0.0)
            p.setdefault("agency_gbp", 0.0)
            p["gross_gbp"] = _r(p["gross_gbp"])
        ln["total_gbp"] = _r(sum(p["value_gbp"] for p in ln["placements"]))
    pr.payable_gbp = _r(pr.total_gbp - pr.agency_gbp)
    pr.rate_card_gbp, pr.added_value_gbp = _r(pr.rate_card_gbp), _r(pr.added_value_gbp)
    return pr


def price_deal(deal: SalesDeal, lines: list[dict] | None = None) -> Priced:
    return price(lines if lines is not None else deal.lines or [], pricing=deal.pricing, package_price=_f(deal.package_price_gbp, None) if deal.package_price_gbp is not None else None,
                 package_split=deal.package_split, discount_pct=_f(deal.discount_pct), agency_pct=_f(deal.agency_pct))


def settings(db: Session) -> DealSettings:
    st = db.get(DealSettings, 1)
    if not st:
        st = DealSettings(id=1, next_number=2437, company_block="", footer="", terms="", artwork_specs="")
        db.add(st)
        db.flush()
    return st


def next_number(db: Session) -> int:
    st = settings(db)
    from sqlalchemy import func
    highest = db.scalar(select(func.max(SalesDeal.number))) or 0
    n = max(int(st.next_number), highest + 1)
    st.next_number = n + 1
    return n


def _credits(db: Session, deal: SalesDeal, order: SalesOrder) -> None:
    for c in db.scalars(select(SalesOrderCredit).where(SalesOrderCredit.order_id == order.id)):
        db.delete(c)
    db.flush()
    value = _f(order.value_gbp)
    if value <= 0:
        return
    split = [s for s in (deal.split or []) if _f(s.get("pct")) > 0]
    if split:
        amounts = _spread(value, [_f(s["pct"]) for s in split])
        merged: dict[str, float] = {}
        for s, a in zip(split, amounts):
            merged[s["rep_id"]] = merged.get(s["rep_id"], 0) + a
        for rep_id, a in merged.items():
            db.add(SalesOrderCredit(id=uuid.uuid4(), order_id=order.id, rep_id=uuid.UUID(str(rep_id)), amount_gbp=_r(a)))
    elif deal.rep_id:
        db.add(SalesOrderCredit(id=uuid.uuid4(), order_id=order.id, rep_id=deal.rep_id, amount_gbp=value))


def sync_items(db: Session, deal: SalesDeal, priced: Priced, user_id: uuid.UUID | None, *, today: date | None = None) -> list[str]:
    """Makes the order's bookings match its lines: one booking per placement. Returns notes for the person
    (e.g. an invoiced item that was taken off the order is cancelled, not deleted)."""
    from app.sales.sor_import import parse_pages

    today = today or date.today()
    notes: list[str] = []
    existing = {o.id: o for o in db.scalars(select(SalesOrder).where(SalesOrder.deal_id == deal.id))}
    titles = {t.id: t for t in db.scalars(select(SalesTitle))}
    keep: set[uuid.UUID] = set()
    status = ITEM_STATUS[deal.status]
    for ln in priced.lines:
        ln.setdefault("key", uuid.uuid4().hex[:12])
        for p in ln["placements"]:
            ed = db.get(SalesEdition, uuid.UUID(str(p["edition_id"]))) if p.get("edition_id") else None
            if not ed:
                raise OrderError(f"Choose the issue, month or event for every item of “{ln.get('description') or 'a line'}”.")
            oid = uuid.UUID(str(p["order_id"])) if p.get("order_id") else None
            o = existing.get(oid) if oid else None
            if not o:
                o = SalesOrder(id=uuid.uuid4(), deal_id=deal.id, created_by_user_id=user_id, extra={})
                # Added to an order that already existed: booked today (matters for "everything on the first invoice").
                p["booked_on"] = p.get("booked_on") or (deal.booked_on if not existing else today).isoformat()
                db.add(o)
            p["order_id"] = str(o.id)
            keep.add(o.id)
            size = (ln.get("size") or ln.get("description") or "").strip()[:120] or None
            o.edition_id, o.client_name, o.company_id, o.rep_id = ed.id, deal.client_name, deal.company_id, deal.rep_id
            o.booked_on = date.fromisoformat(p["booked_on"]) if p.get("booked_on") else deal.booked_on
            o.size, o.description, o.deal_line = size, (ln.get("description") or "").strip()[:300] or None, ln["key"]
            o.quantity, o.unit_price_gbp = ln["qty"], _f(ln.get("unit_price"), None) if ln.get("unit_price") not in (None, "") else None
            o.list_price_gbp = _f(ln.get("list_price"), None) if ln.get("list_price") not in (None, "") else None
            o.added_value = bool(ln.get("added_value"))
            o.value_gbp, o.agency_commission_gbp = p["value_gbp"], (p["agency_gbp"] or None)
            o.item_date = date.fromisoformat(p["item_date"]) if p.get("item_date") else None
            o.copy_due = date.fromisoformat(p["copy_due"]) if p.get("copy_due") else None
            o.notes = (p.get("note") or "").strip() or None
            # A confirmed order doesn't un-cancel an item someone cancelled on its own; everything else follows the order.
            if not (o.status == "cancelled" and status == "booked" and o.id in existing and p.get("keep_cancelled")):
                o.status = status
            o.status_reason = deal.cancelled_reason if status == "cancelled" else None
            t = titles.get(ed.title_id)
            o.pages = parse_pages(size, t.product_line) if t else None
            o.updated_at = datetime.now(timezone.utc)
            db.flush()
            _credits(db, deal, o)
    for oid, o in existing.items():
        if oid in keep:
            continue
        if o.invoice_number or o.xero_invoice_id:
            o.status, o.status_reason = "cancelled", f"Taken off order {deal.number}"
            notes.append(f"{o.description or o.size or 'An item'} was already invoiced, so it's marked cancelled rather than deleted.")
        else:
            db.delete(o)
    db.flush()
    return notes


def insertions_text(db: Session, deal: SalesDeal) -> str:
    """'April, June, September, November' - the months (or issues) the order runs in, in date order."""
    if deal.insertions_label:
        return deal.insertions_label
    seen: list[tuple[date, str]] = []
    for ln in deal.lines or []:
        for p in ln.get("placements") or []:
            ed = db.get(SalesEdition, uuid.UUID(str(p["edition_id"]))) if p.get("edition_id") else None
            d = date.fromisoformat(p["item_date"]) if p.get("item_date") else (ed.edition_date if ed else None)
            if d:
                seen.append((d, d.strftime("%B")))
    seen.sort()
    out: list[str] = []
    for _, m in seen:
        if m not in out:
            out.append(m)
    years = {d.year for d, _ in seen}
    text = ", ".join(out)
    return f"{text} {min(years)}" if len(years) == 1 and text else text


def _day(d: date | None) -> str:
    if not d:
        return ""
    n = d.day
    suffix = "th" if 11 <= n % 100 <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix} {d:%B}"


def document(db: Session, deal: SalesDeal) -> dict:
    """Everything the printed confirmation shows, in BMI's layout (same data for the print page and the Word file)."""
    from app.sales.editorial import issue_label

    st = settings(db)
    pr = price_deal(deal)
    from app.models import User
    rep = db.get(SalesRep, deal.rep_id) if deal.rep_id else None
    rep_user = db.get(User, rep.user_id) if rep and rep.user_id else None
    title = db.get(SalesTitle, deal.title_id) if deal.title_id else None
    rows: list[dict] = []
    for ln in pr.lines:
        when = []
        for p in ln["placements"]:
            ed = db.get(SalesEdition, uuid.UUID(str(p["edition_id"]))) if p.get("edition_id") else None
            if p.get("item_date"):
                when.append(_day(date.fromisoformat(p["item_date"])))
            elif ed:
                when.append(issue_label(ed))
        n = len(ln["placements"])
        each = _f(ln.get("unit_price")) * ln["qty"] * (1 - _f(ln.get("discount_pct")))
        lst = _f(ln.get("list_price"))
        row = {"description": (ln.get("description") or ln.get("size") or "").strip(), "when": when, "insertions": n,
               "detail": (ln.get("detail") or "").strip() or None, "added_value": bool(ln.get("added_value")), "qty": ln["qty"]}
        if ln.get("added_value"):
            normally = f" - normally £{lst * ln['qty']:,.2f} each" if lst else ""
            row.update(each=None, total=None, note=f"Added value{normally}")
        elif deal.pricing == "package":
            row.update(each=None, total=None, note=None)
        else:
            row.update(each=_r(each), total=_r(each * n),
                       note=f"usual rate £{lst * ln['qty']:,.2f}" if lst and lst * ln["qty"] - each > 0.005 else None)
        rows.append(row)
    # "Copy due: 18th May, 18th August (3/4 page column); 1st June (News story)" - one entry per item, its dates in order
    copy_by_item: dict[str, list[str]] = {}
    for ln in pr.lines:
        for p in ln["placements"]:
            if p.get("copy_due"):
                copy_by_item.setdefault((ln.get("description") or ln.get("size") or "").strip(), []).append(p["copy_due"])
    copy_dates = [(sorted(ds)[0], item, ", ".join(_day(date.fromisoformat(d)) for d in sorted(set(ds)))) for item, ds in copy_by_item.items()]
    copy_dates.sort()
    return {
        "number": deal.number, "status": deal.status, "document": deal.document,
        "heading": "SCHEDULE OF WORKS" if deal.document == "schedule" else "ORDER CONFIRMATION",
        "subheading": None if deal.document == "schedule" else "(Invoice to follow)",
        "date": (deal.confirmed_at.date() if deal.confirmed_at else date.today()).isoformat(),
        "company_block": st.company_block, "footer": st.footer, "terms": st.terms,
        "artwork_specs": st.artwork_specs if deal.show_artwork_specs else None,
        "confirmation_address": deal.confirmation_address or deal.client_name, "invoice_to": deal.invoice_to or deal.confirmation_address or deal.client_name,
        "your_contact": deal.contact_name, "your_email": deal.contact_email,
        "our_contact": rep.name if rep else None, "our_email": rep_user.email if rep_user else None,
        "order_ref": deal.po_number, "invoice_email": deal.invoice_email,
        "intro": "We are pleased to confirm the following booking:" if deal.document == "schedule" else "We are pleased to confirm the following advertisement booking:",
        "publication": deal.publication_label or (title.name if title else None),
        "insertions": insertions_text(db, deal),
        "package": {"label": deal.package_label or f"Package for {deal.client_name}", "price": _r(_f(deal.package_price_gbp))} if deal.pricing == "package" else None,
        "rows": rows, "gross": pr.gross_gbp, "discount_pct": _f(deal.discount_pct), "discount": pr.discount_gbp, "total": pr.total_gbp,
        "agency_pct": _f(deal.agency_pct), "agency": pr.agency_gbp, "agency_name": deal.agency_name, "payable": pr.payable_gbp,
        "special_instructions": deal.special_instructions, "copy_instructions": deal.copy_instructions,
        "copy_dates": [{"date": d, "item": i, "dates": txt} for d, i, txt in copy_dates], "production_contact": deal.production_contact,
    }


def build_docx(doc: dict) -> bytes:
    """The confirmation as a Word file, laid out like BMI's own."""
    import io

    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    def money(v):
        return f"£{v:,.2f}"

    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(9.5)
    head = d.add_table(rows=1, cols=2)
    left, right = head.rows[0].cells
    p = left.paragraphs[0]
    r = p.add_run(doc["heading"])
    r.bold, r.font.size = True, Pt(18)
    if doc["subheading"]:
        left.add_paragraph().add_run(doc["subheading"]).bold = True
    left.add_paragraph(f"Date.   {date.fromisoformat(doc['date']):%d %B %Y}")
    lines = (doc["company_block"] or "").splitlines()
    rp = right.paragraphs[0]
    if lines:
        rp.add_run(lines[0]).bold = True
        for ln in lines[1:]:
            right.add_paragraph(ln)
    d.add_paragraph()
    addr = d.add_table(rows=1, cols=2)
    addr.style = "Table Grid"
    for cell, label, text in ((addr.rows[0].cells[0], "Confirmation address", doc["confirmation_address"]),
                              (addr.rows[0].cells[1], "Invoice to:", doc["invoice_to"])):
        cell.paragraphs[0].add_run(label).underline = True
        for ln in (text or "").splitlines():
            cell.add_paragraph().add_run(ln).bold = True
    d.add_paragraph()
    info = d.add_table(rows=2, cols=4)
    info.style = "Table Grid"
    vals = [("Your contact.", doc["your_contact"]), ("Email.", doc["your_email"]), ("Our contact.", doc["our_contact"]), ("Order no.", f"{doc['number']}" + (f"  (your ref {doc['order_ref']})" if doc["order_ref"] else ""))]
    for i, (label, val) in enumerate(vals):
        row = info.rows[i // 2].cells
        row[(i % 2) * 2].text = label
        row[(i % 2) * 2 + 1].text = val or ""
    d.add_paragraph()
    d.add_paragraph(doc["intro"])
    if doc["publication"]:
        p = d.add_paragraph("Publication   ")
        p.add_run(doc["publication"]).bold = True
    box = d.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    box.rows[0].cells[0].paragraphs[0].add_run("Insertions booked").underline = True
    box.rows[0].cells[0].add_paragraph().add_run(doc["insertions"] or "").bold = True
    d.add_paragraph()
    t = d.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in zip(t.rows[0].cells, ("Description", "Cost per insertion", "Insertions booked", "Total amount")):
        c.text = h
    if doc["package"]:
        c = t.add_row().cells
        c[0].paragraphs[0].add_run(doc["package"]["label"]).bold = True
        c[1].text, c[2].text, c[3].text = money(doc["package"]["price"]), "1", money(doc["package"]["price"])
    for row in doc["rows"]:
        c = t.add_row().cells
        text = row["description"] + (f" - {', '.join(row['when'])}" if row["when"] and (row["added_value"] or doc["package"] or row["insertions"] == 1) else "")
        c[0].text = text
        if row["detail"]:
            c[0].add_paragraph(row["detail"])
        if row["note"]:
            c[0].add_paragraph().add_run(row["note"]).italic = True
        if row["total"] is not None:
            c[1].text, c[2].text, c[3].text = money(row["each"]), str(row["insertions"]), money(row["total"])
            if row["insertions"] > 1 and row["when"]:
                c[0].add_paragraph(", ".join(row["when"]))
    if doc["discount"]:
        c = t.add_row().cells
        c[0].text = f"Less discount {doc['discount_pct'] * 100:g}%"
        c[3].text = f"({money(doc['discount'])})"
    if doc["agency"]:
        c = t.add_row().cells
        c[0].paragraphs[0].add_run(f"Less agency commission {doc['agency_pct'] * 100:g}%").bold = True
        run = c[3].paragraphs[0].add_run(f"({money(doc['agency'])})")
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    if doc["discount"] or doc["agency"] or doc["package"] or len(doc["rows"]) > 1:
        c = t.add_row().cells
        c[0].paragraphs[0].add_run("Total").bold = True
        c[3].paragraphs[0].add_run(money(doc["payable"])).bold = True
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("VAT is not included")
    run.bold, run.underline = True, True
    extra = [x for x in (doc["special_instructions"], doc["copy_instructions"],
                         f"Please email the invoice to {doc['invoice_email']}" if doc["invoice_email"] else None) if x]
    if doc["copy_dates"]:
        extra.append("Copy due: " + "; ".join(f"{c['dates']} ({c['item']})" if len(doc["copy_dates"]) > 1 else c["dates"] for c in doc["copy_dates"]))
    box = d.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    box.rows[0].cells[0].paragraphs[0].add_run("Special Instructions").underline = True
    for x in extra:
        for ln in x.splitlines():
            box.rows[0].cells[0].add_paragraph(ln)
    if doc["production_contact"]:
        r = d.add_paragraph().add_run(f"Production contact {doc['production_contact']}")
        r.bold, r.font.color.rgb = True, RGBColor(0xC0, 0x00, 0x00)
    if doc["artwork_specs"]:
        sp = d.add_paragraph(doc["artwork_specs"])
        sp.runs[0].font.size = Pt(7.5)
    if doc["terms"]:
        d.add_paragraph(doc["terms"]).runs[0].font.size = Pt(8)
    for ln in (doc["footer"] or "").splitlines():
        fp = d.add_paragraph(ln)
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.runs[0].font.size = Pt(8)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def rebook_lines(db: Session, deal: SalesDeal, target_year: int, today: date | None = None) -> tuple[list[dict], list[str]]:
    """The order's lines moved to next year's matching issues (same kind, nearest date a year on), at next
    year's rate card price where the product is on it. Notes say what couldn't be matched."""
    from app.models import SalesRate
    from app.sales.editorial import shift_year

    from sqlalchemy import func
    notes: list[str] = []
    out: list[dict] = []
    for ln in deal.lines or []:
        new = {k: v for k, v in ln.items() if k not in ("placements", "total_gbp")}
        new["key"] = uuid.uuid4().hex[:12]
        if ln.get("rate_id"):
            old = db.get(SalesRate, uuid.UUID(str(ln["rate_id"])))
            nxt = db.scalars(select(SalesRate).where(SalesRate.title_id == old.title_id, SalesRate.year == target_year,
                                                     SalesRate.product == old.product, SalesRate.archived.is_(False))).first() if old else None
            if nxt and nxt.price_gbp is not None:
                if _f(ln.get("list_price")) and abs(_f(nxt.price_gbp) - _f(ln.get("list_price"))) > 0.005:
                    notes.append(f"{old.product}: rate card is £{_f(nxt.price_gbp):,.2f} in {target_year} (was £{_f(ln.get('list_price')):,.2f}).")
                new["rate_id"], new["list_price"] = str(nxt.id), _f(nxt.price_gbp)
                if _f(ln.get("unit_price")) == _f(ln.get("list_price")):
                    new["unit_price"] = _f(nxt.price_gbp)
        places = []
        for p in ln.get("placements") or []:
            ed = db.get(SalesEdition, uuid.UUID(str(p["edition_id"]))) if p.get("edition_id") else None
            if not ed:
                continue
            want = shift_year(ed.edition_date) if ed.edition_date else None
            q = select(SalesEdition).where(SalesEdition.title_id == ed.title_id, SalesEdition.year == target_year, SalesEdition.kind == ed.kind)
            nxt = db.scalars(q.order_by(func.abs(SalesEdition.edition_date - want)) if want else q).first()
            item_date = shift_year(date.fromisoformat(p["item_date"])) if p.get("item_date") else None
            if not nxt:
                notes.append(f"No {target_year} issue yet for {ed.name} - choose one (or plan next year in the editorial plan).")
            places.append({"edition_id": str(nxt.id) if nxt else None, "item_date": item_date.isoformat() if item_date else None,
                           "copy_due": None, "note": p.get("note")})
        new["placements"] = places
        out.append(new)
    return out, notes
