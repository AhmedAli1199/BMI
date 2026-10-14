"""Writes the proposal editor's HTML into a Word document, in the template's own styles.

Covers what the editor makes: paragraphs, headings, bold / italic / underline, text colour and highlight, links,
bullet and numbered lists, tables (including the ready-made blocks' tables) and pictures (uploaded ones are kept
in the HTML as data: images). Anything else is written as plain text, so nothing a person typed is lost.
"""
from __future__ import annotations

import base64
import io
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser

from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BLOCK_TAGS = {"p", "div", "h1", "h2", "h3", "h4", "ul", "ol", "li", "table", "blockquote"}
VOID = {"br", "img", "hr"}


@dataclass
class Node:
    tag: str
    attrs: dict = field(default_factory=dict)
    children: list = field(default_factory=list)   # Node | str


class _Parser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        n = Node(tag, {k: v or "" for k, v in attrs})
        self.stack[-1].children.append(n)
        if tag not in VOID:
            self.stack.append(n)

    def handle_startendtag(self, tag, attrs):
        self.stack[-1].children.append(Node(tag, {k: v or "" for k, v in attrs}))

    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return

    def handle_data(self, data):
        self.stack[-1].children.append(data)


def parse(html: str) -> Node:
    p = _Parser()
    p.feed(html or "")
    p.close()
    return p.root


def _style(attrs: dict) -> dict:
    out = {}
    for part in (attrs.get("style") or "").split(";"):
        if ":" in part:
            k, v = part.split(":", 1)
            out[k.strip().lower()] = v.strip().lower()
    return out


_HIGHLIGHTS = {"#fff59d": WD_COLOR_INDEX.YELLOW, "yellow": WD_COLOR_INDEX.YELLOW, "#c8f7c5": WD_COLOR_INDEX.BRIGHT_GREEN,
               "#cfe8ff": WD_COLOR_INDEX.TURQUOISE}


def _rgb(v: str) -> RGBColor | None:
    m = re.match(r"#([0-9a-f]{6})$", v or "")
    if m:
        return RGBColor.from_string(m.group(1).upper())
    m = re.match(r"rgb\((\d+),\s*(\d+),\s*(\d+)\)", v or "")
    return RGBColor(int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None


def _add_hyperlink(paragraph, url: str, text: str, fmt: dict) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E9E")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(u)
    if fmt.get("bold"):
        rpr.append(OxmlElement("w:b"))
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


class Writer:
    """Walks the parsed HTML and adds paragraphs / tables / pictures to `container` (a document or a table cell)."""

    def __init__(self, document, body_style: str, heading_style: str, bold_style: str):
        self.doc = document
        self.body_style, self.heading_style, self.bold_style = body_style, heading_style, bold_style

    # -- inline ----------------------------------------------------------------------------------------
    def inline(self, para, nodes, fmt: dict) -> None:
        for n in nodes:
            if isinstance(n, str):
                text = re.sub(r"\s+", " ", n)
                if text.strip() or (text == " " and para.runs):
                    self._run(para, text, fmt)
                continue
            f = dict(fmt)
            st = _style(n.attrs)
            if n.tag in ("b", "strong") or st.get("font-weight") in ("bold", "700", "600"):
                f["bold"] = True
            if n.tag in ("i", "em") or st.get("font-style") == "italic":
                f["italic"] = True
            if n.tag == "u" or "underline" in st.get("text-decoration", ""):
                f["underline"] = True
            if st.get("color"):
                f["color"] = st["color"]
            if st.get("background-color") or st.get("background"):
                f["highlight"] = st.get("background-color") or st.get("background")
            if n.tag == "br":
                para.add_run().add_break()
            elif n.tag == "a" and n.attrs.get("href", "").startswith(("http", "mailto:")):
                _add_hyperlink(para, n.attrs["href"], self._text(n) or n.attrs["href"], f)
            elif n.tag == "img":
                self._picture(para, n)
            else:
                self.inline(para, n.children, f)

    def _run(self, para, text, fmt):
        r = para.add_run(text)
        if fmt.get("bold"):
            r.bold = True
        if fmt.get("italic"):
            r.italic = True
        if fmt.get("underline"):
            r.underline = True
        if fmt.get("color") and (c := _rgb(fmt["color"])):
            r.font.color.rgb = c
        if fmt.get("highlight"):
            r.font.highlight_color = _HIGHLIGHTS.get(fmt["highlight"], WD_COLOR_INDEX.YELLOW)
        if fmt.get("size"):
            r.font.size = Pt(fmt["size"])

    def _text(self, n) -> str:
        return "".join(c if isinstance(c, str) else self._text(c) for c in n.children).strip()

    def _picture(self, para, n) -> None:
        src = n.attrs.get("src", "")
        m = re.match(r"data:image/(png|jpe?g|gif);base64,(.+)$", src, re.S)
        if not m:
            if src.startswith("http"):
                self._run(para, n.attrs.get("alt") or src, {"italic": True})
            return
        try:
            data = base64.b64decode(m.group(2))
            width = n.attrs.get("width")
            w = Cm(min(16.0, float(width) / 37.8)) if width and width.isdigit() else Cm(14)
            para.add_run().add_picture(io.BytesIO(data), width=w)
        except Exception:
            self._run(para, "[picture]", {"italic": True})

    # -- blocks ----------------------------------------------------------------------------------------
    def blocks(self, container, nodes, *, list_kind: str | None = None) -> None:
        loose: list = []

        def flush():
            if any((isinstance(x, str) and x.strip()) or not isinstance(x, str) for x in loose):
                p = container.add_paragraph(style=self.body_style)
                self.inline(p, loose, {})
            loose.clear()

        counter = 0
        for n in nodes:
            if isinstance(n, str) or n.tag not in BLOCK_TAGS:
                loose.append(n)
                continue
            flush()
            if n.tag in ("h1", "h2", "h3", "h4"):
                p = container.add_paragraph(style=self.heading_style)
                self.inline(p, n.children, {})
            elif n.tag in ("ul", "ol"):
                self.blocks(container, [c for c in n.children if not isinstance(c, str)], list_kind=n.tag)
            elif n.tag == "li":
                counter += 1
                p = container.add_paragraph(style=self.body_style)
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.first_line_indent = Cm(-0.4)
                p.add_run(f"{counter}. " if list_kind == "ol" else "•  ")
                inner_blocks = [c for c in n.children if not isinstance(c, str) and c.tag in ("ul", "ol")]
                self.inline(p, [c for c in n.children if c not in inner_blocks], {})
                for b in inner_blocks:
                    self.blocks(container, [c for c in b.children if not isinstance(c, str)], list_kind=b.tag)
            elif n.tag == "table":
                self.table(container, n)
            else:   # p, div, blockquote
                if any(not isinstance(c, str) and c.tag in BLOCK_TAGS for c in n.children):
                    self.blocks(container, n.children)
                else:
                    style = self.bold_style if _style(n.attrs).get("font-weight") == "bold" else self.body_style
                    p = container.add_paragraph(style=style)
                    self.inline(p, n.children, {})
        flush()

    def table(self, container, n: Node) -> None:
        rows = []

        def collect(x):
            for c in x.children:
                if isinstance(c, Node):
                    if c.tag == "tr":
                        rows.append(c)
                    elif c.tag in ("thead", "tbody", "tfoot"):
                        collect(c)
        collect(n)
        if not rows:
            return
        cells_per = [[c for c in r.children if isinstance(c, Node) and c.tag in ("td", "th")] for r in rows]
        width = max(sum(int(c.attrs.get("colspan") or 1) for c in cs) for cs in cells_per) or 1
        t = container.add_table(len(rows), width)
        try:
            t.style = "Table Grid"
        except Exception:
            pass
        for ri, cs in enumerate(cells_per):
            ci = 0
            for c in cs:
                span = int(c.attrs.get("colspan") or 1)
                cell = t.cell(ri, ci)
                if span > 1 and ci + span - 1 < width:
                    cell = cell.merge(t.cell(ri, ci + span - 1))
                cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
                bold = c.tag == "th" or _style(c.attrs).get("font-weight") == "bold"
                if any(isinstance(x, Node) and x.tag in BLOCK_TAGS for x in c.children):
                    self.blocks(cell, c.children)
                else:
                    p = cell.add_paragraph(style=self.body_style)
                    self.inline(p, c.children, {"bold": bold})
                    if _style(c.attrs).get("text-align") == "right":
                        from docx.enum.text import WD_ALIGN_PARAGRAPH
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if not cell.paragraphs:
                    cell.add_paragraph()
                ci += span
        if container is self.doc:
            self.doc.add_paragraph(style=self.body_style)   # a little space after the table


def write(document, html: str, *, body_style: str, heading_style: str, bold_style: str) -> None:
    Writer(document, body_style, heading_style, bold_style).blocks(document, parse(html).children)
