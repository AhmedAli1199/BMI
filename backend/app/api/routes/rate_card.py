"""The rate card, organised the way BMI's media packs are: brand -> section ->
products, plus each brand's offers and discounts, a price history with
"put back", and a guided "prepare next year's prices" step.

Who can change what: admins and data managers everything; a brand's publishers
their own brand (app/sales/brands.py). Everyone can read.
"""
from __future__ import annotations

import io
import math
import uuid
from datetime import date, datetime

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.identity import Identity, get_identity
from app.db.session import get_db
from app.models import FieldChange, RateOffer, SalesRate, SalesTitle, User
from app.sales import brands as B
from app.sales import rate_card_seed as seed
from app.services.field_audit import record_field_changes

router = APIRouter(prefix="/rate-card", tags=["rate card"])

PRICE_TYPES = ("fixed", "from", "poa")
UNITS = ("each", "month", "week", "year", "event", "entry")
OFFER_KINDS = ("volume", "series", "early_bird", "note")
UNIT_WORDS = {"each": "", "month": " per month", "week": " per week", "year": " per year", "event": " per event", "entry": " per entry"}


# ---- shapes -------------------------------------------------------------------------

class RateItem(BaseModel):
    id: uuid.UUID
    title_id: uuid.UUID
    title_name: str
    year: int
    section: str
    product: str
    price_type: str
    price_gbp: float | None
    unit: str
    specs: str | None
    aliases: list[str]
    notes: str | None
    valid_until: date | None
    sort_order: int
    needs_check: bool
    source: str | None
    archived: bool
    price_label: str
    updated_at: datetime


class OfferOut(BaseModel):
    id: uuid.UUID
    brand: str
    year: int
    kind: str
    label: str
    details: str | None
    rules: dict
    rate_ids: list[str]
    section: str | None
    valid_until: date | None
    needs_check: bool


class SectionOut(BaseModel):
    key: str
    label: str
    hint: str
    default_title_id: uuid.UUID | None
    items: list[RateItem]


class TitleRef(BaseModel):
    id: uuid.UUID
    slug: str
    name: str
    digital_page_url: str | None = None
    digital_issue_url: str | None = None


class BrandSummary(BaseModel):
    key: str
    name: str
    short: str
    website: str
    can_edit: bool
    products: int
    needs_check: int
    on_request: int
    offers: int
    last_updated: datetime | None
    seed_available: int
    sections: list[str]


class BrandPage(BaseModel):
    brand: BrandSummary
    year: int
    years: list[int]
    sections: list[SectionOut]
    offers: list[OfferOut]
    archived: list[RateItem]
    titles: list[TitleRef]
    has_next_year: bool


class Overview(BaseModel):
    year: int
    years: list[int]
    brands: list[BrandSummary]


def price_label(r: SalesRate) -> str:
    if r.price_type == "poa" or r.price_gbp is None:
        return "Price on request"
    v = float(r.price_gbp)
    money = f"£{v:,.0f}" if v == int(v) else f"£{v:,.2f}"
    return ("From " if r.price_type == "from" else "") + money + UNIT_WORDS.get(r.unit, "")


def _item(r: SalesRate, titles: dict[uuid.UUID, SalesTitle]) -> RateItem:
    t = titles.get(r.title_id)
    return RateItem(id=r.id, title_id=r.title_id, title_name=t.name if t else "", year=r.year, section=r.section, product=r.product,
                    price_type=r.price_type, price_gbp=float(r.price_gbp) if r.price_gbp is not None else None, unit=r.unit, specs=r.specs,
                    aliases=r.aliases or [], notes=r.notes, valid_until=r.valid_until, sort_order=r.sort_order, needs_check=r.needs_check,
                    source=r.source, archived=r.archived, price_label=price_label(r), updated_at=r.updated_at)


def _offer(o: RateOffer) -> OfferOut:
    return OfferOut(id=o.id, brand=o.brand, year=o.year, kind=o.kind, label=o.label, details=o.details, rules=o.rules or {},
                    rate_ids=[str(x) for x in (o.rate_ids or [])], section=o.section, valid_until=o.valid_until, needs_check=o.needs_check)


def _brand_or_404(key: str) -> B.Brand:
    b = B.get_brand(key)
    if not b:
        raise HTTPException(404, "Brand not found")
    return b


def _titles(db: Session, brand: B.Brand) -> dict[uuid.UUID, SalesTitle]:
    return {t.id: t for t in B.brand_titles(db, brand)}


def _require_edit(db: Session, identity: Identity, brand: B.Brand) -> None:
    if not B.can_edit_brand(db, identity, brand):
        raise HTTPException(403, f"Only admins, data managers and {brand.name}'s publishers can change this rate card.")


def _rates(db: Session, brand: B.Brand, year: int, archived: bool | None = False) -> list[SalesRate]:
    ids = list(_titles(db, brand))
    if not ids:
        return []
    q = select(SalesRate).where(SalesRate.title_id.in_(ids), SalesRate.year == year)
    if archived is not None:
        q = q.where(SalesRate.archived.is_(archived))
    return list(db.scalars(q.order_by(SalesRate.sort_order, SalesRate.product)))


def _years(db: Session) -> list[int]:
    this = date.today().year
    have = set(db.scalars(select(SalesRate.year).distinct()))
    return sorted(have | {this, this + 1}, reverse=True)


def _summary(db: Session, identity: Identity, brand: B.Brand, year: int) -> BrandSummary:
    rows = _rates(db, brand, year)
    offers = db.scalar(select(func.count()).select_from(RateOffer).where(RateOffer.brand == brand.key, RateOffer.year == year)) or 0
    return BrandSummary(
        key=brand.key, name=brand.name, short=brand.short, website=brand.website, can_edit=B.can_edit_brand(db, identity, brand),
        products=len(rows), needs_check=sum(1 for r in rows if r.needs_check), on_request=sum(1 for r in rows if r.price_type == "poa"),
        offers=offers, last_updated=max((r.updated_at for r in rows), default=None), seed_available=seed.seed_available(db, brand.key, year),
        sections=[s for s in B.SECTION_KEYS if any(r.section == s for r in rows)],
    )


# ---- reading ----------------------------------------------------------------------------

@router.get("", response_model=Overview)
def overview(year: int | None = None, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> Overview:
    y = year or date.today().year
    return Overview(year=y, years=_years(db), brands=[_summary(db, identity, b, y) for b in B.BRANDS])


@router.get("/brands/{key}", response_model=BrandPage)
def brand_page(key: str, year: int | None = None, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> BrandPage:
    brand = _brand_or_404(key)
    y = year or date.today().year
    titles = _titles(db, brand)
    rows = _rates(db, brand, y, archived=None)
    live = [r for r in rows if not r.archived]
    by_slug = {t.slug: t for t in titles.values()}
    sections = [SectionOut(key=s.key, label=s.label, hint=s.hint,
                           default_title_id=by_slug[brand.title_for_section(s.key)].id if brand.title_for_section(s.key) in by_slug else None,
                           items=[_item(r, titles) for r in live if r.section == s.key]) for s in B.SECTIONS]
    offers = db.scalars(select(RateOffer).where(RateOffer.brand == key, RateOffer.year == y).order_by(RateOffer.sort_order, RateOffer.created_at))
    next_ids = list(titles)
    has_next = bool(next_ids and db.scalar(select(SalesRate.id).where(SalesRate.title_id.in_(next_ids), SalesRate.year == y + 1).limit(1)))
    return BrandPage(
        brand=_summary(db, identity, brand, y), year=y, years=_years(db), sections=sections, offers=[_offer(o) for o in offers],
        archived=[_item(r, titles) for r in rows if r.archived],
        titles=[TitleRef(id=t.id, slug=t.slug, name=t.name, digital_page_url=t.digital_page_url, digital_issue_url=t.digital_issue_url) for t in titles.values()],
        has_next_year=has_next,
    )


# ---- editing prices ---------------------------------------------------------------------

class ItemIn(BaseModel):
    brand: str
    year: int = Field(ge=2000, le=2100)
    section: str
    title_id: uuid.UUID | None = None
    product: str = Field(min_length=1, max_length=120)
    price_type: str = "fixed"
    price_gbp: float | None = Field(default=None, ge=0)
    unit: str = "each"
    specs: str | None = Field(default=None, max_length=200)
    aliases: list[str] = []
    notes: str | None = Field(default=None, max_length=300)
    valid_until: date | None = None


class ItemPatch(BaseModel):
    section: str | None = None
    title_id: uuid.UUID | None = None
    product: str | None = Field(default=None, min_length=1, max_length=120)
    price_type: str | None = None
    price_gbp: float | None = Field(default=None, ge=0)
    unit: str | None = None
    specs: str | None = Field(default=None, max_length=200)
    aliases: list[str] | None = None
    notes: str | None = Field(default=None, max_length=300)
    valid_until: date | None = None
    needs_check: bool | None = None


AUDITED = ("product", "price_type", "price_gbp", "unit", "specs", "aliases", "notes", "valid_until", "section", "title_id", "archived")


def _clean_aliases(a: list[str]) -> list[str]:
    return list(dict.fromkeys(x.strip() for x in a if x and x.strip()))[:12]


def _validate(section: str, price_type: str, price: float | None, unit: str) -> None:
    if section not in B.SECTION_KEYS:
        raise HTTPException(422, "Choose a section.")
    if price_type not in PRICE_TYPES:
        raise HTTPException(422, "Choose how the price is quoted.")
    if unit not in UNITS:
        raise HTTPException(422, "Choose what the price is for.")
    if price_type != "poa" and price is None:
        raise HTTPException(422, "Type the price, or choose “Price on request”.")


def _rate_brand(db: Session, r: SalesRate) -> B.Brand:
    t = db.get(SalesTitle, r.title_id)
    b = B.brand_for_title_slug(t.slug) if t else None
    if not b:
        raise HTTPException(404, "Price not found")
    return b


def _snapshot(r: SalesRate) -> dict:
    return {k: (str(getattr(r, k)) if k == "title_id" else getattr(r, k)) for k in AUDITED}


def _dup_check(db: Session, title_id, year, section, product, exclude: uuid.UUID | None = None) -> None:
    q = select(SalesRate).where(SalesRate.title_id == title_id, SalesRate.year == year, SalesRate.section == section,
                                func.lower(SalesRate.product) == product.lower())
    if exclude:
        q = q.where(SalesRate.id != exclude)
    if db.scalars(q).first():
        raise HTTPException(409, f"“{product}” is already on this rate card - edit that one instead.")


@router.post("/items", response_model=RateItem, status_code=201)
def add_item(p: ItemIn, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> RateItem:
    brand = _brand_or_404(p.brand)
    _require_edit(db, identity, brand)
    _validate(p.section, p.price_type, p.price_gbp, p.unit)
    titles = _titles(db, brand)
    title_id = p.title_id or B.title_id_for_section(db, brand, p.section)
    if title_id not in titles:
        raise HTTPException(422, "That part of the business isn't in this brand.")
    product = p.product.strip()
    _dup_check(db, title_id, p.year, p.section, product)
    last = db.scalar(select(func.max(SalesRate.sort_order)).where(SalesRate.title_id.in_(list(titles)), SalesRate.year == p.year, SalesRate.section == p.section)) or 0
    r = SalesRate(id=uuid.uuid4(), title_id=title_id, year=p.year, section=p.section, product=product, price_type=p.price_type,
                  price_gbp=None if p.price_type == "poa" else p.price_gbp, unit=p.unit, specs=(p.specs or "").strip() or None,
                  aliases=_clean_aliases(p.aliases), notes=(p.notes or "").strip() or None, valid_until=p.valid_until, sort_order=last + 10)
    db.add(r)
    db.flush()
    record_field_changes(db, entity_type="sales_rate", entity_id=r.id, before={"created": None}, updates={"created": price_label(r)},
                         changed_by_user_id=identity.user_uuid)
    db.commit()
    return _item(r, titles)


@router.patch("/items/{rid}", response_model=RateItem)
def edit_item(rid: uuid.UUID, p: ItemPatch, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> RateItem:
    r = db.get(SalesRate, rid)
    if not r:
        raise HTTPException(404, "Price not found")
    brand = _rate_brand(db, r)
    _require_edit(db, identity, brand)
    titles = _titles(db, brand)
    before = _snapshot(r)
    data = p.model_dump(exclude_unset=True)
    if "title_id" in data and data["title_id"] not in titles:
        raise HTTPException(422, "That part of the business isn't in this brand.")
    if "product" in data:
        data["product"] = data["product"].strip()
    if "aliases" in data:
        data["aliases"] = _clean_aliases(data["aliases"] or [])
    for k in ("specs", "notes"):
        if k in data:
            data[k] = (data[k] or "").strip() or None
    for k, v in data.items():
        setattr(r, k, v)
    if r.price_type == "poa":
        r.price_gbp = None
    _validate(r.section, r.price_type, float(r.price_gbp) if r.price_gbp is not None else None, r.unit)
    _dup_check(db, r.title_id, r.year, r.section, r.product, exclude=r.id)
    if "needs_check" not in data and any(k in data for k in ("price_gbp", "price_type", "product")):
        r.needs_check = False  # someone just set it by hand - that counts as checked
    after = _snapshot(r)
    record_field_changes(db, entity_type="sales_rate", entity_id=r.id, before=before, updates=after, changed_by_user_id=identity.user_uuid)
    db.commit()
    return _item(r, titles)


@router.post("/items/{rid}/archive", response_model=RateItem)
def archive_item(rid: uuid.UUID, restore: bool = False, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> RateItem:
    """Takes a product off the rate card (it stays in "Removed" and can be put back)."""
    r = db.get(SalesRate, rid)
    if not r:
        raise HTTPException(404, "Price not found")
    brand = _rate_brand(db, r)
    _require_edit(db, identity, brand)
    record_field_changes(db, entity_type="sales_rate", entity_id=r.id, before={"archived": r.archived}, updates={"archived": not restore},
                         changed_by_user_id=identity.user_uuid)
    r.archived = not restore
    db.commit()
    return _item(r, _titles(db, brand))


class ConfirmIn(BaseModel):
    brand: str
    year: int


@router.post("/confirm-all")
def confirm_all(p: ConfirmIn, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    """"I've checked them all": clears every "please check" mark for a brand's year."""
    brand = _brand_or_404(p.brand)
    _require_edit(db, identity, brand)
    n = 0
    for r in _rates(db, brand, p.year):
        if r.needs_check:
            r.needs_check, n = False, n + 1
    for o in db.scalars(select(RateOffer).where(RateOffer.brand == brand.key, RateOffer.year == p.year, RateOffer.needs_check.is_(True))):
        o.needs_check, n = False, n + 1
    db.commit()
    return {"confirmed": n}


class ReorderIn(BaseModel):
    ids: list[uuid.UUID] = Field(min_length=1, max_length=200)


@router.post("/items/reorder")
def reorder(p: ReorderIn, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    rows = [db.get(SalesRate, i) for i in p.ids]
    if any(r is None for r in rows):
        raise HTTPException(404, "Price not found")
    brands = {_rate_brand(db, r).key for r in rows}
    if len(brands) != 1:
        raise HTTPException(422, "Reorder one brand at a time.")
    _require_edit(db, identity, B.get_brand(brands.pop()))
    for i, r in enumerate(rows):
        r.sort_order = (i + 1) * 10
    db.commit()
    return {"ok": True}


@router.post("/brands/{key}/load-media-pack")
def load_media_pack(key: str, year: int = Query(seed.SEED_YEAR), db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    """Loads the prices taken from the brand's media pack, all marked "please check"."""
    brand = _brand_or_404(key)
    _require_edit(db, identity, brand)
    n = seed.load_seed(db, key, year)
    db.commit()
    return {"added": n}


# ---- offers -----------------------------------------------------------------------------

class OfferIn(BaseModel):
    brand: str
    year: int
    kind: str = "note"
    label: str = Field(min_length=1, max_length=300)
    details: str | None = None
    rules: dict = {}
    rate_ids: list[str] = []
    section: str | None = None
    valid_until: date | None = None


def _validate_offer(p: OfferIn) -> None:
    if p.kind not in OFFER_KINDS:
        raise HTTPException(422, "Choose what kind of offer it is.")
    tiers = (p.rules or {}).get("tiers", [])
    if p.kind == "volume" and not all(isinstance(t.get("qty"), int) and t["qty"] >= 2 and 0 < float(t.get("discount_pct", 0)) < 100 for t in tiers):
        raise HTTPException(422, "Each discount needs a number of bookings (2 or more) and a percentage.")
    if p.kind == "series" and not all(isinstance(t.get("qty"), int) and t["qty"] >= 2 and (t.get("unit_price") or t.get("total")) for t in tiers):
        raise HTTPException(422, "Each step needs a number of bookings and a price.")
    if p.kind == "early_bird" and not p.valid_until:
        raise HTTPException(422, "Say when the early-bird price ends.")
    if p.section and p.section not in B.SECTION_KEYS:
        raise HTTPException(422, "Choose a section.")


@router.post("/offers", response_model=OfferOut, status_code=201)
def add_offer(p: OfferIn, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> OfferOut:
    brand = _brand_or_404(p.brand)
    _require_edit(db, identity, brand)
    _validate_offer(p)
    o = RateOffer(id=uuid.uuid4(), brand=brand.key, year=p.year, kind=p.kind, label=p.label.strip(), details=(p.details or "").strip() or None,
                  rules=p.rules or {}, rate_ids=p.rate_ids, section=p.section, valid_until=p.valid_until,
                  sort_order=(db.scalar(select(func.max(RateOffer.sort_order)).where(RateOffer.brand == brand.key, RateOffer.year == p.year)) or 0) + 10)
    db.add(o)
    db.flush()
    record_field_changes(db, entity_type="rate_offer", entity_id=o.id, before={"created": None}, updates={"created": o.label}, changed_by_user_id=identity.user_uuid)
    db.commit()
    return _offer(o)


@router.put("/offers/{oid}", response_model=OfferOut)
def edit_offer(oid: uuid.UUID, p: OfferIn, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> OfferOut:
    o = db.get(RateOffer, oid)
    if not o:
        raise HTTPException(404, "Offer not found")
    brand = _brand_or_404(o.brand)
    _require_edit(db, identity, brand)
    _validate_offer(p)
    before = {"label": o.label, "details": o.details, "rules": o.rules, "valid_until": o.valid_until}
    o.kind, o.label, o.details, o.rules, o.rate_ids, o.section, o.valid_until = (p.kind, p.label.strip(), (p.details or "").strip() or None,
                                                                               p.rules or {}, p.rate_ids, p.section, p.valid_until)
    o.needs_check = False
    record_field_changes(db, entity_type="rate_offer", entity_id=o.id, before=before,
                         updates={"label": o.label, "details": o.details, "rules": o.rules, "valid_until": o.valid_until}, changed_by_user_id=identity.user_uuid)
    db.commit()
    return _offer(o)


@router.delete("/offers/{oid}", status_code=204, response_model=None)
def delete_offer(oid: uuid.UUID, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> None:
    o = db.get(RateOffer, oid)
    if not o:
        return
    _require_edit(db, identity, _brand_or_404(o.brand))
    record_field_changes(db, entity_type="rate_offer", entity_id=o.id, before={"deleted": None}, updates={"deleted": o.label}, changed_by_user_id=identity.user_uuid)
    db.delete(o)
    db.commit()


# ---- next year ----------------------------------------------------------------------------

class NextYearIn(BaseModel):
    brand: str
    from_year: int
    raise_pct: float = Field(default=0, ge=-50, le=100)
    round_to: int = Field(default=5, ge=1, le=1000)
    overrides: dict[str, float | None] = {}


class NextYearRow(BaseModel):
    id: uuid.UUID
    section: str
    product: str
    old_label: str
    old: float | None
    new: float | None
    price_type: str
    unit: str


class NextYearPreview(BaseModel):
    to_year: int
    rows: list[NextYearRow]
    already_there: int
    offers: int


def _bump(v: float | None, pct: float, round_to: int) -> float | None:
    if v is None:
        return None
    raw = v * (1 + pct / 100)
    return float(max(round_to, math.floor(raw / round_to + 0.5) * round_to)) if pct else float(v)


@router.post("/next-year/preview", response_model=NextYearPreview)
def next_year_preview(p: NextYearIn, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> NextYearPreview:
    brand = _brand_or_404(p.brand)
    _require_edit(db, identity, brand)
    to = p.from_year + 1
    existing = {(r.title_id, r.section, r.product.lower()) for r in _rates(db, brand, to, archived=None)}
    rows, skipped = [], 0
    order = {k: i for i, k in enumerate(B.SECTION_KEYS)}
    for r in sorted(_rates(db, brand, p.from_year), key=lambda r: (order.get(r.section, 99), r.sort_order, r.product)):
        if (r.title_id, r.section, r.product.lower()) in existing:
            skipped += 1
            continue
        old = float(r.price_gbp) if r.price_gbp is not None else None
        rows.append(NextYearRow(id=r.id, section=r.section, product=r.product, old_label=price_label(r), old=old,
                                new=_bump(old, p.raise_pct, p.round_to), price_type=r.price_type, unit=r.unit))
    offers = db.scalar(select(func.count()).select_from(RateOffer).where(RateOffer.brand == brand.key, RateOffer.year == p.from_year)) or 0
    return NextYearPreview(to_year=to, rows=rows, already_there=skipped, offers=offers)


@router.post("/next-year")
def next_year_apply(p: NextYearIn, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    """Copies a brand's prices (and offers) into the next year, with the chosen rise and any prices typed over."""
    preview = next_year_preview(p, db, identity)
    to = preview.to_year
    copied = 0
    id_map: dict[str, str] = {}
    for row in preview.rows:
        src = db.get(SalesRate, row.id)
        new = p.overrides.get(str(row.id), row.new) if str(row.id) in p.overrides else row.new
        r = SalesRate(id=uuid.uuid4(), title_id=src.title_id, year=to, section=src.section, product=src.product, price_type=src.price_type,
                      price_gbp=None if src.price_type == "poa" else new, unit=src.unit, specs=src.specs, aliases=list(src.aliases or []),
                      notes=src.notes, valid_until=_shift(src.valid_until), sort_order=src.sort_order, needs_check=False,
                      source=f"Copied from {p.from_year}" + (f" (+{p.raise_pct:g}%)" if p.raise_pct else ""))
        db.add(r)
        id_map[str(src.id)] = str(r.id)
        copied += 1
    if not db.scalar(select(RateOffer.id).where(RateOffer.brand == p.brand, RateOffer.year == to).limit(1)):
        for o in db.scalars(select(RateOffer).where(RateOffer.brand == p.brand, RateOffer.year == p.from_year)):
            db.add(RateOffer(id=uuid.uuid4(), brand=o.brand, year=to, kind=o.kind, label=o.label, details=o.details, rules=o.rules,
                             rate_ids=[id_map[x] for x in (o.rate_ids or []) if x in id_map], section=o.section,
                             valid_until=_shift(o.valid_until), sort_order=o.sort_order, needs_check=True))
    db.commit()
    return {"copied": copied, "year": to}


def _shift(d: date | None) -> date | None:
    if not d:
        return None
    try:
        return d.replace(year=d.year + 1)
    except ValueError:  # 29 Feb
        return d.replace(year=d.year + 1, day=28)


# ---- history ------------------------------------------------------------------------------

class HistoryRow(BaseModel):
    id: uuid.UUID
    entity_type: str
    entity_id: uuid.UUID
    what: str
    field: str
    old_value: str | None
    new_value: str | None
    by: str | None
    at: datetime
    can_put_back: bool


FIELD_WORDS = {"price_gbp": "price", "price_type": "how it's priced", "unit": "what the price is for", "specs": "size / specs", "aliases": "also called",
               "notes": "notes", "valid_until": "valid until", "section": "section", "title_id": "booked under", "archived": "removed",
               "product": "name", "created": "added", "deleted": "deleted", "label": "offer", "details": "offer details", "rules": "offer steps"}


@router.get("/history", response_model=list[HistoryRow])
def history(brand: str, year: int | None = None, db: Session = Depends(get_db)) -> list[HistoryRow]:
    b = _brand_or_404(brand)
    titles = _titles(db, b)
    rate_q = select(SalesRate).where(SalesRate.title_id.in_(list(titles))) if titles else None
    if rate_q is not None and year:
        rate_q = rate_q.where(SalesRate.year == year)
    rates = {r.id: r for r in db.scalars(rate_q)} if rate_q is not None else {}
    oq = select(RateOffer).where(RateOffer.brand == b.key)
    if year:
        oq = oq.where(RateOffer.year == year)
    offers = {o.id: o for o in db.scalars(oq)}
    rows = db.scalars(select(FieldChange).where(
        ((FieldChange.entity_type == "sales_rate") & FieldChange.entity_id.in_(list(rates) or [uuid.uuid4()]))
        | ((FieldChange.entity_type == "rate_offer") & FieldChange.entity_id.in_(list(offers) or [uuid.uuid4()]))
    ).order_by(FieldChange.changed_at.desc()).limit(200)).all()
    users = {u.id: u.name for u in db.scalars(select(User).where(User.id.in_({r.changed_by_user_id for r in rows if r.changed_by_user_id})))}
    out = []
    for c in rows:
        if c.entity_type == "sales_rate":
            r = rates.get(c.entity_id)
            what = f"{r.product} ({r.year})" if r else "A price"
        else:
            o = offers.get(c.entity_id)
            what = f"Offer: {o.label}" if o else "An offer"
        out.append(HistoryRow(id=c.id, entity_type=c.entity_type, entity_id=c.entity_id, what=what, field=FIELD_WORDS.get(c.field, c.field),
                              old_value=c.old_value, new_value=c.new_value, by=users.get(c.changed_by_user_id), at=c.changed_at,
                              can_put_back=c.entity_type == "sales_rate" and c.field in UNDOABLE and c.entity_id in rates))
    return out


UNDOABLE = {"price_gbp", "price_type", "unit", "specs", "notes", "product", "archived", "valid_until"}


def _parse(field: str, v: str | None):
    if v is None or v == "None":
        return None
    if field == "price_gbp":
        return float(v)
    if field == "archived":
        return v == "True"
    if field == "valid_until":
        return date.fromisoformat(v)
    return v


@router.post("/history/{cid}/put-back", response_model=RateItem)
def put_back(cid: uuid.UUID, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> RateItem:
    """Undo one change: sets the field back to what it was before - unless it has been changed again since."""
    c = db.get(FieldChange, cid)
    if not c or c.entity_type != "sales_rate":
        raise HTTPException(404, "Change not found")
    field = c.field
    if field not in UNDOABLE:
        raise HTTPException(422, "That change can't be put back automatically.")
    r = db.get(SalesRate, c.entity_id)
    if not r:
        raise HTTPException(404, "Price not found")
    brand = _rate_brand(db, r)
    _require_edit(db, identity, brand)
    current = getattr(r, field)
    cur_s = None if current is None else str(current)
    if cur_s != c.new_value and not (field == "price_gbp" and current is not None and c.new_value and float(current) == float(c.new_value)):
        raise HTTPException(409, "It has been changed again since - edit it directly instead.")
    before = _snapshot(r)
    setattr(r, field, _parse(field, c.old_value))
    if field == "price_type" and r.price_type == "poa":
        r.price_gbp = None
    record_field_changes(db, entity_type="sales_rate", entity_id=r.id, before=before, updates=_snapshot(r), changed_by_user_id=identity.user_uuid)
    db.commit()
    return _item(r, _titles(db, brand))


# ---- download -------------------------------------------------------------------------------

@router.get("/export.xlsx")
def export(year: int, brand: str | None = None, db: Session = Depends(get_db)) -> StreamingResponse:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for b in [_brand_or_404(brand)] if brand else B.BRANDS:
        ws = wb.create_sheet(b.short)
        ws.append([f"{b.name} - {year} rate card (prices before VAT)"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([])
        titles = _titles(db, b)
        rows = _rates(db, b, year)
        for s in B.SECTIONS:
            items = [r for r in rows if r.section == s.key]
            if not items:
                continue
            ws.append([s.label])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            ws.append(["Product", "Price", "Size / specs", "Notes", "Booked under"])
            for c in ws[ws.max_row]:
                c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="DDE4F0")
            for r in items:
                ws.append([r.product, price_label(r), r.specs or "", r.notes or "", titles[r.title_id].name if r.title_id in titles else ""])
            ws.append([])
        offers = list(db.scalars(select(RateOffer).where(RateOffer.brand == b.key, RateOffer.year == year).order_by(RateOffer.sort_order)))
        if offers:
            ws.append(["Offers & discounts"])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            for o in offers:
                ws.append([o.label, "", "", o.details or "", f"Until {o.valid_until:%d %b %Y}" if o.valid_until else ""])
        for col, w in zip("ABCDE", (44, 24, 32, 60, 30)):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"{B.get_brand(brand).name if brand else 'BMI'} rate card {year}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})
