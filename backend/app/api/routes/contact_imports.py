"""Import contacts from a spreadsheet or PDF (Act! feedback B4).

Steps the screen walks through, each one a call here:
  upload -> (columns are matched automatically) -> adjust matching/options ->
  review every row -> import -> undo if it was a mistake.

Nothing is written to the contact book until /commit.
"""
from __future__ import annotations

import io
import uuid
from datetime import datetime

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.routes.messaging import current_user
from app.contacts import fields as F
from app.contacts import import_engine as E
from app.contacts import import_mapper as M
from app.contacts import import_parse as P
from app.core.identity import Identity, get_identity
from app.db.session import get_db
from app.models import ContactImport, Publication, User

router = APIRouter(prefix="/contact-imports", tags=["contact import"])


# ---- targets ---------------------------------------------------------------------------

class TargetOut(BaseModel):
    key: str
    label: str
    group: str


def _targets(db: Session) -> list[TargetOut]:
    out = [TargetOut(key=t.key, label=t.label, group=t.group) for t in M.TARGETS]
    out += [TargetOut(key=f.key, label=f.label, group="Custom fields") for f in F.custom_field_defs(db)]
    return out


@router.get("/targets", response_model=list[TargetOut])
def targets(db: Session = Depends(get_db)) -> list[TargetOut]:
    """Every field a column can be sent to (the dropdown on the matching screen)."""
    return _targets(db)


def _custom_aliases(db: Session) -> dict[str, list[str]]:
    return {f.key: [f.label, f.column] for f in F.custom_field_defs(db)}


# ---- shapes ----------------------------------------------------------------------------

class ColumnOut(BaseModel):
    index: int
    header: str
    samples: list[str]
    field: str                 # what it is mapped to now ("skip" = not imported)
    name: str | None = None    # for new custom fields
    phone_type: str | None = None
    suggested: str | None      # what the matcher proposed
    level: str                 # high | medium | ambiguous | none | empty | skip | manual
    confidence: float
    reason: str
    candidates: list[dict]
    warning: str | None


class ImportOut(BaseModel):
    id: uuid.UUID
    filename: str
    kind: str
    status: str
    sheets: list[dict]
    sheet_name: str | None
    has_header: bool
    header_row: int
    raw_top: list[list[str]]
    preview_rows: list[list[str]]
    row_count: int
    columns: list[ColumnOut]
    options: dict
    excluded: list[int]
    notes: list[str]
    ai_read: bool
    problems: list[str]
    matched: int
    result: dict | None
    created_at: datetime
    by: str | None = None


class ImportRow(BaseModel):
    id: uuid.UUID
    filename: str
    status: str
    row_count: int
    created: int | None
    updated: int | None
    created_at: datetime
    by: str | None


def _problems(db: Session, imp: ContactImport) -> list[str]:
    out = []
    used: dict[str, int] = {}
    for i, spec in imp.mapping.items():
        f = spec.get("field")
        if f in (None, "skip"):
            continue
        if f == "new_custom":
            if not (spec.get("name") or "").strip():
                out.append(f"Name the new field for the column “{imp.headers[int(i)]}”.")
            continue
        if f in M.SINGLE or f.startswith("custom:"):
            if f in used:
                label = next((t.label for t in M.TARGETS if t.key == f), f)
                out.append(f"“{imp.headers[used[f]]}” and “{imp.headers[int(i)]}” are both set to {label}. Each field can only come from one column.")
            used[f] = int(i)
    names = {"first_name", "last_name", "full_name"}
    if not (names & set(used) or "email" in used or "email_2" in used):
        out.append("Match at least a name column (First name, Last name or Full name) or an Email column.")
    opts = imp.options or {}
    if not opts.get("source_db"):
        out.append("Choose which database these contacts belong to.")
    return out


def _out(db: Session, imp: ContactImport) -> ImportOut:
    sug = imp.suggestions or {}
    cols = []
    by_idx = {c["index"]: c for c in sug.get("columns", [])}
    for i, h in enumerate(imp.headers):
        spec = imp.mapping.get(str(i), {"field": "skip"})
        s = by_idx.get(i, {})
        field = spec.get("field") or "skip"
        same = field == (s.get("field") or "skip")
        cols.append(ColumnOut(
            index=i, header=h, samples=s.get("samples", []), field=field, name=spec.get("name"), phone_type=spec.get("type"),
            suggested=s.get("field"), level=s.get("level", "none") if same else "manual", confidence=s.get("confidence", 0.0),
            reason=s.get("reason", ""), candidates=s.get("candidates", []), warning=s.get("warning")))
    u = db.get(User, imp.created_by_user_id) if imp.created_by_user_id else None
    return ImportOut(
        id=imp.id, filename=imp.filename, kind=imp.file_kind, status=imp.status, sheets=imp.sheets, sheet_name=imp.sheet_name,
        has_header=imp.has_header, header_row=imp.header_row, raw_top=sug.get("raw_top", []), preview_rows=imp.rows[:5], row_count=imp.row_count, columns=cols,
        options={**E.DEFAULT_OPTIONS, **imp.options}, excluded=imp.excluded, notes=imp.notes, ai_read=bool(sug.get("ai_read")),
        problems=_problems(db, imp) if imp.status == "mapping" else [],
        matched=sum(1 for c in cols if c.field not in ("skip",)), result={k: v for k, v in imp.result.items() if k != "rows"} or None,
        created_at=imp.created_at, by=u.name if u else None)


def _get(db: Session, iid: uuid.UUID, user: User, identity: Identity) -> ContactImport:
    imp = db.get(ContactImport, iid)
    if not imp or (identity.is_known and identity.role not in ("admin", "data_manager") and imp.created_by_user_id != user.id):
        raise HTTPException(404, "Import not found")
    return imp


# ---- parsing + automatic matching --------------------------------------------------------

def _load(db: Session, imp: ContactImport, *, sheet: str | None = None, header_row: int | None = None, has_header: bool | None = None) -> None:
    """(Re)reads the file and proposes a column matching. Replaces any matching done so far."""
    try:
        kind, parsed = P.parse(imp.filename, imp.file_data or b"", sheet or imp.sheet_name)
    except P.ImportFileError as exc:
        raise HTTPException(422, str(exc))
    if header_row is None or has_header is None:
        det_row, det_has = P.detect_header(parsed.rows)
        header_row = det_row if header_row is None else header_row
        has_header = det_has if has_header is None else has_header
    headers, body = P.split_rows(parsed.rows, header_row, has_header)
    results = M.match_columns(headers, body, _custom_aliases(db))
    imp.file_kind, imp.sheets, imp.sheet_name = kind, parsed.sheets, parsed.sheet_name
    imp.header_row, imp.has_header = header_row, has_header
    imp.headers, imp.rows, imp.row_count = headers, body, len(body)
    imp.notes = parsed.notes + ([f"Skipped {header_row} line{'' if header_row == 1 else 's'} above the column headings."] if has_header and header_row else [])
    imp.suggestions = {
        "columns": [{"index": r.index, "field": r.field if r.field else "skip", "samples": r.samples, "level": r.level, "confidence": round(r.confidence, 2),
                     "reason": r.reason, "candidates": r.candidates, "warning": r.warning} for r in results],
        "raw_top": parsed.rows[:12], "ai_read": parsed.ai_read,
    }
    imp.mapping = {str(r.index): {"field": r.field if r.field else "skip"} for r in results}
    imp.excluded = []


def _allowed(identity: Identity, db: Session) -> list[str]:
    slugs = [p for p in db.scalars(select(Publication.slug))]
    allowed = identity.allowed_source_dbs()
    return slugs if allowed is None else [s for s in slugs if s in allowed]


# ---- endpoints -------------------------------------------------------------------------

@router.post("", response_model=ImportOut, status_code=201)
async def upload(file: UploadFile = File(...), source_db: str | None = Form(None), db: Session = Depends(get_db),
                 user: User = Depends(current_user), identity: Identity = Depends(get_identity)) -> ImportOut:
    data = await file.read()
    name = (file.filename or "contacts").strip()[:300]
    try:
        P.file_kind(name, data)
    except P.ImportFileError as exc:
        raise HTTPException(422, str(exc))
    imp = ContactImport(id=uuid.uuid4(), created_by_user_id=user.id, filename=name, file_kind="xlsx", file_data=data, status="mapping",
                        options={})
    choices = _allowed(identity, db)
    if source_db in choices:
        imp.options = {"source_db": source_db}
    elif len(choices) == 1:
        imp.options = {"source_db": choices[0]}
    _load(db, imp)
    db.add(imp)
    db.commit()
    return _out(db, imp)


@router.get("", response_model=list[ImportRow])
def recent(db: Session = Depends(get_db), user: User = Depends(current_user), identity: Identity = Depends(get_identity)) -> list[ImportRow]:
    q = select(ContactImport).where(ContactImport.status != "mapping").order_by(ContactImport.created_at.desc()).limit(20)
    if identity.is_known and identity.role not in ("admin", "data_manager"):
        q = q.where(ContactImport.created_by_user_id == user.id)
    out = []
    for i in db.scalars(q):
        u = db.get(User, i.created_by_user_id) if i.created_by_user_id else None
        out.append(ImportRow(id=i.id, filename=i.filename, status=i.status, row_count=i.row_count, created=i.result.get("created"),
                             updated=i.result.get("updated"), created_at=i.created_at, by=u.name if u else None))
    return out


@router.get("/template.xlsx")
def template() -> StreamingResponse:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contacts"
    ws.append(["First name", "Last name", "Job title", "Company", "Email", "Phone", "Mobile", "Address line 1", "Address line 2", "City", "County", "Postcode", "Country", "Notes"])
    ws.append(["Ann", "Lee", "Travel Counsellor", "Sunny Travel Ltd", "ann@sunnytravel.example", "020 7946 0001", "07700 900123", "1 High Street", "", "London", "", "SW1A 1AA", "United Kingdom", "Met at the trade show"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="DDE4F0")
    for col, w in zip("ABCDEFGHIJKLMN", (14, 14, 22, 24, 30, 16, 16, 22, 16, 14, 12, 12, 16, 28)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="BMI contact import template.xlsx"'})


@router.get("/{iid}", response_model=ImportOut)
def get_import(iid: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(current_user), identity: Identity = Depends(get_identity)) -> ImportOut:
    return _out(db, _get(db, iid, user, identity))


class SpecIn(BaseModel):
    field: str = "skip"
    name: str | None = Field(default=None, max_length=80)
    type: str | None = Field(default=None, max_length=30)


class PatchIn(BaseModel):
    sheet: str | None = None
    header_row: int | None = Field(default=None, ge=0, le=50)
    has_header: bool | None = None
    mapping: dict[str, SpecIn] | None = None
    options: dict | None = None
    excluded: list[int] | None = None


@router.patch("/{iid}", response_model=ImportOut)
def patch(iid: uuid.UUID, p: PatchIn, db: Session = Depends(get_db), user: User = Depends(current_user),
          identity: Identity = Depends(get_identity)) -> ImportOut:
    imp = _get(db, iid, user, identity)
    if imp.status != "mapping":
        raise HTTPException(409, "This import has already been done.")
    if p.sheet is not None or p.header_row is not None or p.has_header is not None:
        _load(db, imp, sheet=p.sheet, header_row=p.header_row, has_header=p.has_header)
    if p.mapping is not None:
        valid = {t.key for t in _targets(db)} | {"skip", "new_custom"}
        new = dict(imp.mapping)
        for idx, spec in p.mapping.items():
            if not idx.isdigit() or int(idx) >= len(imp.headers):
                raise HTTPException(422, "Unknown column.")
            if spec.field not in valid:
                raise HTTPException(422, f"“{spec.field}” isn't a field you can import into.")
            new[idx] = {k: v for k, v in spec.model_dump().items() if v}
        imp.mapping = new
    if p.options is not None:
        o = {**imp.options}
        for k, v in p.options.items():
            if k not in E.DEFAULT_OPTIONS:
                continue
            if k == "on_duplicate" and v not in E.ON_DUPLICATE:
                raise HTTPException(422, "Choose what to do with people already in the CRM.")
            if k == "source_db" and v and v not in _allowed(identity, db):
                raise HTTPException(403, "You can't import into that database.")
            o[k] = v
        imp.options = o
    if p.excluded is not None:
        imp.excluded = sorted(set(p.excluded))
    db.commit()
    return _out(db, imp)


@router.post("/{iid}/automap", response_model=ImportOut)
def automap(iid: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(current_user), identity: Identity = Depends(get_identity)) -> ImportOut:
    """Throws away manual changes and matches the columns again."""
    imp = _get(db, iid, user, identity)
    if imp.status != "mapping":
        raise HTTPException(409, "This import has already been done.")
    _load(db, imp, header_row=imp.header_row, has_header=imp.has_header)
    db.commit()
    return _out(db, imp)


class ReviewIn(BaseModel):
    status: str | None = None      # new | update | skip_existing | skip_repeat | error | warnings | possible | excluded
    q: str | None = None
    page: int = Field(default=1, ge=1)
    page_size: int = Field(default=50, ge=1, le=200)


class ReviewOut(BaseModel):
    totals: dict
    rows: list[dict]
    total_rows: int
    page: int
    page_size: int
    problems: list[str]


@router.post("/{iid}/review", response_model=ReviewOut)
def review(iid: uuid.UUID, r: ReviewIn, db: Session = Depends(get_db), user: User = Depends(current_user),
           identity: Identity = Depends(get_identity)) -> ReviewOut:
    """What will happen to every row, with the totals. Recalculated each time so it always matches the current matching and options."""
    imp = _get(db, iid, user, identity)
    problems = _problems(db, imp) if imp.status == "mapping" else []
    if imp.status == "mapping" and problems:
        return ReviewOut(totals={}, rows=[], total_rows=0, page=1, page_size=r.page_size, problems=problems)
    a = E.analyse(db, imp) if imp.status == "mapping" else {"totals": imp.result.get("totals", {}), "rows": imp.result.get("rows", [])}
    rows = a["rows"]
    if r.status == "warnings":
        rows = [x for x in rows if any(i["level"] == "warning" for i in x["issues"])]
    elif r.status == "possible":
        rows = [x for x in rows if x["match"] and x["match"]["kind"] == "possible"]
    elif r.status:
        rows = [x for x in rows if x["status"] == r.status]
    if r.q and r.q.strip():
        needle = r.q.strip().lower()
        rows = [x for x in rows if needle in " ".join(str(x.get(k) or "") for k in ("name", "email", "company", "phone")).lower()]
    start = (r.page - 1) * r.page_size
    return ReviewOut(totals=a["totals"], rows=rows[start:start + r.page_size], total_rows=len(rows), page=r.page, page_size=r.page_size, problems=[])


@router.post("/{iid}/commit", response_model=ImportOut)
def commit(iid: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(current_user), identity: Identity = Depends(get_identity)) -> ImportOut:
    imp = _get(db, iid, user, identity)
    if imp.status != "mapping":
        raise HTTPException(409, "This import has already been done.")
    problems = _problems(db, imp)
    if problems:
        raise HTTPException(422, problems[0])
    if imp.options.get("source_db") not in _allowed(identity, db):
        raise HTTPException(403, "You can't import into that database.")
    try:
        E.commit(db, imp, user.id)
    except HTTPException:
        raise
    except Exception as exc:
        db.rollback()
        raise HTTPException(500, "The import failed and nothing was saved. Please try again, or contact support if it keeps happening.")
    return _out(db, imp)


class UndoOut(BaseModel):
    deleted: int
    kept: int
    reverted: int


@router.post("/{iid}/undo", response_model=UndoOut)
def undo(iid: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(current_user), identity: Identity = Depends(get_identity)) -> UndoOut:
    imp = _get(db, iid, user, identity)
    if imp.status != "imported":
        raise HTTPException(409, "Only an import that has been done can be undone.")
    return UndoOut(**E.undo_import(db, imp))


@router.delete("/{iid}", status_code=204, response_model=None)
def discard(iid: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(current_user), identity: Identity = Depends(get_identity)) -> None:
    imp = _get(db, iid, user, identity)
    if imp.status != "mapping":
        raise HTTPException(409, "An import that has been done is kept as a record.")
    db.delete(imp)
    db.commit()


@router.get("/{iid}/report.xlsx")
def report(iid: uuid.UUID, db: Session = Depends(get_db), user: User = Depends(current_user), identity: Identity = Depends(get_identity)) -> StreamingResponse:
    """Row-by-row outcome (and problems) as a spreadsheet - fix the file and import it again."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    imp = _get(db, iid, user, identity)
    rows = imp.result.get("rows") if imp.status != "mapping" else E.analyse(db, imp)["rows"]
    label = {"new": "Will be added", "update": "Will update existing", "skip_existing": "Already in the CRM - skipped", "skip_repeat": "Repeated in the file - skipped",
             "error": "Can't import", "excluded": "Left out by you"}
    done = {"new": "Added", "update": "Updated existing"}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Import report"
    ws.append(["Row", "What happened", "Name", "Email", "Company", "Phone", "Already in the CRM as", "Notes"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="DDE4F0")
    for x in rows or []:
        ws.append([x["n"], (done if imp.status == "imported" else label).get(x["status"], label.get(x["status"], x["status"])), x["name"], x["email"], x["company"], x["phone"],
                   (x["match"] or {}).get("name"), "; ".join(i["text"] for i in x["issues"])])
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="Import report.xlsx"'})
