"""Commission: monthly statements per salesperson, plans, new-business decisions,
new contract-publishing guides and event cost sign-off. See app/sales/commission.py.

Salespeople see only their own statements. Admins manage plans, settings and
approval; admins and data managers decide new business and mark guides.
"""
from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timezone

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.identity import Identity, get_identity
from app.db.session import get_db
from app.models import CommissionRule, CommissionStatement, SalesEdition, SalesOrder, SalesRep, SalesTitle, User
from app.roles import CAN_USE_AUTOMATIONS
from app.sales import commission as cm
from app.sales import commission_seed as seed
from app.services.field_audit import record_field_changes

router = APIRouter(prefix="/commission", tags=["commission"])


def _staff(identity: Identity) -> bool:
    return not identity.is_known or identity.role in CAN_USE_AUTOMATIONS


def _admin(identity: Identity) -> bool:
    return not identity.is_known or identity.role == "admin"


def _need_admin(identity: Identity) -> None:
    if not _admin(identity):
        raise HTTPException(403, "Only administrators can change commission plans or approve statements.")


def _need_staff(identity: Identity) -> None:
    if not _staff(identity):
        raise HTTPException(403, "Only administrators and data managers can do this.")


def _my_rep(db: Session, identity: Identity) -> SalesRep | None:
    return db.scalars(select(SalesRep).where(SalesRep.user_id == identity.user_uuid)).first() if identity.user_uuid else None


def _rep_or_404(db: Session, identity: Identity, rep_id: uuid.UUID) -> SalesRep:
    rep = db.get(SalesRep, rep_id)
    if not rep:
        raise HTTPException(404, "Salesperson not found")
    if not _staff(identity):
        me = _my_rep(db, identity)
        if not me or me.id != rep.id:
            raise HTTPException(404, "Salesperson not found")
    return rep


def _rep_ref(r: SalesRep) -> dict:
    return {"id": str(r.id), "code": r.code, "name": r.name, "active": r.active, "has_login": r.user_id is not None}


def _period_ok(period: str) -> str:
    try:
        datetime.strptime(period, "%Y-%m")
    except ValueError:
        raise HTTPException(422, "The month should look like 2026-10.")
    return period


# ---- statements --------------------------------------------------------------------------

@router.get("/overview")
def overview(year: int | None = None, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    """Each salesperson's commission month by month for a year (approved months as paid, others as they stand)."""
    year = year or date.today().year
    staff = _staff(identity)
    me = None if staff else _my_rep(db, identity)
    if not staff and not me:
        return {"year": year, "reps": [], "scoped_to_me": True, "can_manage": False, "has_plans": seed.plan_count(db) > 0}
    ctx = cm.Ctx.build(db)
    reps = [me] if me else db.scalars(select(SalesRep).where(SalesRep.active.is_(True)).order_by(SalesRep.name)).all()
    approved = {(s.rep_id, s.period): s for s in db.scalars(select(CommissionStatement))}
    last = date.today().month if year == date.today().year else 12
    out = []
    for rep in reps:
        months = []
        for m in range(1, last + 1):
            period = f"{year}-{m:02d}"
            st = approved.get((rep.id, period))
            if st:
                months.append({"period": period, "total_gbp": float(st.total_gbp), "approved": True, "checks": 0, "bookings": len(st.snapshot.get("lines", []))})
            else:
                c = cm.core(ctx, rep, period)
                months.append({"period": period, "total_gbp": c["totals"]["core_gbp"], "approved": False, "checks": c["checks"], "bookings": len(c["lines"])})
        if staff and not any(x["bookings"] or x["total_gbp"] for x in months):
            continue
        out.append({"rep": _rep_ref(rep), "months": months, "year_total_gbp": round(sum(x["total_gbp"] for x in months), 2),
                    "checks": sum(x["checks"] for x in months), "has_plan": bool(cm._rules(db, rep))})
    return {"year": year, "reps": out, "scoped_to_me": not staff, "can_manage": _admin(identity), "has_plans": seed.plan_count(db) > 0}


@router.get("/statement")
def get_statement(rep_id: uuid.UUID, period: str, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    rep = _rep_or_404(db, identity, rep_id)
    ctx = cm.Ctx.build(db)
    s = cm.statement(ctx, rep, _period_ok(period))
    if s.get("approved") and s["approved"].get("by"):
        u = db.get(User, uuid.UUID(s["approved"]["by"]))
        s["approved"]["by_name"] = u.name if u else None
    db.commit()  # settings row may have been created
    return {**s, "rep": _rep_ref(rep), "can_approve": _admin(identity) and not s.get("approved") and period < date.today().strftime("%Y-%m"),
            "can_decide": _staff(identity) and not s.get("approved"), "plan": [_rule_out(r) for r in cm._rules(db, rep)],
            "settings": _settings_out(ctx.st)}


class ApproveIn(BaseModel):
    rep_id: uuid.UUID
    period: str


@router.post("/statement/approve")
def approve_statement(payload: ApproveIn, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    _need_admin(identity)
    rep = _rep_or_404(db, identity, payload.rep_id)
    try:
        cm.approve(cm.Ctx.build(db), rep, _period_ok(payload.period), identity.user_uuid)
    except ValueError as exc:
        raise HTTPException(409, str(exc)) from exc
    db.commit()
    return get_statement(rep.id, payload.period, db, identity)


@router.get("/statement.xlsx")
def statement_xlsx(rep_id: uuid.UUID, period: str, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> Response:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    rep = _rep_or_404(db, identity, rep_id)
    s = cm.statement(cm.Ctx.build(db), rep, _period_ok(period))
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    month = datetime.strptime(period, "%Y-%m").strftime("%B %Y")
    ws.append([f"Commission statement: {rep.name}, {month}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Approved" if s.get("approved") else "Not approved yet - figures can still change"])
    ws.append([])
    head = ["Product group", "Issue / event", "Published", "Client", "Booking value £", "Their share £", "Rate", "Commission £",
            "New business", "New business rate", "New business £", "Total £", "Why"]
    ws.append(head)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for ln in s["lines"]:
        ws.append([ln["group"], ln["edition"], ln["publication"], ln["client"], ln["value_gbp"], ln["share_gbp"], ln["base_rate"], ln["base_gbp"],
                   {"new": "Yes", "returning": "No", "check": "Needs a decision"}[ln["new_business"]], ln["new_business_rate"] or None,
                   ln["new_business_gbp"] or None, ln["commission_gbp"], ln["new_business_reason"]])
        ws.cell(ws.max_row, 7).number_format = "0.0%"
        ws.cell(ws.max_row, 10).number_format = "0.0%"
    ws.append([])
    for b in s["bonuses"]:
        ws.append(["Bonus", b["label"], None, None, None, None, None, None, None, None, None, b["amount_gbp"], b["reason"]])
    for e in s["events"]:
        ws.append(["Event profit share", e["edition"], e["event_date"], None, e["income_gbp"], None, e["rate"], None, None, None, None, e["amount_gbp"],
                   f"£{e['income_gbp']:,.2f} income - £{e['costs_gbp']:,.2f} costs = £{e['profit_gbp']:,.2f} profit"])
    for a in s.get("adjustments", []):
        ws.append(["Adjustment", a["label"], None, None, None, None, None, None, None, None, None, a["amount_gbp"], None])
    ws.append([])
    ws.append(["Total", None, None, None, None, s["totals"]["share_gbp"], None, None, None, None, None, s["totals"].get("total_gbp", s["totals"]["core_gbp"])])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEFGHIJKLM", (34, 26, 12, 30, 14, 13, 8, 13, 15, 10, 13, 12, 70)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    name = f"Commission {rep.name} {period}.xlsx"
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


# ---- new business decisions ----------------------------------------------------------------

@router.get("/orders/{order_id}/new-business")
def booking_new_business(order_id: uuid.UUID, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    o = db.get(SalesOrder, order_id)
    if not o:
        raise HTTPException(404, "Booking not found")
    ctx = cm.Ctx.build(db)
    v = ctx.ix.classify(o, db.get(SalesEdition, o.edition_id))
    db.commit()
    return {"status": v.status, "reason": v.reason, "decided_by": v.decided_by, "last": v.last, "similar": v.similar, "can_decide": _staff(identity)}


class DecisionIn(BaseModel):
    decision: str | None = Field(default=None, pattern="^(new|returning)$")  # None = work it out from history again
    reason: str | None = Field(default=None, max_length=300)


@router.put("/orders/{order_id}/new-business")
def decide_new_business(order_id: uuid.UUID, payload: DecisionIn, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    _need_staff(identity)
    o = db.get(SalesOrder, order_id)
    if not o:
        raise HTTPException(404, "Booking not found")
    before = {"new_business": None if o.new_business_override is None else ("new" if o.new_business_override else "returning")}
    o.new_business_override = None if payload.decision is None else payload.decision == "new"
    o.new_business_reason = (payload.reason or "").strip() or None if payload.decision else None
    o.new_business_set_by_user_id = identity.user_uuid
    o.new_business_set_at = datetime.now(timezone.utc)
    record_field_changes(db, entity_type="sales_order", entity_id=o.id, before=before,
                         updates={"new_business": payload.decision or "worked out from history"}, changed_by_user_id=identity.user_uuid)
    db.commit()
    return booking_new_business(order_id, db, identity)


# ---- plans and settings ----------------------------------------------------------------------

class RuleIn(BaseModel):
    rep_id: uuid.UUID
    name: str = Field(min_length=1, max_length=160)
    title_slugs: list[str] = Field(default_factory=list)
    base_rate: float = Field(ge=0, le=1)
    new_business_rate: float = Field(default=0, ge=0, le=1)
    new_business_rate_change_on: date | None = None
    new_business_rate_after: float | None = Field(default=None, ge=0, le=1)
    new_guide_bonus_gbp: float | None = Field(default=None, ge=0)
    threshold_bonus_gbp: float | None = Field(default=None, ge=0)
    threshold_gbp: float | None = Field(default=None, ge=0)
    new_client_bonus_gbp: float | None = Field(default=None, ge=0)
    event_profit_rate: float | None = Field(default=None, ge=0, le=1)
    valid_from: date | None = None
    valid_until: date | None = None
    notes: str | None = None


def _f(v):
    return None if v is None else float(v)


def _rule_out(r: CommissionRule) -> dict:
    return {"id": str(r.id), "rep_id": str(r.rep_id), "name": r.name, "title_slugs": r.title_slugs or [], "base_rate": float(r.base_rate),
            "new_business_rate": float(r.new_business_rate), "new_business_rate_change_on": r.new_business_rate_change_on.isoformat() if r.new_business_rate_change_on else None,
            "new_business_rate_after": _f(r.new_business_rate_after), "new_guide_bonus_gbp": _f(r.new_guide_bonus_gbp),
            "threshold_bonus_gbp": _f(r.threshold_bonus_gbp), "threshold_gbp": _f(r.threshold_gbp), "new_client_bonus_gbp": _f(r.new_client_bonus_gbp),
            "event_profit_rate": _f(r.event_profit_rate), "valid_from": r.valid_from.isoformat() if r.valid_from else None,
            "valid_until": r.valid_until.isoformat() if r.valid_until else None, "notes": r.notes}


def _settings_out(st) -> dict:
    return {"earned_on": st.earned_on, "lookback_months": st.lookback_months, "first_deal_days": st.first_deal_days,
            "event_profit_basis": st.event_profit_basis}


@router.get("/plans")
def plans(db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    _need_staff(identity)
    st = cm.settings(db)
    db.commit()
    reps = db.scalars(select(SalesRep).order_by(SalesRep.active.desc(), SalesRep.name)).all()
    rules = db.scalars(select(CommissionRule).order_by(CommissionRule.sort_order, CommissionRule.name)).all()
    return {"reps": [{**_rep_ref(r), "default_rate": float(r.commission_rate or 0), "rules": [_rule_out(x) for x in rules if x.rep_id == r.id]} for r in reps],
            "titles": [{"slug": t.slug, "name": t.name} for t in db.scalars(select(SalesTitle).order_by(SalesTitle.sort_order))],
            "settings": _settings_out(st), "can_edit": _admin(identity), "seed_available": seed.plan_count(db) == 0}


@router.post("/plans/load")
def load_plans(replace: bool = False, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    _need_admin(identity)
    n = seed.load_structure(db, replace=replace)
    db.commit()
    return {"rules": n}


def _check_titles(db: Session, slugs: list[str]) -> None:
    known = {t.slug for t in db.scalars(select(SalesTitle))}
    bad = [s for s in slugs if s not in known]
    if bad:
        raise HTTPException(422, f"Unknown title: {', '.join(bad)}")


@router.post("/rules", status_code=201)
def add_rule(payload: RuleIn, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    _need_admin(identity)
    if not db.get(SalesRep, payload.rep_id):
        raise HTTPException(422, "Unknown salesperson")
    _check_titles(db, payload.title_slugs)
    r = CommissionRule(id=uuid.uuid4(), **payload.model_dump())
    db.add(r)
    db.commit()
    return _rule_out(r)


@router.patch("/rules/{rule_id}")
def edit_rule(rule_id: uuid.UUID, payload: RuleIn, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    _need_admin(identity)
    r = db.get(CommissionRule, rule_id)
    if not r:
        raise HTTPException(404, "Rule not found")
    _check_titles(db, payload.title_slugs)
    for k, v in payload.model_dump().items():
        setattr(r, k, v)
    db.commit()
    return _rule_out(r)


@router.delete("/rules/{rule_id}", status_code=204, response_model=None)
def delete_rule(rule_id: uuid.UUID, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> None:
    _need_admin(identity)
    r = db.get(CommissionRule, rule_id)
    if r:
        db.delete(r)
        db.commit()


class SettingsIn(BaseModel):
    earned_on: str = Field(pattern="^(publication|booked)$")
    lookback_months: int = Field(ge=1, le=120)
    first_deal_days: int = Field(ge=0, le=366)
    event_profit_basis: str = Field(pattern="^(all|own)$")


@router.put("/settings")
def save_settings(payload: SettingsIn, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    _need_admin(identity)
    st = cm.settings(db)
    for k, v in payload.model_dump().items():
        setattr(st, k, v)
    st.updated_at = datetime.now(timezone.utc)
    db.commit()
    return _settings_out(st)


# ---- editions: new contract-publishing guides, event cost sign-off -----------------------------

def _edition(db: Session, eid: uuid.UUID) -> SalesEdition:
    ed = db.get(SalesEdition, eid)
    if not ed:
        raise HTTPException(404, "Edition not found")
    return ed


def edition_commission(db: Session, ed: SalesEdition, identity: Identity) -> dict:
    t = db.get(SalesTitle, ed.title_id)
    rules = db.scalars(select(CommissionRule)).all()
    profit_reps = {r.rep_id for r in rules if r.event_profit_rate and t.slug in (r.title_slugs or [])}
    guide_title = any(r.new_guide_bonus_gbp and t.slug in (r.title_slugs or []) for r in rules)
    me = _my_rep(db, identity)
    users = {u.id: u.name for u in db.scalars(select(User).where(User.id.in_([x for x in (ed.costs_final_by_user_id, ed.costs_signed_off_by_user_id) if x])))}
    rep = db.get(SalesRep, ed.new_contract_rep_id) if ed.new_contract_rep_id else None
    return {
        "new_contract_guide": ed.new_contract_guide, "new_contract_rep": _rep_ref(rep) if rep else None, "guide_title": guide_title,
        "profit_share": bool(profit_reps), "profit_share_reps": [_rep_ref(r) for r in db.scalars(select(SalesRep).where(SalesRep.id.in_(profit_reps)))] if profit_reps else [],
        "costs_final_at": ed.costs_final_at.isoformat() if ed.costs_final_at else None, "costs_final_by": users.get(ed.costs_final_by_user_id),
        "costs_signed_off_at": ed.costs_signed_off_at.isoformat() if ed.costs_signed_off_at else None, "costs_signed_off_by": users.get(ed.costs_signed_off_by_user_id),
        "can_mark_guide": _staff(identity), "can_mark_final": _staff(identity) or bool(me and me.id in profit_reps), "can_sign_off": _admin(identity),
    }


def can_edit_costs(db: Session, ed: SalesEdition, identity: Identity) -> None:
    """Costs are locked once signed off; until then admins, data managers and the person whose profit share it is can edit them."""
    if ed.costs_signed_off_at:
        raise HTTPException(409, "These costs are signed off for commission. An administrator can reopen them first.")
    if _staff(identity):
        return
    info = edition_commission(db, ed, identity)
    if not info["can_mark_final"]:
        raise HTTPException(403, "Only administrators, data managers and the event's organiser can change its costs.")


@router.get("/editions/{edition_id}")
def get_edition(edition_id: uuid.UUID, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    return edition_commission(db, _edition(db, edition_id), identity)


class GuideIn(BaseModel):
    on: bool
    rep_id: uuid.UUID | None = None


@router.put("/editions/{edition_id}/new-guide")
def mark_new_guide(edition_id: uuid.UUID, payload: GuideIn, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    _need_staff(identity)
    ed = _edition(db, edition_id)
    if payload.on and not payload.rep_id:
        raise HTTPException(422, "Choose who sold the guide.")
    if payload.rep_id and not db.get(SalesRep, payload.rep_id):
        raise HTTPException(422, "Unknown salesperson")
    ed.new_contract_guide, ed.new_contract_rep_id = payload.on, payload.rep_id if payload.on else None
    db.commit()
    return edition_commission(db, ed, identity)


class FlagIn(BaseModel):
    on: bool


@router.post("/editions/{edition_id}/costs-final")
def mark_costs_final(edition_id: uuid.UUID, payload: FlagIn, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    ed = _edition(db, edition_id)
    info = edition_commission(db, ed, identity)
    if not info["can_mark_final"]:
        raise HTTPException(403, "Only administrators, data managers and the event's organiser can mark its costs final.")
    if ed.costs_signed_off_at:
        raise HTTPException(409, "These costs are already signed off.")
    ed.costs_final_at = datetime.now(timezone.utc) if payload.on else None
    ed.costs_final_by_user_id = identity.user_uuid if payload.on else None
    db.commit()
    return edition_commission(db, ed, identity)


@router.post("/editions/{edition_id}/costs-signoff")
def sign_off_costs(edition_id: uuid.UUID, payload: FlagIn, db: Session = Depends(get_db), identity: Identity = Depends(get_identity)) -> dict:
    _need_admin(identity)
    ed = _edition(db, edition_id)
    if payload.on and not ed.costs_final_at:
        raise HTTPException(409, "The costs need to be marked final before they're signed off.")
    ed.costs_signed_off_at = datetime.now(timezone.utc) if payload.on else None
    ed.costs_signed_off_by_user_id = identity.user_uuid if payload.on else None
    if not payload.on:
        ed.costs_final_at = ed.costs_final_by_user_id = None  # reopened: they're editable again
    db.commit()
    return edition_commission(db, ed, identity)
